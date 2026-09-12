from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException, WebSocket, WebSocketDisconnect, Body
from fastapi.responses import StreamingResponse, Response
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
import uuid
from datetime import datetime, timezone
import csv
import io
import pandas as pd
import json
import re
import ast
from pydantic import BaseModel


# ---------------------------------------------------------------------------
# DSL template AST safety validator
# ---------------------------------------------------------------------------
# Generated DSL templates are executed via exec(). The template scaffolding
# itself is trusted (we generate it), but it embeds user-supplied DSL/Custom
# Code lines verbatim. Without validation, a Custom Code step could write
# something like  __import__('os').system('rm -rf /')  and bypass the small
# 6-name builtin blacklist used at the exec sites.
#
# This validator walks the AST of the *full generated template* and rejects:
#   1. Any `import` / `from ... import` whose module is not in
#      _ALLOWED_IMPORT_MODULES (the fixed set the scaffolding actually uses).
#   2. Any reference to dunder names (e.g. __import__, __class__,
#      __subclasses__, __globals__, __builtins__, __getattribute__, __mro__,
#      __code__, __bases__, __dict__) either as a bare Name or as an
#      attribute access. A small allowlist of safe dunders the scaffolding
#      itself uses (__file__, __name__) is permitted.
#
# Raised errors are surfaced to the caller as DSLSecurityError so existing
# error-handling code paths can present a clear message to the user.
# ---------------------------------------------------------------------------

class DSLSecurityError(Exception):
    """Raised when generated DSL code contains a forbidden construct."""


_ALLOWED_IMPORT_MODULES = frozenset({
    'sys',
    'os',
    'json',
    'datetime',
    'inspect',
    'backend.dsl_functions',
    'dsl_functions',
})

_ALLOWED_DUNDER_NAMES = frozenset({
    '__file__',
    '__name__',
})


def _validate_template_ast(source: str, label: str = '<dsl_template>') -> None:
    """Validate a generated DSL template before exec.

    Raises DSLSecurityError if the AST contains a forbidden import or any
    dunder reference outside the small allowlist.
    """
    try:
        tree = ast.parse(source)
    except SyntaxError:
        # Let the caller's existing syntax-error handling produce the message.
        return

    def _is_forbidden_dunder(name: str) -> bool:
        if not isinstance(name, str):
            return False
        if not (name.startswith('__') and name.endswith('__') and len(name) > 4):
            return False
        return name not in _ALLOWED_DUNDER_NAMES

    for node in ast.walk(tree):
        # Block disallowed `import x` / `import x.y`
        if isinstance(node, ast.Import):
            for alias in node.names:
                root = (alias.name or '').split('.')[0]
                if alias.name not in _ALLOWED_IMPORT_MODULES and root not in _ALLOWED_IMPORT_MODULES:
                    raise DSLSecurityError(
                        f"Disallowed import '{alias.name}' in {label}. "
                        f"Custom Code may not import arbitrary modules."
                    )
            continue

        # Block disallowed `from x import ...`
        if isinstance(node, ast.ImportFrom):
            module = node.module or ''
            root = module.split('.')[0]
            if module not in _ALLOWED_IMPORT_MODULES and root not in _ALLOWED_IMPORT_MODULES:
                raise DSLSecurityError(
                    f"Disallowed import 'from {module} import ...' in {label}. "
                    f"Custom Code may not import arbitrary modules."
                )
            continue

        # Block bare references to dunder names (e.g. __import__, __builtins__)
        if isinstance(node, ast.Name) and _is_forbidden_dunder(node.id):
            raise DSLSecurityError(
                f"Use of '{node.id}' is not allowed in DSL Custom Code."
            )

        # Block dunder attribute access (e.g. obj.__class__.__subclasses__())
        if isinstance(node, ast.Attribute) and _is_forbidden_dunder(node.attr):
            raise DSLSecurityError(
                f"Access to attribute '{node.attr}' is not allowed in DSL Custom Code."
            )


# ---------------------------------------------------------------------------
# Sandbox builtins for DSL exec sites
# ---------------------------------------------------------------------------
# Every generated template is run via exec() with a restricted __builtins__.
# Beyond removing the obvious code-exec / file / IO escape hatches, we replace
# __import__ with a guard bound to _ALLOWED_IMPORT_MODULES. Previously __import__
# was left fully available, so a bypass of the AST validator could reach
# `__import__('os').system(...)`. The scaffolding's legitimate `import json` /
# `import datetime` still work because those modules are on the allow-list.
# ---------------------------------------------------------------------------

_SANDBOX_BLOCKED_BUILTINS = ('exec', 'eval', 'compile', 'open', 'input', 'breakpoint')


def _make_sandbox_builtins() -> dict:
    """Return the restricted __builtins__ mapping used at every DSL exec site."""
    if isinstance(__builtins__, dict):
        base = dict(__builtins__)
    else:
        base = {k: getattr(__builtins__, k) for k in dir(__builtins__)}
    real_import = base.get('__import__')
    for _k in _SANDBOX_BLOCKED_BUILTINS:
        base.pop(_k, None)

    def _guarded_import(name, _globals=None, _locals=None, fromlist=(), level=0):
        root = (name or '').split('.')[0]
        if name in _ALLOWED_IMPORT_MODULES or root in _ALLOWED_IMPORT_MODULES:
            return real_import(name, _globals, _locals, fromlist, level)
        raise ImportError(
            f"Import of '{name}' is not permitted in the DSL execution sandbox."
        )

    if real_import is not None:
        base['__import__'] = _guarded_import
    return base


# ---------------------------------------------------------------------------
# DSL *user code* AST safety validator (strict)
# ---------------------------------------------------------------------------
# The template scaffolding we generate is trusted and legitimately uses
# `globals().update(...)`, `getattr(...)`, `hasattr(...)`, `type(...)`, etc.
# User-supplied DSL/Custom Code, however, has no business calling those —
# they are well-known Python sandbox-escape primitives. This validator runs
# only over the user-code section (after DSL→Python translation, before
# wrapping in the scaffolding) so we can be aggressive without breaking the
# scaffolding.
#
# Rejects:
#   * any `import` / `from ... import`
#   * any reference to or attribute named __dunder__ (allowlist: __file__,
#     __name__) — same rule as the template-level validator
#   * any subscript with a dunder string literal:  obj['__class__']
#   * walrus assignments to a dunder target
#   * calls to known dangerous builtins, by call-name OR attribute-name:
#       eval, exec, compile, open, input, breakpoint, __import__,
#       getattr, setattr, delattr, globals, locals, vars
# ---------------------------------------------------------------------------

_USER_FORBIDDEN_CALL_NAMES = frozenset({
    'eval', 'exec', 'compile',
    'open', 'input', 'breakpoint',
    '__import__',
    'getattr', 'setattr', 'delattr',
    'globals', 'locals', 'vars',
})


def _validate_dsl_user_code(user_python_body: str, label: str = '<dsl_user_code>') -> None:
    """Validate the user-supplied (post-translation) DSL python body.

    `user_python_body` is the assembled, indented body that will be embedded
    inside `def process_standalone():` / `def process_event_data():`. We wrap
    it in a synthetic function so the indentation parses cleanly.
    """
    if not user_python_body or not user_python_body.strip():
        return

    # Detect the leading indent so we can wrap correctly. The two callers
    # produce 4-space (standalone) or 8-space (multi-event) indentation, but
    # we don't need to know which — wrapping inside a fresh `def` re-anchors
    # the indentation as long as every non-blank line is indented at least
    # once.
    wrapped = "def __dsl_user__():\n" + user_python_body
    try:
        tree = ast.parse(wrapped)
    except SyntaxError:
        # Let the existing syntax-error path produce the user-facing message.
        return

    def _is_forbidden_dunder_local(name):
        if not isinstance(name, str):
            return False
        if not (name.startswith('__') and name.endswith('__') and len(name) > 4):
            return False
        return name not in _ALLOWED_DUNDER_NAMES

    for node in ast.walk(tree):
        # No imports in user code.
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            raise DSLSecurityError(
                f"'import' statements are not allowed in DSL Custom Code."
            )

        # Dunder name reference.
        if isinstance(node, ast.Name) and _is_forbidden_dunder_local(node.id):
            raise DSLSecurityError(
                f"Use of '{node.id}' is not allowed in DSL Custom Code."
            )

        # Dunder attribute access (obj.__class__).
        if isinstance(node, ast.Attribute) and _is_forbidden_dunder_local(node.attr):
            raise DSLSecurityError(
                f"Access to attribute '.{node.attr}' is not allowed in DSL Custom Code."
            )

        # Dunder subscript with a string constant: obj['__class__'].
        if isinstance(node, ast.Subscript):
            slice_node = getattr(node, 'slice', None)
            # Python 3.9+: slice is the expression directly.
            const_val = None
            if isinstance(slice_node, ast.Constant):
                const_val = slice_node.value
            # Python <=3.8 uses ast.Index wrapping a Constant.
            elif hasattr(ast, 'Index') and isinstance(slice_node, ast.Index):
                inner = getattr(slice_node, 'value', None)
                if isinstance(inner, ast.Constant):
                    const_val = inner.value
            if isinstance(const_val, str) and _is_forbidden_dunder_local(const_val):
                raise DSLSecurityError(
                    f"Subscripting with '{const_val}' is not allowed in DSL Custom Code."
                )

        # Walrus assignment to a dunder target: (__import__ := f).
        if hasattr(ast, 'NamedExpr') and isinstance(node, ast.NamedExpr):
            tgt = node.target
            if isinstance(tgt, ast.Name) and _is_forbidden_dunder_local(tgt.id):
                raise DSLSecurityError(
                    f"Walrus assignment to '{tgt.id}' is not allowed in DSL Custom Code."
                )

        # Direct or attribute calls to dangerous builtins:
        #   getattr(obj, '__class__')   -> blocked by call-name
        #   x.getattr(...)              -> also blocked (defensive)
        if isinstance(node, ast.Call):
            func = node.func
            fname = None
            if isinstance(func, ast.Name):
                fname = func.id
            elif isinstance(func, ast.Attribute):
                fname = func.attr
            if fname in _USER_FORBIDDEN_CALL_NAMES:
                raise DSLSecurityError(
                    f"Call to '{fname}()' is not allowed in DSL Custom Code."
                )


def _normalize_ingest_date_value(value):
    """Normalize a value (scalar or list or JSON-list-string) to yyyy-mm-dd or list of such strings."""
    if value is None:
        return ''
    # Lists -> normalize each
    if isinstance(value, list):
        out = []
        for v in value:
            try:
                nv = normalize_date(v)
            except Exception:
                nv = ''
            if nv:
                out.append(nv)
        return out

    s = value
    # Try parsing JSON arrays
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return ''
        try:
            parsed = json.loads(s)
            if isinstance(parsed, list):
                return [normalize_date(p) for p in parsed if normalize_date(p)]
        except Exception:
            pass
        # Delimited lists
        if ',' in s or ';' in s or '|' in s:
            parts = [p.strip() for p in re.split('[,;|]', s) if p.strip()]
            return [normalize_date(p) for p in parts if normalize_date(p)]

    # Fallback: normalize scalar
    try:
        return normalize_date(s)
    except Exception:
        return ''


def _sort_activity_rows(rows):
    """Enforce canonical activity-data ordering:
        instrumentid ASC, postingdate ASC, effectivedate ASC, subinstrumentid ASC

    Applied at every ingestion entry point (direct upload + JSON import
    transform) and again at rule-step execution, so every step inside a rule
    (Schedule, Condition, Iteration, Calculation, Custom Code, Create
    Transaction) sees the same deterministic order.

    NOTE: must only be called for ACTIVITY data. Reference / custom / static
    data has no instrumentid/postingdate/effectivedate/subinstrumentid and
    must be passed through untouched.
    """
    if not isinstance(rows, list) or len(rows) <= 1:
        return rows

    def _ci(row, name):
        if not isinstance(row, dict):
            return ''
        if name in row:
            v = row[name]
        else:
            lname = name.lower()
            v = ''
            for k, val in row.items():
                if str(k).lower() == lname:
                    v = val
                    break
        if v is None:
            return ''
        return str(v)

    try:
        rows.sort(key=lambda r: (
            _ci(r, 'instrumentid'),
            _ci(r, 'postingdate'),
            _ci(r, 'effectivedate'),
            _ci(r, 'subinstrumentid') or '1',
        ))
    except Exception:
        # Never let sorting mask an ingestion / execution error.
        pass
    return rows
# AI provider abstraction layer
try:
    from backend.ai_providers import (
        get_provider, PROVIDER_INFO, build_agent_context,
        encrypt_key, decrypt_key, AIError,
    )
except Exception:
    try:
        from ai_providers import (
            get_provider, PROVIDER_INFO, build_agent_context,
            encrypt_key, decrypt_key, AIError,
        )
    except Exception:
        from .ai_providers import (
            get_provider, PROVIDER_INFO, build_agent_context,
            encrypt_key, decrypt_key, AIError,
        )

# Autonomous agent runtime (tools + plan/act/observe loop)
try:
    from backend.agent import (
        run_agent as agent_run, submit_approval as agent_submit_approval,
        cancel_run as agent_cancel_run, configure_bridge as agent_configure_bridge,
        DESTRUCTIVE_TOOLS as AGENT_DESTRUCTIVE_TOOLS,
    )
except Exception:
    try:
        from agent import (
            run_agent as agent_run, submit_approval as agent_submit_approval,
            cancel_run as agent_cancel_run, configure_bridge as agent_configure_bridge,
            DESTRUCTIVE_TOOLS as AGENT_DESTRUCTIVE_TOOLS,
        )
    except Exception:
        from .agent import (
            run_agent as agent_run, submit_approval as agent_submit_approval,
            cancel_run as agent_cancel_run, configure_bridge as agent_configure_bridge,
            DESTRUCTIVE_TOOLS as AGENT_DESTRUCTIVE_TOOLS,
        )
# Support running in different execution contexts: prefer package import, fallback to module-level
try:
    from backend.dsl_functions import DSL_FUNCTIONS, DSL_FUNCTION_METADATA, normalize_date
except Exception:
    try:
        from dsl_functions import DSL_FUNCTIONS, DSL_FUNCTION_METADATA, normalize_date
    except Exception:
        # Last resort: try relative import (works when executed as package)
        from .dsl_functions import DSL_FUNCTIONS, DSL_FUNCTION_METADATA, normalize_date

try:
    from bson import ObjectId
except Exception:
    ObjectId = None

# Load configuration
try:
    from backend.config import settings
except Exception:
    try:
        from config import settings
    except Exception:
        from .config import settings

ROOT_DIR = Path(__file__).parent

# MongoDB connection
client = AsyncIOMotorClient(settings.mongo_url, serverSelectionTimeoutMS=settings.mongo_timeout_ms)
db = client[settings.db_name]

# --- Shared error message table for AI chat endpoints ---
ERROR_MESSAGES = {
    "no_provider": "You haven't set up an AI provider yet. Go to Settings \u2192 AI Agent Setup to get started.",
    "invalid_key": "Your API key appears to be invalid or has expired. Please update it in Settings \u2192 AI Agent Setup.",
    "quota_exceeded": "You've reached the usage limit for your {provider} account. Please check your plan or billing.",
    "rate_limited": "You're sending messages too quickly. Please wait a moment before trying again.",
    "model_premium": "The selected model ({model}) requires a paid subscription on {provider}. Switch to a free-tier model or upgrade your account.",
    "network": "Couldn't reach {provider} right now. Check your internet connection and try again.",
    "model_deprecated": "The model '{model}' is no longer available on {provider}. Please select a different model in the chatbot settings.",
}

# Create the main app
app = FastAPI()
# Router without /api prefix - proxy will handle the /api part
api_router = APIRouter()

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# ============= In-Memory Storage (for when MongoDB is unavailable) =============
# This allows the app to work without MongoDB
in_memory_data = {
    "event_definitions": [],
    "event_data": [],
    "templates": [],
    "template_artifacts": [],
    "custom_functions": [],
    "transaction_reports": [],
    "transaction_definitions": [],
}
# Flag to track if we should use in-memory storage
USE_IN_MEMORY = False

# ============= Models (imported from models.py) =============
try:
    from backend.models import (
        EventDefinition, DSLFunction, EventData, DSLTemplate, DSLTemplateArtifact,
        TransactionOutput, TransactionReport, ChatMessage, ChatResponse,
        AIProviderTestRequest, AIProviderSaveRequest, DSLValidationRequest,
        SaveTemplateRequest, DSLRunRequest, TemplateExecuteRequest,
        TemplateDeployRequest,
    )
except Exception:
    try:
        from models import (
            EventDefinition, DSLFunction, EventData, DSLTemplate, DSLTemplateArtifact,
            TransactionOutput, TransactionReport, ChatMessage, ChatResponse,
            AIProviderTestRequest, AIProviderSaveRequest, DSLValidationRequest,
            SaveTemplateRequest, DSLRunRequest, TemplateExecuteRequest,
            TemplateDeployRequest,
        )
    except Exception:
        from .models import (
            EventDefinition, DSLFunction, EventData, DSLTemplate, DSLTemplateArtifact,
            TransactionOutput, TransactionReport, ChatMessage, ChatResponse,
            AIProviderTestRequest, AIProviderSaveRequest, DSLValidationRequest,
            SaveTemplateRequest, DSLRunRequest, TemplateExecuteRequest,
            TemplateDeployRequest,
        )

# ============= Sample Data (for when MongoDB is unavailable) =============
SAMPLE_EVENTS = [
    {
        "id": "evt1",
        "event_name": "LoanEvent",
        "fields": [
            {"name": "principal", "datatype": "decimal"},
            {"name": "rate", "datatype": "decimal"},
            {"name": "term", "datatype": "decimal"}
        ],
        "created_at": datetime.now(timezone.utc),
        "eventType": "activity",
        "eventTable": "standard"
    },
    {
        "id": "evt2",
        "event_name": "PaymentEvent",
        "fields": [
            {"name": "payment_amount", "datatype": "decimal"},
            {"name": "payment_date", "datatype": "date"},
            {"name": "payment_type", "datatype": "string"}
        ],
        "created_at": datetime.now(timezone.utc),
        "eventType": "activity",
        "eventTable": "standard"
    },
    {
        "id": "evt3",
        "event_name": "InvestmentEvent",
        "fields": [
            {"name": "initial_investment", "datatype": "decimal"},
            {"name": "return_rate", "datatype": "decimal"},
            {"name": "years", "datatype": "decimal"}
        ],
        "created_at": datetime.now(timezone.utc),
        "eventType": "activity",
        "eventTable": "standard"
    }
]

SAMPLE_TEMPLATES = []

# ============= Helper Functions =============

def parse_csv_content(content: str) -> List[List[str]]:
    """Parse CSV content and return list of rows"""
    # Remove BOM (Byte Order Mark) if present
    if content.startswith('\ufeff'):
        content = content[1:]
    reader = csv.reader(io.StringIO(content))
    return list(reader)

def get_field_case_insensitive(row: Dict[str, Any], field_name: str, default: Any = '') -> Any:
    """Get field value with case-insensitive key matching"""
    # First try exact match
    if field_name in row:
        return row[field_name]
    # Try case-insensitive match
    field_lower = field_name.lower()
    for key in row:
        if key.lower() == field_lower:
            return row[key]
    return default

def get_latest_data_per_instrument(data_rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    """Get latest postingdate per instrumentid (case-insensitive field matching).

    Defensively skips any row that is not a dict (e.g. a stringified JSON object
    that slipped through during import). Such rows are logged so the user can fix
    the source data instead of seeing a cryptic ``'str' object has no attribute 'items'``.
    """
    latest_data = {}
    for idx, row in enumerate(data_rows):
        if not isinstance(row, dict):
            logger.warning(
                "Skipping non-dict row at index %d in event data (got %s). "
                "Re-import the source file — each row must be a JSON object.",
                idx, type(row).__name__,
            )
            continue
        instrument_id = get_field_case_insensitive(row, 'instrumentid', '')
        posting_date = get_field_case_insensitive(row, 'postingdate', '')
        
        if not instrument_id:
            continue
            
        if instrument_id not in latest_data:
            latest_data[instrument_id] = row
        else:
            existing_date = get_field_case_insensitive(latest_data[instrument_id], 'postingdate', '')
            if posting_date > existing_date:
                latest_data[instrument_id] = row
    
    return latest_data

def extract_event_names_from_dsl(dsl_code: str) -> List[str]:
    """Extract all event names referenced in DSL code (EVENT_NAME.field pattern).

    Event names may be ANY identifier casing — including lowercase / snake_case
    like `line_items`. In the DSL a dot only ever means an event-field access
    (`EVENT.field`), so we match any `identifier.field`. (A previous version
    required an UPPERCASE first letter, which silently dropped lowercase event
    names → the whole run fell into standalone mode and failed with
    `name '<event>' is not defined`.) A matched name that isn't a real event is
    caught downstream when its definition is looked up.
    """
    import re
    pattern = r'\b([A-Za-z_][A-Za-z0-9_]*)\.[A-Za-z_][A-Za-z0-9_]*'
    matches = re.findall(pattern, dsl_code)
    # Drop DSL/Python builtins that can appear before a dot but are never events.
    _NOT_EVENTS = {"self", "math", "datetime", "os", "sys", "json", "re"}
    return list({m for m in matches if m not in _NOT_EVENTS})

def merge_event_data_by_instrument(event_data_dict: Dict[str, List[Dict]]) -> List[Dict]:
    """
    Merge data from multiple events by instrumentid.
    Each event's fields are prefixed with EVENT_NAME_ to avoid conflicts.
    Also provides event-specific postingdate, effectivedate, and subinstrumentid.
    
    Hierarchy: postingDate → instrumentId → subInstrumentId → effectiveDates
    
    If subInstrumentId is missing or null, it defaults to "1".
    """
    merged_data = {}
    bad_row_events = []
    
    for event_name, data_rows in event_data_dict.items():
        # Pre-flight check: ensure every row is a dict. Surface a clear error pointing
        # at the offending event/row so the user knows where to look.
        if isinstance(data_rows, list):
            for idx, row in enumerate(data_rows):
                if not isinstance(row, dict):
                    bad_row_events.append((event_name, idx, type(row).__name__))
        latest_data = get_latest_data_per_instrument(data_rows if isinstance(data_rows, list) else [])
        
        for instrument_id, row in latest_data.items():
            if instrument_id not in merged_data:
                # Get subinstrumentid with default of "1" if missing
                subinstrument_id = get_field_case_insensitive(row, 'subinstrumentid', '')
                if not subinstrument_id or subinstrument_id == 'None' or str(subinstrument_id).strip() == '':
                    subinstrument_id = '1'
                
                merged_data[instrument_id] = {
                    'instrumentid': instrument_id,
                    'subinstrumentid': str(subinstrument_id),
                    'postingdate': get_field_case_insensitive(row, 'postingdate', ''),
                    'effectivedate': get_field_case_insensitive(row, 'effectivedate', '')
                }
            
            # Get event-specific standard fields
            event_postingdate = get_field_case_insensitive(row, 'postingdate', '')
            event_effectivedate = get_field_case_insensitive(row, 'effectivedate', '')
            event_subinstrumentid = get_field_case_insensitive(row, 'subinstrumentid', '')
            if not event_subinstrumentid or event_subinstrumentid == 'None' or str(event_subinstrumentid).strip() == '':
                event_subinstrumentid = '1'
            
            # Add event-prefixed standard fields (e.g., INT_ACC_postingdate, INT_ACC_subinstrumentid)
            merged_data[instrument_id][f"{event_name}_postingdate"] = event_postingdate
            merged_data[instrument_id][f"{event_name}_effectivedate"] = event_effectivedate
            merged_data[instrument_id][f"{event_name}_subinstrumentid"] = str(event_subinstrumentid)
            
            # Add other fields with event prefix (EVENT_FIELD) for clarity
            # Also add without prefix for direct field access
            if not isinstance(row, dict):
                # Already logged above; skip safely.
                continue
            for key, value in row.items():
                key_lower = key.lower()
                if key_lower not in ['instrumentid', 'postingdate', 'effectivedate', 'subinstrumentid']:
                    # Store with event prefix: PMT_TRANSACTIONS_AMOUNT_REMIT
                    prefixed_key = f"{event_name}_{key}"
                    merged_data[instrument_id][prefixed_key] = value
                    # Also store the original field name for backward compatibility
                    merged_data[instrument_id][key] = value
    
    if bad_row_events:
        # Raise a single descriptive error pointing at the first bad row so the user
        # knows which event needs to be re-imported.
        evt, idx, kind = bad_row_events[0]
        raise ValueError(
            f"Event '{evt}' has malformed data: row #{idx} is a {kind}, not an object. "
            f"Re-import the source file — each row must be a JSON object "
            f"(total bad rows: {len(bad_row_events)})."
        )
    
    return list(merged_data.values())


def filter_event_data_by_posting_date(
    event_data_dict: Dict[str, List[Dict]],
    posting_date: str,
    event_metadata: Optional[Dict[str, Dict]] = None,
) -> Dict[str, List[Dict]]:
    """
    Return a copy of event_data_dict where each event's rows are restricted to those
    whose postingdate (case-insensitive) matches the requested posting_date string.
    Events with no matching rows keep an empty list (not removed, so callers can log
    a warning instead of crashing).

    Reference events (custom tables without postingdate) are passed through unchanged
    when their metadata says eventType == 'reference'. Without this, every CATALOG
    row would be filtered out and `collect_all(CATALOG.field)` would return [].
    """
    filtered: Dict[str, List[Dict]] = {}
    target = posting_date.strip()
    for event_name, rows in event_data_dict.items():
        safe_rows = rows if isinstance(rows, list) else []
        meta = (event_metadata or {}).get(event_name) or {}
        if str(meta.get('eventType', 'activity')).lower() == 'reference':
            # Reference tables have no postingdate — keep all rows.
            filtered[event_name] = list(safe_rows)
            continue
        filtered[event_name] = [
            row for row in safe_rows
            if isinstance(row, dict)
            and str(get_field_case_insensitive(row, "postingdate", "")).strip() == target
        ]
    return filtered


def _extract_dsl_line_from_exception(python_code: str, exc: Exception) -> Optional[int]:
    """Extract the DSL line number from a Python exception using DSL_LINE comments.
    
    Looks at the traceback to find the Python line that failed, then reads the
    corresponding source line from python_code to find a # DSL_LINE:N marker.
    Returns the DSL line number (1-based) or None if it can't be determined.
    """
    import traceback as tb_mod
    try:
        tb = exc.__traceback__
        if tb is None:
            return None
        # Walk to the innermost frame
        while tb.tb_next:
            tb = tb.tb_next
        py_lineno = tb.tb_lineno
        code_lines = python_code.split('\n')
        if 1 <= py_lineno <= len(code_lines):
            source_line = code_lines[py_lineno - 1]
            m = re.search(r'# DSL_LINE:(\d+)', source_line)
            if m:
                return int(m.group(1))
    except Exception:
        pass
    return None


def dsl_to_python_standalone(dsl_code: str) -> str:
    """Convert DSL code to Python for standalone execution (no events required)"""
    
    imports = '''
import sys, os
# Ensure backend package folder is on path so imports work when executed from different cwd
try:
    ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__)))
except Exception:
    ROOT_DIR = os.getcwd()
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)
try:
    from backend.dsl_functions import DSL_FUNCTIONS, _set_current_instrumentid, _set_current_postingdate, _set_current_subinstrumentid, _clear_transaction_results, _get_transaction_results, _set_dsl_print
except Exception:
    from dsl_functions import DSL_FUNCTIONS, _set_current_instrumentid, _set_current_postingdate, _set_current_subinstrumentid, _clear_transaction_results, _get_transaction_results, _set_dsl_print
from datetime import datetime
import json

# Preserve Python built-ins before updating with DSL functions
_builtin_min = min
_builtin_max = max
_builtin_sum = sum
_builtin_len = len
_builtin_range = range
_builtin_print = print

# Make all DSL functions available globally
globals().update(DSL_FUNCTIONS)

# Expose safe aliases for DSL functions whose names are Python keywords
and_op = DSL_FUNCTIONS.get('and', lambda a, b: a and b)
or_op = DSL_FUNCTIONS.get('or', lambda a, b: a or b)
not_op = DSL_FUNCTIONS.get('not', lambda a: not a)

# Restore Python built-ins (needed for native Python syntax)
min = _builtin_min
max = _builtin_max
sum = _builtin_sum
len = _builtin_len
# Smart range: DSL range(list)->max-min; Python range(int,...) for iterations
_dsl_range_val = DSL_FUNCTIONS.get('range', lambda col: (_builtin_max(col) - _builtin_min(col)) if col else 0)
def range(*args):
    if len(args) == 1 and isinstance(args[0], list):
        return _dsl_range_val(args[0])
    return _builtin_range(*args)

# Global list to capture print outputs
_print_outputs = []

def dsl_print(*args, **kwargs):
    try:
        # If a single argument looks like schedule(s), delegate to print_all_schedules
        if len(args) == 1:
            obj = args[0]
            if isinstance(obj, list) and obj:
                first = obj[0]
                if isinstance(first, dict) and 'schedule' in first:
                    try:
                        print_all_schedules(obj)
                        return
                    except Exception:
                        pass
                if isinstance(first, list):
                    inner_first = first[0] if first else None
                    if isinstance(inner_first, dict) and ('period_date' in inner_first or 'period_revenue' in inner_first or 'period_amount' in inner_first):
                        try:
                            print_all_schedules(obj)
                            return
                        except Exception:
                            pass
                    try:
                        print_all_schedules(obj)
                        return
                    except Exception:
                        pass
                if isinstance(first, dict) and ('period_date' in first or 'period_revenue' in first or 'period_amount' in first):
                    try:
                        # treat as array of rows (single schedule)
                        print_all_schedules([{"schedule": obj}])
                        return
                    except Exception:
                        pass
            if isinstance(obj, dict) and 'schedule' in obj:
                try:
                    print_all_schedules([obj])
                    return
                except Exception:
                    pass

        output_parts = []
        for arg in args:
            if isinstance(arg, (list, dict)):
                try:
                    output_parts.append(json.dumps(arg, indent=2, default=str))
                except Exception:
                    output_parts.append(str(arg))
            else:
                output_parts.append(str(arg))

        sep = kwargs.get('sep', ' ')
        output = sep.join(output_parts)
        _print_outputs.append(output)
    except Exception:
        try:
            _builtin_print(' '.join(map(str, args)))
        except Exception:
            pass

print = dsl_print

# Set the DSL print function for use by dsl_functions module (e.g., print_schedule)
_set_dsl_print(dsl_print)

def get_print_outputs():
    return _print_outputs

def clear_print_outputs():
    global _print_outputs
    _print_outputs = []
'''
    
    # Process DSL code
    import re
    processed_lines = []
    
    dsl_lines = dsl_code.split('\n')
    for dsl_line_num, line in enumerate(dsl_lines, start=1):
        stripped = line.strip()
        
        if not stripped or stripped.startswith('#') or stripped.startswith('//'):
            if stripped.startswith('//'):
                stripped = '#' + stripped[2:]
            processed_lines.append(f"    {stripped}" if stripped else "")
            continue
        
        # Replace Python keyword function calls with safe aliases
        stripped = re.sub(r'\band\s*\(', 'and_op(', stripped)
        stripped = re.sub(r'\bor\s*\(', 'or_op(', stripped)
        stripped = re.sub(r'\bnot\s*\(', 'not_op(', stripped)
        stripped = re.sub(r'\bif\s*\(', 'iif(', stripped)
        # ^ is exponentiation in DSL (Excel-style); convert to Python **
        stripped = re.sub(r'(?<!["\'])\^(?!["\'])', '**', stripped)

        # Simply add the line with DSL line marker
        processed_lines.append(f"    {stripped}  # DSL_LINE:{dsl_line_num}")
    
    python_body = '\n'.join(processed_lines)

    # Defense-in-depth: strict validation of user-supplied DSL code before it
    # is embedded in the trusted scaffolding. The scaffolding itself is then
    # also validated at exec time, but this catches obvious sandbox-escape
    # attempts (getattr/__class__/etc.) with a clearer error message and
    # before any code generation work downstream.
    _validate_dsl_user_code(python_body, label='<dsl_standalone_user_code>')
    
    template = f'''
{imports}

def process_standalone(override_postingdate=None, override_effectivedate=None):
    # Clear any previous transaction results
    _clear_transaction_results()
    
    # Set instrumentid for standalone mode
    _set_current_instrumentid('STANDALONE')
    _set_current_subinstrumentid('1')
    
    # Expose posting_date in scope so schedule column formulas can reference it
    postingdate = override_postingdate or ''
    posting_date = postingdate
    _set_current_postingdate(postingdate)
    effectivedate = override_effectivedate or ''
    effective_date = effectivedate
    
    # Execute DSL logic - transactions are created via createTransaction()
{python_body}
    
    # Get transactions created via createTransaction()
    results = _get_transaction_results()
    
    return results, get_print_outputs()
'''
    return template


# Helper to sanitize DB documents for JSON serialization
def sanitize_for_json(obj):
    """Recursively convert ObjectId to str and datetimes to ISO strings."""
    if isinstance(obj, dict):
        new = {}
        for k, v in obj.items():
            new[k] = sanitize_for_json(v)
        return new
    if isinstance(obj, list):
        return [sanitize_for_json(v) for v in obj]
    # ObjectId handling
    try:
        if ObjectId is not None and isinstance(obj, ObjectId):
            return str(obj)
    except Exception:
        pass
    # datetime handling
    from datetime import datetime
    if isinstance(obj, datetime):
        return obj.isoformat()
    return obj

def dsl_to_python_multi_event(dsl_code: str, all_event_fields: Dict[str, List[Dict[str, str]]]) -> str:
    """Convert DSL code to Python code template supporting multiple events and multiple transactions per row"""
    
    imports = '''
import sys, os
# Ensure backend package folder is on path so imports work when executed from different cwd
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)
try:
    from backend.dsl_functions import DSL_FUNCTIONS, _set_current_instrumentid, _set_current_postingdate, _set_current_subinstrumentid, _clear_transaction_results, _get_transaction_results, _set_dsl_print
except Exception:
    from dsl_functions import DSL_FUNCTIONS, _set_current_instrumentid, _set_current_postingdate, _set_current_subinstrumentid, _clear_transaction_results, _get_transaction_results, _set_dsl_print
from datetime import datetime
import json

# Preserve Python built-ins before updating with DSL functions
_builtin_min = min
_builtin_max = max
_builtin_sum = sum
_builtin_len = len
_builtin_range = range
_builtin_print = print

# Make all DSL functions available globally
globals().update(DSL_FUNCTIONS)

# Expose safe aliases for DSL functions whose names are Python keywords
and_op = DSL_FUNCTIONS.get('and', lambda a, b: a and b)
or_op = DSL_FUNCTIONS.get('or', lambda a, b: a or b)
not_op = DSL_FUNCTIONS.get('not', lambda a: not a)

# Restore Python built-ins (needed for native Python syntax)
min = _builtin_min
max = _builtin_max
sum = _builtin_sum
len = _builtin_len
# Smart range: DSL range(list)->max-min; Python range(int,...) for iterations
_dsl_range_val = DSL_FUNCTIONS.get('range', lambda col: (_builtin_max(col) - _builtin_min(col)) if col else 0)
def range(*args):
    if len(args) == 1 and isinstance(args[0], list):
        return _dsl_range_val(args[0])
    return _builtin_range(*args)

# Global list to capture print outputs
_print_outputs = []

def dsl_print(*args, **kwargs):
    """Custom print function that captures output for display in console"""
    try:
        # If a single argument looks like schedule(s), delegate to print_all_schedules
        if len(args) == 1:
            obj = args[0]
            if isinstance(obj, list) and obj:
                first = obj[0]
                if isinstance(first, dict) and 'schedule' in first:
                    try:
                        print_all_schedules(obj)
                        return
                    except Exception:
                        pass
                if isinstance(first, list):
                    inner_first = first[0] if first else None
                    if isinstance(inner_first, dict) and ('period_date' in inner_first or 'period_revenue' in inner_first or 'period_amount' in inner_first):
                        try:
                            print_all_schedules(obj)
                            return
                        except Exception:
                            pass
                    try:
                        print_all_schedules(obj)
                        return
                    except Exception:
                        pass
                if isinstance(first, dict) and ('period_date' in first or 'period_revenue' in first or 'period_amount' in first):
                    try:
                        # treat as array of rows (single schedule)
                        print_all_schedules([{"schedule": obj}])
                        return
                    except Exception:
                        pass
            if isinstance(obj, dict) and 'schedule' in obj:
                try:
                    print_all_schedules([obj])
                    return
                except Exception:
                    pass

        output_parts = []
        for arg in args:
            if isinstance(arg, (list, dict)):
                # Pretty print complex objects
                try:
                    output_parts.append(json.dumps(arg, indent=2, default=str))
                except Exception:
                    output_parts.append(str(arg))
            else:
                output_parts.append(str(arg))

        sep = kwargs.get('sep', ' ')
        output = sep.join(output_parts)
        _print_outputs.append(output)
    except Exception:
        try:
            _builtin_print(' '.join(map(str, args)))
        except Exception:
            pass

# Override print with our custom version
print = dsl_print

# Set the DSL print function for use by dsl_functions module (e.g., print_schedule)
_set_dsl_print(dsl_print)

def get_field_case_insensitive(row, field_name, default=''):
    \"\"\"Get field value with case-insensitive key matching\"\"\"
    if field_name in row:
        return row[field_name]
    field_lower = field_name.lower()
    for key in row:
        if key.lower() == field_lower:
            return row[key]
    return default

def get_print_outputs():
    \"\"\"Return all captured print outputs\"\"\"
    return _print_outputs

def clear_print_outputs():
    \"\"\"Clear captured print outputs\"\"\"
    global _print_outputs
    _print_outputs = []

# Global reference to all event data for collect() function
_all_event_data = []
_raw_event_data = {}  # Raw data by event name: {'ECF': [...], 'PMT': [...]}
_current_context = {}

def set_all_event_data(data):
    \"\"\"Set the global event data reference\"\"\"
    global _all_event_data
    _all_event_data = data

def set_raw_event_data(data):
    \"\"\"Set the raw event data (unmerged) for collect() functions\"\"\"
    global _raw_event_data
    if not isinstance(data, dict):
        # Refuse to corrupt global state — something upstream passed the wrong type.
        # Reset to empty so collect_*() functions return [] instead of crashing later
        # with the cryptic ``'str' object has no attribute 'items'``.
        try:
            _builtin_print(
                f"[dsl-template warning] set_raw_event_data got {type(data).__name__}; expected dict. Resetting to empty."
            )
        except Exception:
            pass
        _raw_event_data = {}
        return
    _raw_event_data = data

def set_current_context(instrumentid, postingdate, effectivedate, subinstrumentid='1'):
    \"\"\"Set the current row context for filtering collect_by_* functions\"\"\"
    global _current_context
    _current_context = {
        'instrumentid': instrumentid,
        'subinstrumentid': subinstrumentid or '1',
        'postingdate': postingdate,
        'effectivedate': effectivedate
    }

# Fields that are IDENTIFIERS, not measures. Coercing these to float turned
# subinstrumentid '1' into 1.0, so a natural join like
#   lookup(amounts, sub_ids, subinstrumentid)
# silently returned None -- the row built-in `subinstrumentid` is the STRING
# '1'. Everything else on the platform (row built-ins, TransactionOutput,
# merged event data) keeps these as strings, so collect_*() does too.
_IDENTIFIER_FIELDS = ('instrumentid', 'subinstrumentid')


def _is_identifier_field(actual_field, field_name):
    \"\"\"True when the collected field is an id rather than a measure.\"\"\"
    for candidate in (actual_field, field_name):
        if isinstance(candidate, str) and candidate.lower() in _IDENTIFIER_FIELDS:
            return True
    return False


def _row_has_field(row, name):
    \"\"\"True when `row` carries `name` (case-insensitive).\"\"\"
    if not isinstance(row, dict) or not isinstance(name, str):
        return False
    if name in row:
        return True
    lowered = name.lower()
    for key in row:
        if str(key).lower() == lowered:
            return True
    return False


def _no_such_collect_field(fn_name, field_name, actual_field):
    \"\"\"
    Message for a collect_*() whose field exists in no loaded event.

    This used to return one blank per scanned row -- an array of '' sized to
    the ACTIVITY row count, which looks like real data and quietly zeroed
    every downstream total. It happens when a reference event is named only
    inside a quoted collector argument: nothing detects the reference, so
    the event is never loaded for the run.
    \"\"\"
    loaded = sorted(_raw_event_data.keys())
    known = []
    for evt in loaded:
        rows = _raw_event_data.get(evt) or []
        if rows and isinstance(rows[0], dict):
            known.append(evt + '(' + ', '.join(sorted(rows[0].keys())) + ')')
        else:
            known.append(evt + '(no rows)')
    return (
        fn_name + '(' + repr(field_name) + '): no loaded event supplies a '
        'field named ' + repr(actual_field) + '. Loaded events: '
        + ('; '.join(known) if known else '(none)') + '. '
        'If the event name is part of that string, reference it in DOTTED '
        'form instead -- ' + fn_name + '(EVENTNAME.fieldname) -- so the run '
        'actually loads the event. A quoted name is invisible to the '
        'event loader.')


def _split_event_field(field_name):
    \"\"\"
    Split a flattened 'EVENTNAME_fieldname' reference into (event, field).

    An event name may itself contain underscores (SO_EVENT, line_items,
    sales_order). A naive field_name.split('_', 1) then picks the WRONG
    boundary -- 'SO_EVENT_line_amount' parsed as event 'SO' + field
    'EVENT_line_amount' -- which matches no event, so every collect_*()
    call silently returned []. Resolve against the event names we actually
    hold, longest first, so 'SO_EVENT' wins over a hypothetical 'SO'.

    Returns (None, field_name) when no known event prefixes the name: that
    means 'a bare field, look in every event', which is what a caller who
    passed an unprefixed name intends.
    \"\"\"
    if not isinstance(field_name, str):
        return None, field_name
    lowered = field_name.lower()
    for evt in sorted(_raw_event_data.keys(), key=len, reverse=True):
        prefix = str(evt).lower() + '_'
        if lowered.startswith(prefix) and len(field_name) > len(prefix):
            return evt, field_name[len(prefix):]
    return None, field_name


def collect_by_instrument(field_name):
    \"\"\"
    Collect all values of a field for the current instrumentid only (ignores dates).
    Useful for time-series data across multiple periods for same instrument.
    Returns numeric values as floats, non-numeric (dates, strings) as strings.

    Results are sorted by subinstrumentid (numeric-aware) so arrays produced
    by separate collect_by_instrument() calls in the same rule line up index
    for index across instruments. Without this sort, collect_by_instrument(REV.x)
    and collect_by_instrument(REV.y) could end up in different orders for
    different instruments and break index-based joins.
    \"\"\"
    pairs = []
    found_field = False
    current_instrument = _current_context.get('instrumentid', '')

    # Parse field_name (event names may contain underscores)
    event_name, actual_field = _split_event_field(field_name)

    for evt_name, rows in _raw_event_data.items():
        if event_name and evt_name.upper() != event_name.upper():
            continue

        for row in rows:
            row_instrument = get_field_case_insensitive(row, 'instrumentid', '')

            if row_instrument == current_instrument:
                if _row_has_field(row, actual_field) or _row_has_field(row, field_name):
                    found_field = True
                val = get_field_case_insensitive(row, actual_field, None)
                if val is None:
                    val = get_field_case_insensitive(row, field_name, None)
                # Always emit a row per subinstrument so parallel arrays stay
                # index-aligned. Type-aware placeholder is decided after the
                # scan so dates/strings don't get coerced to 0.
                sub = get_field_case_insensitive(row, 'subinstrumentid', '') or ''
                pairs.append((str(sub), val))

    # Scanned rows but the field was on none of them -> the caller named a
    # field (or an event) this run never loaded. Say so instead of handing
    # back a plausible-looking array of blanks.
    if pairs and not found_field:
        raise ValueError(_no_such_collect_field(
            'collect_by_instrument', field_name, actual_field))

    # Decide whether this is a numeric field. If every non-null value parses
    # as a number, missing entries become 0; otherwise they become ''. This
    # preserves subinstrument alignment without polluting date/string arrays
    # with a meaningless 0.
    all_numeric = True
    has_value = False
    for _s, v in pairs:
        if v is None or v == '':
            continue
        has_value = True
        try:
            float(v)
        except (ValueError, TypeError):
            all_numeric = False
            break
    # Identifier arrays stay textual end-to-end, so a missing id must be an
    # empty string too - never an int 0 sitting among string ids.
    _keep_as_text = _is_identifier_field(actual_field, field_name)
    null_placeholder = 0 if (has_value and all_numeric and not _keep_as_text) else ''

    converted = []
    for s, v in pairs:
        if v is None or v == '':
            converted.append((s, null_placeholder))
        elif _keep_as_text:
            converted.append((s, str(v)))
        else:
            try:
                converted.append((s, float(v)))
            except (ValueError, TypeError):
                converted.append((s, str(v)))
    pairs = converted

    def _sort_key(p):
        s = p[0]
        try:
            return (0, float(s))
        except (ValueError, TypeError):
            return (1, s)

    pairs.sort(key=_sort_key)
    sub_ids = [s for s, _v in pairs]
    values = [v for _s, v in pairs]
    try:
        from dsl_functions import _ScheduleValueList
        return _ScheduleValueList(values, subinstrument_ids=sub_ids)
    except Exception:
        return values

def collect_all(field_name):
    \"\"\"
    Collect ALL values of a field across all data rows (no filtering).
    Returns numeric values as floats, non-numeric (dates, strings) as strings.

    Results are sorted by subinstrumentid (numeric-aware) where present so
    parallel collect_all() arrays stay aligned by index. Reference tables
    without subinstrumentid keep their natural row order.
    \"\"\"
    pairs = []
    found_field = False

    # Parse field_name (event names may contain underscores)
    event_name, actual_field = _split_event_field(field_name)

    for evt_name, rows in _raw_event_data.items():
        if event_name and evt_name.upper() != event_name.upper():
            continue

        for idx, row in enumerate(rows):
            if _row_has_field(row, actual_field) or _row_has_field(row, field_name):
                found_field = True
            val = get_field_case_insensitive(row, actual_field, None)
            if val is None:
                val = get_field_case_insensitive(row, field_name, None)
            # Always emit a row so parallel collect_all() arrays stay
            # index-aligned. Type-aware placeholder is decided after scan.
            sub = get_field_case_insensitive(row, 'subinstrumentid', '') or ''
            pairs.append((str(sub), idx, val))

    if pairs and not found_field:
        raise ValueError(_no_such_collect_field(
            'collect_all', field_name, actual_field))

    all_numeric = True
    has_value = False
    for _s, _i, v in pairs:
        if v is None or v == '':
            continue
        has_value = True
        try:
            float(v)
        except (ValueError, TypeError):
            all_numeric = False
            break
    # Identifier arrays stay textual end-to-end, so a missing id must be an
    # empty string too - never an int 0 sitting among string ids.
    _keep_as_text = _is_identifier_field(actual_field, field_name)
    null_placeholder = 0 if (has_value and all_numeric and not _keep_as_text) else ''

    converted = []
    for s, i, v in pairs:
        if v is None or v == '':
            converted.append((s, i, null_placeholder))
        elif _keep_as_text:
            converted.append((s, i, str(v)))
        else:
            try:
                converted.append((s, i, float(v)))
            except (ValueError, TypeError):
                converted.append((s, i, str(v)))
    pairs = converted

    def _sort_key(p):
        s = p[0]
        if s == '':
            # Reference/no-sub rows keep insertion order via the idx tiebreaker.
            return (2, p[1])
        try:
            return (0, float(s), p[1])
        except (ValueError, TypeError):
            return (1, s, p[1])

    pairs.sort(key=_sort_key)
    return [v for _s, _i, v in pairs]

def collect_by_subinstrument(field_name):
    \"\"\"
    Collect all values of a field for the current instrumentid AND subinstrumentid.
    Useful when you need to filter by both parent and child entity.
    
    Hierarchy: postingDate → instrumentId → subInstrumentId → effectiveDates
    \"\"\"
    values = []
    found_field = False
    scanned = False
    current_instrument = _current_context.get('instrumentid', '')
    current_subinstrument = _current_context.get('subinstrumentid', '1')
    
    # Parse field_name (event names may contain underscores)
    event_name, actual_field = _split_event_field(field_name)
    
    for evt_name, rows in _raw_event_data.items():
        if event_name and evt_name.upper() != event_name.upper():
            continue
            
        for row in rows:
            row_instrument = get_field_case_insensitive(row, 'instrumentid', '')
            row_subinstrument = get_field_case_insensitive(row, 'subinstrumentid', '1') or '1'
            
            if row_instrument == current_instrument and row_subinstrument == current_subinstrument:
                scanned = True
                if _row_has_field(row, actual_field) or _row_has_field(row, field_name):
                    found_field = True
                val = get_field_case_insensitive(row, actual_field, None)
                if val is None:
                    val = get_field_case_insensitive(row, field_name, None)
                if val is not None and val != '':
                    if _is_identifier_field(actual_field, field_name):
                        values.append(str(val))
                    else:
                        try:
                            values.append(float(val))
                        except (ValueError, TypeError):
                            # For non-numeric values, store as string
                            values.append(val)
    if scanned and not found_field:
        raise ValueError(_no_such_collect_field(
            'collect_by_subinstrument', field_name, actual_field))
    return values

def collect_effectivedates_for_subinstrument(subinstrument_id=None):
    \"\"\"
    Collect all unique effectiveDates for a specific subInstrumentId within current instrumentId.
    If subinstrument_id is None, uses current context's subinstrumentid.
    \"\"\"
    current_instrument = _current_context.get('instrumentid', '')
    target_subinstrument = subinstrument_id or _current_context.get('subinstrumentid', '1')
    effective_dates = set()
    
    for evt_name, rows in _raw_event_data.items():
        for row in rows:
            row_instrument = get_field_case_insensitive(row, 'instrumentid', '')
            row_subinstrument = get_field_case_insensitive(row, 'subinstrumentid', '1') or '1'
            
            if row_instrument == current_instrument and row_subinstrument == target_subinstrument:
                edate = get_field_case_insensitive(row, 'effectivedate', '')
                if edate:
                    effective_dates.add(edate)
    
    return sorted(list(effective_dates))
'''
    
    # Process DSL code - convert EVENT.field to EVENT_field variable name
    import re
    processed_lines = []

    # Determine which events are reference events so we can alter collect() semantics
    reference_events = set()
    for ename, meta in all_event_fields.items():
        if isinstance(meta, dict):
            if str(meta.get('eventType', 'activity')).lower() == 'reference':
                reference_events.add(ename)

    # Precise EVENT.field flattening driven by the ACTUAL event names.
    # The legacy regex below required an UPPERCASE first letter and silently
    # skipped lowercase / snake_case event names (e.g. `line_items`), leaving
    # `line_items.postingdate` unconverted → `NameError: name 'line_items' is
    # not defined` at run/test time. We convert known event names first
    # (case-insensitive, canonicalised to the defined spelling), then keep the
    # old uppercase-CamelCase pass as a fallback for any ref not in the metadata.
    _event_canon = {en.lower(): en for en in all_event_fields.keys()}

    def _canon_evt(evt):
        return _event_canon.get(evt.lower(), evt)

    _known_event_re = None
    if all_event_fields:
        _event_alt = "|".join(
            re.escape(en) for en in
            sorted(all_event_fields.keys(), key=len, reverse=True)
        )
        _known_event_re = re.compile(
            rf"\b({_event_alt})\.([A-Za-z_][A-Za-z0-9_]*)", re.IGNORECASE)

    lines = dsl_code.split('\n')
    i = 0
    dsl_line_num = 0
    while i < len(lines):
        dsl_line_num = i + 1  # 1-based DSL line number
        line = lines[i].strip()

        if not line or line.startswith('#') or line.startswith('//'):
            if line.startswith('//'):
                line = '#' + line[2:]
            processed_lines.append(f"        {line}" if line else "")
            i += 1
            continue

        # Replace collect_by_instrument(...) -> use collect_all for reference events
        def _collect_by_inst_repl(m):
            evt, fld = _canon_evt(m.group(1)), m.group(2)
            if evt in reference_events:
                return f"collect_all('{evt}_{fld}')"
            return f"collect_by_instrument('{evt}_{fld}')"

        line = re.sub(r"collect_by_instrument\(\s*([A-Za-z_][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)\s*\)", _collect_by_inst_repl, line)

        # Same treatment for collect_by_subinstrument. Without this the
        # generic EVENT.field -> EVENT_field pass below rewrote the
        # argument into the flattened row VARIABLE (a float), and the
        # function died with "'float' object has no attribute 'split'".
        def _collect_by_sub_repl(m):
            evt, fld = _canon_evt(m.group(1)), m.group(2)
            if evt in reference_events:
                # Reference rows carry no instrument/sub-instrument scope.
                return f"collect_all('{evt}_{fld}')"
            return f"collect_by_subinstrument('{evt}_{fld}')"

        line = re.sub(r"collect_by_subinstrument\(\s*([A-Za-z_][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)\s*\)", _collect_by_sub_repl, line)

        # collect_all(EVENT.field) - always becomes collect_all('EVENT_field')
        line = re.sub(r"collect_all\(\s*([A-Za-z_][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)\s*\)",
                      lambda m: f"collect_all('{_canon_evt(m.group(1))}_{m.group(2)}')", line)

        # Convert EVENT.field -> EVENT_field. Known event names first (handles
        # lowercase / snake_case), then the legacy uppercase-CamelCase fallback.
        if _known_event_re is not None:
            line = _known_event_re.sub(
                lambda m: f"{_canon_evt(m.group(1))}_{m.group(2)}", line)
        line = re.sub(r"\b([A-Z][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)", r"\1_\2", line)

        # Replace Python keyword function calls with safe aliases
        line = re.sub(r'\band\s*\(', 'and_op(', line)
        line = re.sub(r'\bor\s*\(', 'or_op(', line)
        line = re.sub(r'\bnot\s*\(', 'not_op(', line)
        line = re.sub(r'\bif\s*\(', 'iif(', line)
        # ^ is exponentiation in DSL (Excel-style); convert to Python **
        line = re.sub(r'(?<!["\'])\^(?!["\'])', '**', line)

        # Add the line with DSL line marker
        processed_lines.append(f"        {line}  # DSL_LINE:{dsl_line_num}")

        i += 1

    python_body = '\n'.join(processed_lines)

    # Defense-in-depth: strict validation of user-supplied DSL code before it
    # is embedded in the multi-event scaffolding (see dsl_to_python_standalone
    # for rationale).
    _validate_dsl_user_code(python_body, label='<dsl_multi_event_user_code>')
    
    # Generate field extraction code for ALL events
    field_extraction_lines = []
    for event_name, meta in all_event_fields.items():
        # meta may be a dict with 'fields' and 'eventType' or a simple list
        if isinstance(meta, dict):
            fields = meta.get('fields', [])
            etype = str(meta.get('eventType', 'activity')).lower()
        else:
            fields = meta
            etype = 'activity'

        field_extraction_lines.append(f"        # Fields from {event_name} ({etype})")

        # Add event-specific standard fields only for activity events
        if etype == 'activity':
            field_extraction_lines.append(
                f"        {event_name}_postingdate = str(get_field_case_insensitive(row, '{event_name}_postingdate', ''))"
            )
            field_extraction_lines.append(
                f"        {event_name}_effectivedate = str(get_field_case_insensitive(row, '{event_name}_effectivedate', ''))"
            )
            field_extraction_lines.append(
                f"        {event_name}_subinstrumentid = str(get_field_case_insensitive(row, '{event_name}_subinstrumentid', '1'))"
            )

        for field in fields:
            field_name = field['name']
            field_type = (field.get('datatype', 'string') or 'string').strip().lower()
            # Variable name: EVENT_FIELD
            var_name = f"{event_name}_{field_name}"

            # Normalise common numeric type aliases to 'decimal'
            if field_type in ('decimal', 'number', 'numeric', 'float', 'double',
                              'currency', 'money', 'amount', 'percent', 'rate'):
                field_extraction_lines.append(
                    f"        {var_name} = float(get_field_case_insensitive(row, '{var_name}', 0) or 0)"
                )
            elif field_type in ('integer', 'int', 'long', 'count'):
                field_extraction_lines.append(
                    f"        {var_name} = int(float(get_field_case_insensitive(row, '{var_name}', 0) or 0))"
                )
            elif field_type == 'date':
                field_extraction_lines.append(
                    f"        {var_name} = str(get_field_case_insensitive(row, '{var_name}', ''))"
                )
            elif field_type == 'boolean':
                field_extraction_lines.append(
                    f"        {var_name} = str(get_field_case_insensitive(row, '{var_name}', '')).lower() in ['true', '1', 'yes']"
                )
            else:
                # 'string' or unrecognised type: smart-cast — if the stored
                # value is already numeric keep it; if it looks like a number
                # coerce to float so arithmetic formulas don't fail with
                # "unsupported operand type(s) for /: 'str' and 'str'".
                field_extraction_lines.append(
                    f"        _fv = get_field_case_insensitive(row, '{var_name}', '')"
                )
                field_extraction_lines.append(
                    f"        if isinstance(_fv, (int, float)):"
                )
                field_extraction_lines.append(
                    f"            {var_name} = _fv"
                )
                field_extraction_lines.append(
                    f"        else:"
                )
                field_extraction_lines.append(
                    f"            _s = str(_fv if _fv is not None else '')"
                )
                field_extraction_lines.append(
                    f"            try: {var_name} = float(_s) if _s.strip() else _s"
                )
                field_extraction_lines.append(
                    f"            except (ValueError, TypeError): {var_name} = _s"
                )
    
    field_extraction_code = '\n'.join(field_extraction_lines)
    
    template = f"""
{imports}
def process_event_data(event_data, raw_event_data=None, override_postingdate=None, override_effectivedate=None):
    # Clear any previous transaction results
    _clear_transaction_results()
    
    _override_postingdate = override_postingdate
    _override_effectivedate = override_effectivedate
    
    # If raw event data provided by the caller, set it for collect() functions
    if raw_event_data is not None:
        set_raw_event_data(raw_event_data)

    # Set global event data for collect() function
    set_all_event_data(event_data)

    # Activity-data ordering guarantee: enforce
    #   instrumentid ASC, postingdate ASC, effectivedate ASC, subinstrumentid ASC
    # so every step inside this rule (Schedule, Condition, Iteration,
    # Calculation, Custom Code, Create Transaction) sees rows in the same
    # canonical order. event_data here is the merged ACTIVITY dataset only;
    # reference/custom rows live in raw_event_data and are not touched.
    try:
        if isinstance(event_data, list) and len(event_data) > 1:
            event_data.sort(key=lambda _r: (
                str(get_field_case_insensitive(_r, 'instrumentid', '') or ''),
                str(get_field_case_insensitive(_r, 'postingdate', '') or ''),
                str(get_field_case_insensitive(_r, 'effectivedate', '') or ''),
                str(get_field_case_insensitive(_r, 'subinstrumentid', '1') or '1'),
            ))
    except Exception:
        pass
    
    for row in event_data:
        # Extract standard fields (case-insensitive)
        postingdate = get_field_case_insensitive(row, 'postingdate', '')
        effectivedate = get_field_case_insensitive(row, 'effectivedate', '') or postingdate
        instrumentid = get_field_case_insensitive(row, 'instrumentid', '')
        subinstrumentid = get_field_case_insensitive(row, 'subinstrumentid', '1') or '1'
        # Expose underscore aliases so schedule column formulas can reference them
        posting_date = postingdate
        effective_date = effectivedate
        
        # Set current instrumentid for createTransaction()
        _set_current_instrumentid(instrumentid)
        # Set current sub-instrument so schedule() can bind the
        # `subinstrument_id` column built-in to this row.
        _set_current_subinstrumentid(subinstrumentid)
        # Set current postingdate so print_schedule() can tag emitted rows
        # with (_instrumentid, _postingdate) for the Business Preview filter.
        _set_current_postingdate(postingdate)
        
        # Set current context for collect() filtering
        set_current_context(instrumentid, postingdate, effectivedate, subinstrumentid)
        
        # Extract fields from all events with proper datatype conversion
{field_extraction_code}
        
        # Execute DSL logic - transactions are created via createTransaction()
{python_body}
    
    # Get all transactions created via createTransaction()
    results = _get_transaction_results()
    return results
"""
    return template

def dsl_to_python(dsl_code: str, event_fields: List[Dict[str, str]]) -> str:
    """Wrapper for backward compatibility - uses multi-event version"""
    # Convert single event fields to dict format
    all_event_fields = {"DEFAULT": event_fields}
    return dsl_to_python_multi_event(dsl_code, all_event_fields)

async def execute_python_template(python_code: str, event_data: List[Dict[str, Any]], raw_event_data: Dict[str, List[Dict]] = None, override_postingdate: str = None, override_effectivedate: str = None) -> Dict[str, Any]:
    """Execute Python template on event data and return transactions + print outputs"""
    # Execute the generated python template in a restricted context and return results.
    try:
        # When executed as package, templates expect to import dsl_functions; ensure package-qualified import
        if "from dsl_functions import" in python_code:
            python_code = python_code.replace("from dsl_functions import", "from backend.dsl_functions import")

        # Provide a minimal execution globals mapping including __file__ so
        # template code that uses os.path.dirname(__file__) will work when
        # executed via exec(). Use the server file path as a sensible base.
        exec_globals = {
            '__file__': os.path.abspath(__file__),
            '__name__': '__dsl_template__',
            '__builtins__': _make_sandbox_builtins(),
        }
        # Defense-in-depth: AST-validate the generated template before exec
        # so user-injected Custom Code cannot reach __import__, dunder
        # introspection, or import disallowed modules.
        _validate_template_ast(python_code, label='<dsl_template>')
        # Execute the template which defines helper functions like process_event_data, get_print_outputs
        exec(compile(python_code, '<dsl_template>', 'exec'), exec_globals)

        # Prefer calling process_event_data (multi-event template) and pass raw_event_data.
        # Inspect the signature explicitly so we never swallow internal TypeErrors as a
        # "wrong signature" — that previously caused the 3-arg fallback to bind
        # raw_event_data = override_postingdate (a string), corrupting global state and
        # producing the cryptic "'str' object has no attribute 'items'" error from
        # collect_by_instrument on subsequent calls.
        if 'process_event_data' in exec_globals:
            import inspect as _inspect
            _proc = exec_globals['process_event_data']
            try:
                _sig = _inspect.signature(_proc)
                _param_count = len(_sig.parameters)
            except (TypeError, ValueError):
                _param_count = 4
            if _param_count >= 4:
                transactions = _proc(event_data, raw_event_data, override_postingdate, override_effectivedate)
            else:
                # Older template signature without raw_event_data
                transactions = _proc(event_data, override_postingdate, override_effectivedate)
        elif 'process_standalone' in exec_globals:
            transactions = exec_globals['process_standalone'](override_postingdate, override_effectivedate)
            # process_standalone returns (transactions, print_outputs) whereas
            # process_event_data returns just the transactions. Treating the
            # tuple as a list of transactions meant EVERY standalone rule came
            # back with zero transactions -- both entries failed
            # TransactionOutput(**...) and were quietly dropped by the
            # normalisation loop below.
            if isinstance(transactions, tuple):
                transactions = transactions[0] if transactions else []
        else:
            raise RuntimeError('Template did not define a process function')
        # Normalize transactions into TransactionOutput models if needed
        # Some DSL helpers (createTransaction) return plain dicts; convert them to
        # TransactionOutput so callers can call `model_dump()` uniformly.
        normalized_transactions = []
        for t in transactions or []:
            try:
                if hasattr(t, 'model_dump'):
                    normalized_transactions.append(t)
                else:
                    normalized_transactions.append(TransactionOutput(**t))
            except Exception:
                # If conversion fails, skip the transaction but continue
                logger.debug(f"Skipping invalid transaction object during normalization: {t}")

        print_outputs = []
        if 'get_print_outputs' in exec_globals:
            try:
                print_outputs = exec_globals['get_print_outputs']()
            except Exception:
                print_outputs = []

        # Surface how many transactions the zero-amount guard suppressed, so a
        # run whose row count is lower than its input can explain the gap
        # instead of the rows just not being there.
        _zero_skipped = 0
        try:
            from backend.dsl_functions import _get_skipped_zero_amount
        except Exception:
            try:
                from dsl_functions import _get_skipped_zero_amount
            except Exception:
                _get_skipped_zero_amount = None
        if _get_skipped_zero_amount is not None:
            try:
                _zero_skipped = _get_skipped_zero_amount()
            except Exception:
                _zero_skipped = 0

        return {"transactions": normalized_transactions,
                "print_outputs": print_outputs,
                "zero_amount_skipped": _zero_skipped}
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        # Dump the generated template so we can inspect the offending line by number
        try:
            with open('/tmp/last_dsl_template.py', 'w') as _f:
                _f.write(python_code)
        except Exception:
            pass
        dsl_line = _extract_dsl_line_from_exception(python_code, e)
        error_msg = str(e)
        if dsl_line:
            error_msg = f"[Line {dsl_line}] {error_msg}"
        logger.error(f"Error executing python template: {error_msg}\nFull traceback:\n{tb}")
        raise HTTPException(status_code=500, detail=error_msg)

# ============= API Endpoints =============

@api_router.get("/")
async def root():
    return {"message": "Fyntrac DSL Studio API"}

@api_router.post("/load-simple-sample")
async def load_simple_sample():
    """Load a small, focused sample dataset for the Settings → Load Sample Data menu.

    Contains exactly two instruments and two event definitions that together
    exercise the standard/custom event-table split:

      - LoanActivity  (eventTable=standard, eventType=activity)   — per-instrument activity rows
      - RateSchedule  (eventTable=custom,   eventType=reference)  — shared reference data

    Existing event_definitions, event_data and dsl_functions collections are
    cleared first so the user starts from a clean slate.
    """
    try:
        await db.event_definitions.delete_many({})
        await db.dsl_functions.delete_many({})
        await db.event_data.delete_many({})

        # ── Event Definitions ───────────────────────────────────────────────
        loan_activity_def = EventDefinition(
            event_name="LoanActivity",
            fields=[
                {"name": "principal", "datatype": "decimal"},
                {"name": "rate_code", "datatype": "string"},
                {"name": "term_months", "datatype": "integer"},
                {"name": "origination_date", "datatype": "date"},
            ],
            eventType="activity",
            eventTable="standard",
        )
        rate_schedule_def = EventDefinition(
            event_name="RateSchedule",
            fields=[
                {"name": "rate_code", "datatype": "string"},
                {"name": "rate_value", "datatype": "decimal"},
                {"name": "effective_date", "datatype": "date"},
                {"name": "expiry_date", "datatype": "date"},
            ],
            eventType="reference",
            eventTable="custom",
        )
        for evt in (loan_activity_def, rate_schedule_def):
            doc = evt.model_dump()
            doc['created_at'] = doc['created_at'].isoformat()
            await db.event_definitions.insert_one(doc)

        # ── Activity Data — 2 instruments ───────────────────────────────────
        loan_activity_data = EventData(
            event_name="LoanActivity",
            data_rows=[
                {
                    "postingdate": "2026-01-01",
                    "effectivedate": "2026-01-01",
                    "instrumentid": "INST-001",
                    "principal": "100000",
                    "rate_code": "PRIME",
                    "term_months": "60",
                    "origination_date": "2026-01-01",
                },
                {
                    "postingdate": "2026-01-01",
                    "effectivedate": "2026-01-01",
                    "instrumentid": "INST-002",
                    "principal": "250000",
                    "rate_code": "BASE",
                    "term_months": "120",
                    "origination_date": "2026-01-01",
                },
            ],
        )
        doc = loan_activity_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)

        # ── Reference Data ──────────────────────────────────────────────────
        rate_schedule_data = EventData(
            event_name="RateSchedule",
            data_rows=[
                {"rate_code": "PRIME", "rate_value": "0.0525", "effective_date": "2025-01-01", "expiry_date": "2025-12-31"},
                {"rate_code": "PRIME", "rate_value": "0.0500", "effective_date": "2026-01-01", "expiry_date": "2026-12-31"},
                {"rate_code": "BASE",  "rate_value": "0.0400", "effective_date": "2025-01-01", "expiry_date": "2025-12-31"},
                {"rate_code": "BASE",  "rate_value": "0.0375", "effective_date": "2026-01-01", "expiry_date": "2026-12-31"},
            ],
        )
        doc = rate_schedule_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)

        # ── Transaction Definitions ─────────────────────────────────────────
        await db.transaction_definitions.delete_many({})
        sample_txn_types = [
            "InterestAccrual",
            "PrincipalPayment",
            "FeeAmortization",
            "Revenue",
            "LeaseExpense",
            "NPVAnalysis",
        ]
        for txn_type in sample_txn_types:
            await db.transaction_definitions.insert_one({"transactiontype": txn_type})

        return {
            "message": "Simple sample data loaded successfully",
            "events": ["LoanActivity", "RateSchedule"],
            "instruments": ["INST-001", "INST-002"],
            "transaction_types": sample_txn_types,
        }
    except Exception as e:
        logger.exception("Failed to load simple sample data")
        raise HTTPException(status_code=500, detail=f"Failed to load sample data: {str(e)}")


@api_router.post("/load-sample-data")
async def load_sample_data():
    """Load sample data for testing"""
    try:
        # Clear existing data
        await db.event_definitions.delete_many({})
        await db.dsl_functions.delete_many({})
        await db.event_data.delete_many({})
        
        # Sample Event Definitions with datatypes
        sample_events = [
            EventDefinition(event_name="LoanEvent", fields=[
                {"name": "principal", "datatype": "decimal"},
                {"name": "rate", "datatype": "decimal"},
                {"name": "term", "datatype": "decimal"}
            ]),
            EventDefinition(event_name="PaymentEvent", fields=[
                {"name": "payment_amount", "datatype": "decimal"},
                {"name": "payment_date", "datatype": "date"},
                {"name": "payment_type", "datatype": "string"}
            ]),
            EventDefinition(event_name="InvestmentEvent", fields=[
                {"name": "initial_investment", "datatype": "decimal"},
                {"name": "return_rate", "datatype": "decimal"},
                {"name": "years", "datatype": "decimal"}
            ]),
            # Custom reference tables
            EventDefinition(
                event_name="RateTable",
                fields=[
                    {"name": "rate_code", "datatype": "string"},
                    {"name": "rate_value", "datatype": "decimal"},
                    {"name": "effective_date", "datatype": "date"},
                    {"name": "expiry_date", "datatype": "date"},
                ],
                eventType="reference",
                eventTable="custom",
            ),
            EventDefinition(
                event_name="ProductConfig",
                fields=[
                    {"name": "product_code", "datatype": "string"},
                    {"name": "product_name", "datatype": "string"},
                    {"name": "max_term", "datatype": "integer"},
                    {"name": "min_principal", "datatype": "decimal"},
                    {"name": "max_principal", "datatype": "decimal"},
                    {"name": "fee_percent", "datatype": "decimal"},
                ],
                eventType="reference",
                eventTable="custom",
            ),
        ]
        
        for event in sample_events:
            doc = event.model_dump()
            doc['created_at'] = doc['created_at'].isoformat()
            await db.event_definitions.insert_one(doc)
        
        # Sample Event Data - LoanEvent
        loan_data = EventData(
            event_name="LoanEvent",
            data_rows=[
                {
                    "postingdate": "2026-01-01",
                    "effectivedate": "2026-01-01",
                    "instrumentid": "LOAN-001",
                    "principal": "100000",
                    "rate": "0.05",
                    "term": "12"
                },
                {
                    "postingdate": "2026-01-15",
                    "effectivedate": "2026-01-15",
                    "instrumentid": "LOAN-002",
                    "principal": "50000",
                    "rate": "0.04",
                    "term": "6"
                },
                {
                    "postingdate": "2026-02-01",
                    "effectivedate": "2026-02-01",
                    "instrumentid": "LOAN-003",
                    "principal": "250000",
                    "rate": "0.06",
                    "term": "24"
                }
            ]
        )
        
        doc = loan_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)
        
        # Sample Event Data - PaymentEvent (instrumentids match LoanEvent for join)
        payment_data = EventData(
            event_name="PaymentEvent",
            data_rows=[
                {
                    "postingdate": "2026-01-01",
                    "effectivedate": "2026-01-01",
                    "instrumentid": "LOAN-001",
                    "payment_amount": "5000",
                    "payment_date": "2026-01-01",
                    "payment_type": "Principal"
                },
                {
                    "postingdate": "2026-01-15",
                    "effectivedate": "2026-01-15",
                    "instrumentid": "LOAN-002",
                    "payment_amount": "2000",
                    "payment_date": "2026-01-15",
                    "payment_type": "Interest"
                }
            ]
        )
        
        doc = payment_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)
        
        # Sample Event Data - InvestmentEvent
        investment_data = EventData(
            event_name="InvestmentEvent",
            data_rows=[
                {
                    "postingdate": "2026-01-01",
                    "effectivedate": "2026-01-01",
                    "instrumentid": "LOAN-001",
                    "initial_investment": "10000",
                    "return_rate": "0.08",
                    "years": "5"
                },
                {
                    "postingdate": "2026-01-15",
                    "effectivedate": "2026-01-15",
                    "instrumentid": "LOAN-002",
                    "initial_investment": "25000",
                    "return_rate": "0.10",
                    "years": "10"
                }
            ]
        )
        
        doc = investment_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)

        # Sample Custom Reference Data - RateTable
        rate_table_data = EventData(
            event_name="RateTable",
            data_rows=[
                {"rate_code": "PRIME", "rate_value": "0.0525", "effective_date": "2025-01-01", "expiry_date": "2025-06-30"},
                {"rate_code": "PRIME", "rate_value": "0.0500", "effective_date": "2025-07-01", "expiry_date": "2025-12-31"},
                {"rate_code": "PRIME", "rate_value": "0.0475", "effective_date": "2026-01-01", "expiry_date": "2026-12-31"},
                {"rate_code": "BASE",  "rate_value": "0.0400", "effective_date": "2025-01-01", "expiry_date": "2025-12-31"},
                {"rate_code": "BASE",  "rate_value": "0.0375", "effective_date": "2026-01-01", "expiry_date": "2026-12-31"},
                {"rate_code": "LIBOR", "rate_value": "0.0310", "effective_date": "2025-01-01", "expiry_date": "2025-12-31"},
                {"rate_code": "LIBOR", "rate_value": "0.0290", "effective_date": "2026-01-01", "expiry_date": "2026-12-31"},
            ]
        )
        doc = rate_table_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)

        # Sample Custom Reference Data - ProductConfig
        product_config_data = EventData(
            event_name="ProductConfig",
            data_rows=[
                {"product_code": "HL-STD",  "product_name": "Standard Home Loan",    "max_term": "360", "min_principal": "50000",  "max_principal": "2000000", "fee_percent": "0.005"},
                {"product_code": "HL-FIX",  "product_name": "Fixed Rate Home Loan",  "max_term": "300", "min_principal": "100000", "max_principal": "1500000", "fee_percent": "0.0075"},
                {"product_code": "PL-UNSEC","product_name": "Unsecured Personal Loan","max_term": "84",  "min_principal": "5000",   "max_principal": "100000",  "fee_percent": "0.010"},
                {"product_code": "BL-SME",  "product_name": "SME Business Loan",     "max_term": "120", "min_principal": "20000",  "max_principal": "500000",  "fee_percent": "0.008"},
                {"product_code": "INV-TERM","product_name": "Term Investment",        "max_term": "60",  "min_principal": "10000",  "max_principal": "5000000", "fee_percent": "0.000"},
            ]
        )
        doc = product_config_data.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.event_data.insert_one(doc)

        # Sample DSL Code - Loan Validation, Fee Calculation & Investment Projection
        sample_dsl_code = """## 1. Loan Validation and Fee Calculation
## Reference data access for single-row config
min_p = ProductConfig.min_principal or 0
max_p = ProductConfig.max_principal or 1000000
principal = LoanEvent.principal or 0

## Check if loan principal is within allowed range
is_valid = and(gte(principal, min_p), lte(principal, max_p))

## Calculate fee using percentage from ProductConfig
fee_pct = ProductConfig.fee_percent or 0.01
loan_fee = if(is_valid, multiply(principal, fee_pct), 0)
print(concat("Calculated Loan Fee: ", loan_fee))

## 2. Loan Payment Calculation
annual_rate = LoanEvent.rate or 0.05
monthly_rate = divide(annual_rate, 12)
term_months = LoanEvent.term or 360

## Use multiply to handle negation for the PV argument in pmt()
neg_principal = multiply(principal, -1)
monthly_pmt = pmt(monthly_rate, term_months, neg_principal)
print(concat("Expected Monthly Payment: ", monthly_pmt))

## 3. Investment Growth Projection
init_inv = InvestmentEvent.initial_investment or 0
ret_rate = InvestmentEvent.return_rate or 0
inv_years = InvestmentEvent.years or 0

## Future value calculation
neg_inv = multiply(init_inv, -1)
future_val = fv(ret_rate, inv_years, 0, neg_inv)
print(concat("Projected Investment Value: ", future_val))

## 4. Create Transactions
## Use global postingdate and effectivedate (no prefixes needed)

## Only create fee transaction if the amount is greater than 0
if(gt(loan_fee, 0), createTransaction(postingdate, effectivedate, "LoanProcessingFee", loan_fee), 0)

## Record the monthly interest accrual
monthly_interest = multiply(principal, monthly_rate)
createTransaction(postingdate, effectivedate, "InterestAccrual", monthly_interest)"""
        
        return {
            "message": "Sample data loaded successfully",
            "events": ["LoanEvent", "PaymentEvent", "InvestmentEvent", "RateTable", "ProductConfig"],
            "sample_dsl_code": sample_dsl_code
        }
    except Exception as e:
        # If database is not available, fall back to in-memory sample data
        logger.warning(f"Could not load sample data into MongoDB, falling back to in-memory: {str(e)}")
        try:
            # Populate in-memory structures for tests
            global USE_IN_MEMORY, in_memory_data
            USE_IN_MEMORY = True
            in_memory_data['event_definitions'] = SAMPLE_EVENTS
            in_memory_data['templates'] = SAMPLE_TEMPLATES
            # Create sample event_data entries similar to DB documents
            # Build simple event_data entries from SAMPLE_EVENTS for tests
            simple_event_docs = []
            for evt in SAMPLE_EVENTS:
                doc = {
                    'id': evt.get('id', str(uuid.uuid4())),
                    'event_name': evt['event_name'],
                    'data_rows': [],
                    'created_at': evt.get('created_at', datetime.now(timezone.utc)).isoformat()
                }
                simple_event_docs.append(doc)
            in_memory_data['event_data'] = simple_event_docs
        except Exception:
            logger.exception("Failed to populate in-memory sample data")

        sample_dsl_code = """## 1. Loan Validation and Fee Calculation
## Reference data access for single-row config
min_p = ProductConfig.min_principal or 0
max_p = ProductConfig.max_principal or 1000000
principal = LoanEvent.principal or 0

## Check if loan principal is within allowed range
is_valid = and(gte(principal, min_p), lte(principal, max_p))

## Calculate fee using percentage from ProductConfig
fee_pct = ProductConfig.fee_percent or 0.01
loan_fee = if(is_valid, multiply(principal, fee_pct), 0)
print(concat("Calculated Loan Fee: ", loan_fee))

## 2. Loan Payment Calculation
annual_rate = LoanEvent.rate or 0.05
monthly_rate = divide(annual_rate, 12)
term_months = LoanEvent.term or 360

## Use multiply to handle negation for the PV argument in pmt()
neg_principal = multiply(principal, -1)
monthly_pmt = pmt(monthly_rate, term_months, neg_principal)
print(concat("Expected Monthly Payment: ", monthly_pmt))

## 3. Investment Growth Projection
init_inv = InvestmentEvent.initial_investment or 0
ret_rate = InvestmentEvent.return_rate or 0
inv_years = InvestmentEvent.years or 0

## Future value calculation
neg_inv = multiply(init_inv, -1)
future_val = fv(ret_rate, inv_years, 0, neg_inv)
print(concat("Projected Investment Value: ", future_val))

## 4. Create Transactions
## Use global postingdate and effectivedate (no prefixes needed)

## Only create fee transaction if the amount is greater than 0
if(gt(loan_fee, 0), createTransaction(postingdate, effectivedate, "LoanProcessingFee", loan_fee), 0)

## Record the monthly interest accrual
monthly_interest = multiply(principal, monthly_rate)
createTransaction(postingdate, effectivedate, "InterestAccrual", monthly_interest)"""
        return {
            "message": "Sample data loaded into memory (DB unavailable)",
            "events": [e['event_name'] for e in SAMPLE_EVENTS],
            "sample_dsl_code": sample_dsl_code
        }

@api_router.delete("/clear-all-data")
async def clear_all_data():
    """Clear all data from the system except templates"""
    try:
        # Delete all collections EXCEPT templates
        await db.event_definitions.delete_many({})
        await db.event_data.delete_many({})
        await db.transaction_reports.delete_many({})
        await db.custom_functions.delete_many({})
        await db.saved_rules.delete_many({})
        await db.saved_schedules.delete_many({})
        await db.transaction_definitions.delete_many({})

        # Also clear in-memory fallback data so stale entries don't survive
        global in_memory_data
        in_memory_data['event_definitions'] = []
        in_memory_data['event_data'] = []
        in_memory_data['transaction_reports'] = []
        in_memory_data['custom_functions'] = []
        in_memory_data['transaction_definitions'] = []
        in_memory_data.pop('saved_rules', None)
        in_memory_data.pop('saved_schedules', None)

        return {
            "message": "All data cleared successfully (templates preserved).",
            "cleared": ["event_definitions", "event_data", "transaction_reports", "custom_functions", "saved_rules", "saved_schedules", "transaction_definitions"],
            "preserved": ["templates"]
        }
    except Exception as e:
        logger.error(f"Error clearing data: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@api_router.post("/events/upload")
async def upload_event_definitions(file: UploadFile = File(...)):
    """Upload Reference Data File (.xlsx) with two sheets:
    - 'events'       : EventName, EventField, DataType[, EventType[, EventTable]]
    - 'transactions' : transactiontype (single column, optional)
    """
    try:
        if not (file.filename or '').lower().endswith('.xlsx'):
            raise HTTPException(
                status_code=400,
                detail="File must be an Excel file (.xlsx). The Reference Data File format is .xlsx with two sheets: 'events' and 'transactions'."
            )
        content = await file.read()
        try:
            xl = pd.ExcelFile(io.BytesIO(content))
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {exc}")

        sheet_names_lower = [s.lower() for s in xl.sheet_names]

        # ── Events sheet ───────────────────────────────────────────────────────
        if 'events' not in sheet_names_lower:
            raise HTTPException(
                status_code=400,
                detail="Excel file must have a sheet named 'events' with columns: EventName, EventField, DataType[, EventType[, EventTable]]"
            )
        events_sheet = xl.sheet_names[sheet_names_lower.index('events')]
        try:
            events_df = pd.read_excel(xl, sheet_name=events_sheet, header=0, dtype=str)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Could not read 'events' sheet: {exc}")

        # Normalise column names (strip whitespace)
        events_df.columns = [str(c).strip() for c in events_df.columns]
        col_lower_map = {c.lower(): c for c in events_df.columns}

        for req in ('eventname', 'eventfield', 'datatype'):
            if req not in col_lower_map:
                raise HTTPException(
                    status_code=400,
                    detail=f"Missing required column in 'events' sheet: '{req}'. Required: EventName, EventField, DataType"
                )

        VALID_EVENT_TYPES = ('activity', 'reference')
        VALID_EVENT_TABLES = ('standard', 'custom')

        events_dict = {}
        event_type_map = {}
        event_table_map = {}

        for _, row in events_df.iterrows():
            def _cell(col_key):
                col = col_lower_map.get(col_key)
                if col is None:
                    return None
                val = row[col]
                return str(val).strip() if val is not None and not (isinstance(val, float) and pd.isna(val)) and str(val).strip() not in ('', 'nan', 'None') else None

            event_name = _cell('eventname')
            event_field = _cell('eventfield')
            data_type = (_cell('datatype') or '').lower()

            if not event_name or not event_field or not data_type:
                continue  # skip blank rows

            # Resolve EventType (optional column)
            event_type = 'activity'
            et_raw = _cell('eventtype')
            if et_raw:
                et_raw_lower = et_raw.lower()
                if et_raw_lower not in VALID_EVENT_TYPES:
                    raise HTTPException(status_code=400, detail=f"Invalid eventType '{et_raw}'. Must be one of: {', '.join(VALID_EVENT_TYPES)}")
                event_type = et_raw_lower

            # Resolve EventTable (optional column)
            event_table = 'standard'
            etbl_raw = _cell('eventtable')
            if etbl_raw:
                etbl_raw_lower = etbl_raw.lower()
                if etbl_raw_lower not in VALID_EVENT_TABLES:
                    raise HTTPException(status_code=400, detail=f"Invalid eventTable '{etbl_raw}'. Must be one of: {', '.join(VALID_EVENT_TABLES)}")
                event_table = etbl_raw_lower

            if event_table == 'standard' and event_type != 'activity':
                raise HTTPException(status_code=400, detail=f"Event '{event_name}': standard event table must have eventType 'activity', got '{event_type}'")

            if data_type not in ['string', 'date', 'boolean', 'decimal', 'integer', 'int']:
                raise HTTPException(status_code=400, detail=f"Invalid datatype '{data_type}'. Must be one of: string, date, boolean, decimal, integer")

            if event_name in event_type_map and event_type_map[event_name] != event_type:
                raise HTTPException(status_code=400, detail=f"Conflicting eventType values for event '{event_name}'")
            event_type_map[event_name] = event_type

            if event_name in event_table_map and event_table_map[event_name] != event_table:
                raise HTTPException(status_code=400, detail=f"Conflicting eventTable values for event '{event_name}'")
            event_table_map[event_name] = event_table

            if event_name not in events_dict:
                events_dict[event_name] = []
            events_dict[event_name].append({"name": event_field, "datatype": data_type})

        if not events_dict:
            raise HTTPException(status_code=400, detail="No valid event definitions found in the 'events' sheet.")

        # ── Transactions sheet (optional) ──────────────────────────────────────
        transaction_types = []
        if 'transactions' in sheet_names_lower:
            txn_sheet = xl.sheet_names[sheet_names_lower.index('transactions')]
            try:
                txn_df = pd.read_excel(xl, sheet_name=txn_sheet, header=0, dtype=str)
                # Locate the transactiontype column (case-insensitive, ignore spaces/underscores)
                txn_col = None
                for col in txn_df.columns:
                    if str(col).strip().lower().replace(' ', '').replace('_', '') == 'transactiontype':
                        txn_col = col
                        break
                if txn_col is None and not txn_df.empty:
                    txn_col = txn_df.columns[0]  # fallback: first column
                if txn_col is not None:
                    transaction_types = [
                        str(v).strip() for v in txn_df[txn_col]
                        if v is not None and not (isinstance(v, float) and pd.isna(v)) and str(v).strip() not in ('', 'nan', 'None')
                    ]
            except Exception as exc:
                logger.warning(f"Could not read 'transactions' sheet: {exc}")

        # ── Store event definitions ────────────────────────────────────────────
        stored_in_db = False
        try:
            await db.event_definitions.delete_many({})
            for event_name, fields in events_dict.items():
                evt_type = event_type_map.get(event_name, 'activity')
                evt_table = event_table_map.get(event_name, 'standard')
                event = EventDefinition(event_name=event_name, fields=fields, eventType=evt_type, eventTable=evt_table)
                doc = event.model_dump()
                doc['created_at'] = doc['created_at'].isoformat()
                await db.event_definitions.insert_one(doc)
            stored_in_db = True
        except Exception as e:
            logger.warning(f"Could not write event definitions to DB, using in-memory storage: {e}")
            in_memory_defs = []
            for event_name, fields in events_dict.items():
                evt_type = event_type_map.get(event_name, 'activity')
                evt_table = event_table_map.get(event_name, 'standard')
                event = EventDefinition(event_name=event_name, fields=fields, eventType=evt_type, eventTable=evt_table)
                doc = event.model_dump()
                doc['created_at'] = doc['created_at'].isoformat()
                in_memory_defs.append(doc)
            in_memory_data['event_definitions'] = in_memory_defs

        # ── Store transaction definitions ──────────────────────────────────────
        try:
            await db.transaction_definitions.delete_many({})
            for txn_type in transaction_types:
                await db.transaction_definitions.insert_one({"transactiontype": txn_type})
        except Exception as e:
            logger.warning(f"Could not write transaction definitions to DB: {e}")
            in_memory_data['transaction_definitions'] = [{"transactiontype": t} for t in transaction_types]

        store_label = "database" if stored_in_db else "in-memory store"
        return {
            "message": f"Uploaded {len(events_dict)} event definition(s) and {len(transaction_types)} transaction type(s) to {store_label}",
            "events": list(events_dict.keys()),
            "transaction_types": transaction_types,
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error uploading reference data: {str(e)}")
        raise HTTPException(status_code=400, detail=str(e))

@api_router.get("/events")
async def get_events():
    """Get all event definitions"""
    try:
        events = await db.event_definitions.find({}, {"_id": 0}).to_list(1000)
        for event in events:
            if isinstance(event.get('created_at'), str):
                event['created_at'] = datetime.fromisoformat(event['created_at'])
        return events
    except Exception as e:
        logger.warning(f"Could not load events from database: {str(e)}")
        # Only return sample data if MongoDB is unavailable (connection error)
        logger.info("Returning sample events due to DB error")
        return SAMPLE_EVENTS

# DSL Functions (Hardcoded + Custom)
@api_router.get("/dsl-functions")
async def get_dsl_functions():
    """Get all DSL functions (hardcoded + custom)"""
    # Get hardcoded functions
    all_functions = list(DSL_FUNCTION_METADATA)
    
    # Get custom functions and convert to same format
    try:
        custom_funcs = await db.custom_functions.find({}, {"_id": 0}).to_list(1000)
        for func in custom_funcs:
            params = ', '.join([f"{p['name']}: {p['type']}" for p in func['parameters']])
            all_functions.append({
                "name": func['name'],
                "params": params,
                "description": func['description'],
                "category": func['category'],
                "is_custom": True
            })
    except Exception as e:
        logger.warning(f"Could not load custom functions from database: {str(e)}")
        # Include in-memory custom functions if DB unavailable
        for func in in_memory_data.get('custom_functions', []):
            try:
                params = ', '.join([f"{p['name']}: {p['type']}" for p in func.get('parameters', [])])
                all_functions.append({
                    "name": func.get('name'),
                    "params": params,
                    "description": func.get('description', ''),
                    "category": func.get('category', 'Custom'),
                    "is_custom": True
                })
            except Exception:
                continue
    
    return all_functions


# Download Event Definitions as CSV
@api_router.get("/transaction-definitions")
async def get_transaction_definitions():
    """Return all loaded transaction types from the Reference Data File."""
    try:
        try:
            docs = await db.transaction_definitions.find({}, {"_id": 0}).to_list(1000)
        except Exception:
            docs = in_memory_data.get('transaction_definitions', [])
        transaction_types = [d.get('transactiontype', '') for d in docs if d.get('transactiontype')]
        return {"transaction_types": transaction_types}
    except Exception as e:
        logger.error(f"Error fetching transaction definitions: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@api_router.get("/events/download")
async def download_event_definitions():
    """Download Reference Data File as .xlsx with two sheets: 'events' and 'transactions'."""
    try:
        # Load events
        try:
            events = await db.event_definitions.find({}, {"_id": 0}).to_list(1000)
        except Exception:
            events = in_memory_data.get('event_definitions', SAMPLE_EVENTS)

        # Load transaction definitions
        try:
            txn_docs = await db.transaction_definitions.find({}, {"_id": 0}).to_list(1000)
        except Exception:
            txn_docs = in_memory_data.get('transaction_definitions', [])
        transaction_types = [d.get('transactiontype', '') for d in txn_docs if d.get('transactiontype')]

        # Build events rows
        events_rows = []
        for event in events:
            evt_type = event.get('eventType', 'activity')
            evt_table = event.get('eventTable', 'standard')
            for field in event.get('fields', []):
                events_rows.append([event.get('event_name'), field.get('name'), field.get('datatype'), evt_type, evt_table])

        import openpyxl
        wb = openpyxl.Workbook()

        # Events sheet
        ws_events = wb.active
        ws_events.title = 'events'
        ws_events.append(['EventName', 'EventField', 'DataType', 'EventType', 'EventTable'])
        for row in events_rows:
            ws_events.append(row)

        # Transactions sheet
        ws_txn = wb.create_sheet('transactions')
        ws_txn.append(['transactiontype'])
        for t in transaction_types:
            ws_txn.append([t])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        xlsx_bytes = output.read()

        return Response(
            content=xlsx_bytes,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={"Content-Disposition": "attachment; filename=reference_data.xlsx"}
        )
    except Exception as e:
        logger.error(f"Error generating reference data xlsx: {e}")
        raise HTTPException(status_code=500, detail=str(e))

# Event Data - Excel Upload (multiple sheets for multiple events)
@api_router.post("/event-data/upload-excel")
async def upload_event_data_excel(file: UploadFile = File(...)):
    """Upload event data from Excel file - each sheet represents one event"""
    try:
        # NOTE: do NOT clear existing event data before validating the incoming file.
        # We will only replace data for specific events after full validation passes.
        if not file.filename.endswith(('.xlsx', '.xls')):
            raise HTTPException(status_code=400, detail="File must be an Excel file (.xlsx or .xls)")
        
        content = await file.read()
        
        # Read Excel file with all sheets
        excel_file = pd.ExcelFile(io.BytesIO(content))
        sheet_names = excel_file.sheet_names
        
        # First pass: Collect all postingdates across all sheets to validate single date
        # Only collect from activity events (not custom reference events which are tenant-level data)
        all_posting_dates = set()
        sheet_data_cache = {}  # Cache sheet data to avoid re-reading
        # Pre-fetch event definitions for each sheet to determine eventType/eventTable
        sheet_event_defs = {}
        
        for sheet_name in sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            sheet_data_cache[sheet_name] = df
            
            # Look up the event definition for this sheet
            try:
                evt_def = await db.event_definitions.find_one(
                    {"event_name": {"$regex": f"^{sheet_name}$", "$options": "i"}},
                    {"_id": 0}
                )
            except Exception:
                evt_def = next((e for e in in_memory_data.get('event_definitions', []) if str(e.get('event_name', '')).lower() == sheet_name.lower()), None)
            sheet_event_defs[sheet_name] = evt_def
            
            # Determine if this is a custom reference event (tenant-level, no instrument/date fields)
            is_reference = (evt_def and evt_def.get('eventTable') == 'custom' and evt_def.get('eventType') == 'reference')
            
            if df.empty or is_reference:
                continue
            
            # Look for postingdate column (case-insensitive)
            posting_col = None
            for col in df.columns:
                if str(col).lower() == 'postingdate':
                    posting_col = col
                    break
            
            if posting_col:
                # Extract all non-null postingdates
                posting_dates = df[posting_col].dropna().unique()
                for pd_val in posting_dates:
                    if pd_val and str(pd_val).strip():
                        # Normalize date format
                        date_str = str(pd_val).strip().split(' ')[0]  # Handle datetime strings
                        all_posting_dates.add(date_str)
        
        # Multiple posting dates are allowed: users often upload a file spanning
        # several periods at once (e.g. month-end snapshots). Downstream, each
        # posting date is processed independently via filter_event_data_by_posting_date.

        # Enforce maximum rows per sheet: do not proceed if any sheet exceeds the limit
        MAX_ROWS_PER_SHEET = 500
        for sheet_name, df in sheet_data_cache.items():
            try:
                row_count = int(df.shape[0])
            except Exception:
                row_count = 0
            if row_count > MAX_ROWS_PER_SHEET:
                raise HTTPException(status_code=400, detail="Upload failed: This file exceeds the allowed row limit. A maximum of 500 rows per table is supported.")
        
        def _normalize(s: str) -> str:
            import re
            return re.sub(r'[^A-Za-z0-9]', '_', (s or '').strip()).strip('_').upper()

        # Standard system columns accepted for activity events (not in event definition fields)
        ACTIVITY_STANDARD_NORM = {'INSTRUMENTID', 'POSTINGDATE', 'EFFECTIVEDATE', 'SUBINSTRUMENTID'}

        # Validation pass: check every sheet's columns against its event definition before saving anything
        header_errors = []
        for sheet_name in sheet_names:
            event = sheet_event_defs.get(sheet_name)
            if not event:
                continue  # will be caught as a missing-definition error below
            df = sheet_data_cache.get(sheet_name)
            if df is None or df.empty:
                continue
            is_reference = (event.get('eventTable') == 'custom' and event.get('eventType') == 'reference')
            field_names = [f['name'] for f in event.get('fields', [])]
            norm_to_field = {_normalize(fn): fn for fn in field_names}
            unknown = []
            for col in df.columns:
                col_norm = _normalize(str(col))
                if col_norm not in norm_to_field:
                    if not is_reference and col_norm in ACTIVITY_STANDARD_NORM:
                        continue  # allowed system column for activity events
                    unknown.append(str(col))
            if unknown:
                event_type_label = "reference" if is_reference else "activity"
                header_errors.append(
                    f"Sheet '{sheet_name}' ({event_type_label} event): unrecognized columns not in event definition — {', '.join(unknown)}"
                )

        if header_errors:
            raise HTTPException(
                status_code=400,
                detail={"message": "Upload rejected: columns do not match event definition", "errors": header_errors}
            )

        uploaded_events = []
        errors = []

        # Note: do not wipe all event data here; only replace data for the target event below.

        for sheet_name in sheet_names:
            # Use pre-fetched event definition from first pass
            event = sheet_event_defs.get(sheet_name)

            if not event:
                errors.append(f"Sheet '{sheet_name}' - No matching event definition found")
                continue

            # Determine if this is a custom reference event
            is_reference = (event.get('eventTable') == 'custom' and event.get('eventType') == 'reference')

            # Use cached data
            df = sheet_data_cache.get(sheet_name)

            if df is None or df.empty:
                errors.append(f"Sheet '{sheet_name}' - No data rows found")
                continue

            # Convert DataFrame to list of dicts and normalize headers to event field names
            df = df.fillna('')
            raw_rows = df.to_dict('records')

            field_names = [f['name'] for f in event.get('fields', [])]
            norm_to_field = { _normalize(fn): fn for fn in field_names }

            # Build header mapping (validation already passed — all columns are known)
            sheet_columns = list(df.columns)
            remapped_headers = []
            col_to_field = {}
            for col in sheet_columns:
                col_norm = _normalize(str(col))
                if col_norm in norm_to_field:
                    canonical = norm_to_field[col_norm]
                    col_to_field[col] = canonical
                    remapped_headers.append({"incoming": str(col), "mapped_to": canonical, "status": "mapped"})
                else:
                    # Must be a standard activity column (already validated above)
                    col_to_field[col] = str(col)
                    remapped_headers.append({"incoming": str(col), "mapped_to": str(col), "status": "standard"})

            data_rows = []
            for raw in raw_rows:
                mapped = {}
                for h, v in raw.items():
                    if h in col_to_field:
                        mapped[col_to_field[h]] = v
                data_rows.append(mapped)

            # Get field types from event definition
            field_types = {f['name']: f.get('datatype', 'string') for f in event.get('fields', [])}

            cleaned_rows = []
            # Track coercions summary: field -> coerced_count
            coercions = {}
            for row in data_rows:
                cleaned_row = {}
                for key, value in row.items():
                    field_type = field_types.get(str(key), 'string')
                    # Normalize NaN / empty / 'None' values
                    if pd.isna(value) or str(value).strip() == '' or str(value).strip().lower() in ('none', 'null'):
                        if field_type in ('decimal', 'float'):
                            cleaned_row[str(key)] = 0.0
                            coercions[str(key)] = coercions.get(str(key), 0) + 1
                        elif field_type in ('integer', 'int'):
                            cleaned_row[str(key)] = 0
                            coercions[str(key)] = coercions.get(str(key), 0) + 1
                        else:
                            cleaned_row[str(key)] = ''
                    # Date fields: normalize to yyyy-mm-dd (scalar or list)
                    elif str(field_type).lower() in ('date', 'datetime', 'timestamp'):
                        nv = _normalize_ingest_date_value(value)
                        cleaned_row[str(key)] = nv
                    elif field_type in ('decimal', 'float'):
                        try:
                            cleaned_row[str(key)] = float(value)
                        except Exception:
                            cleaned_row[str(key)] = 0.0
                            coercions[str(key)] = coercions.get(str(key), 0) + 1
                    elif field_type in ('integer', 'int'):
                        try:
                            cleaned_row[str(key)] = int(float(value))
                        except Exception:
                            cleaned_row[str(key)] = 0
                            coercions[str(key)] = coercions.get(str(key), 0) + 1
                    else:
                        cleaned_row[str(key)] = str(value)
                cleaned_rows.append(cleaned_row)
            # Normalize standard date fields only for non-reference events
            if not is_reference:
                for r in cleaned_rows:
                    for dkey in list(r.keys()):
                        if str(dkey).lower() in ('postingdate', 'effectivedate', 'posting_date', 'effective_date'):
                            r[dkey] = _normalize_ingest_date_value(r.get(dkey))
                    # Ensure standard date fields are normalized even if not declared as date type
                    for dkey in list(cleaned_row.keys()):
                        if str(dkey).lower() in ('postingdate', 'effectivedate', 'posting_date', 'effective_date'):
                            cleaned_row[dkey] = _normalize_ingest_date_value(cleaned_row.get(dkey))

                # Activity-data only: enforce canonical sort
                # (instrumentid ASC, postingdate ASC, effectivedate ASC, subinstrumentid ASC)
                # before the rows are persisted. Reference/custom event data is
                # intentionally skipped — it has no instrument/date axis.
                _sort_activity_rows(cleaned_rows)

            # Store event data
            event_data = EventData(event_name=event['event_name'], data_rows=cleaned_rows)
            doc = event_data.model_dump()
            doc['created_at'] = doc['created_at'].isoformat()
            
            # Replace existing data for this event
            await db.event_data.delete_many({"event_name": event['event_name']})
            await db.event_data.insert_one(doc)
            
            uploaded_events.append({
                "event_name": event['event_name'],
                "sheet_name": sheet_name,
                "rows_uploaded": len(cleaned_rows),
                "remapped_headers": remapped_headers,
                "coercions": coercions if coercions else None
            })
        
        if not all_posting_dates:
            posting_date_info = "No posting dates found"
        else:
            sorted_dates = sorted(all_posting_dates)
            posting_date_info = sorted_dates[0] if len(sorted_dates) == 1 else sorted_dates

        summary = {
            "message": f"Processed {len(sheet_names)} sheets",
            "posting_date": posting_date_info,
            "uploaded_events": uploaded_events,
            "errors": errors if errors else None
        }

        logger.info(f"Excel event data upload summary: {summary}")

        return summary
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error uploading Excel event data: {str(e)}")
        raise HTTPException(status_code=400, detail=str(e))


@api_router.get("/event-data")
async def get_all_event_data():
    """Get summary of all uploaded event data"""
    event_data_list = await db.event_data.find({}, {"_id": 0}).to_list(1000)
    summary = []
    for event_data in event_data_list:
        summary.append({
            "event_name": event_data['event_name'],
            "row_count": len(event_data.get('data_rows', [])),
            "created_at": event_data.get('created_at')
        })
    return summary


@api_router.get("/event-data/posting-dates")
async def get_event_data_posting_dates():
    """
    Return all unique posting dates found across all activity (non-custom/reference) event data,
    sorted ascending.  Custom / reference events have no posting date and are excluded.
    """
    # Identify activity event names
    defs = await db.event_definitions.find({}, {"_id": 0, "event_name": 1, "eventType": 1}).to_list(1000)
    activity_names: set = {
        d["event_name"] for d in defs
        if d.get("eventType", "activity") != "reference"
    }

    # Collect posting dates from activity event data
    unique_dates: set = set()
    event_data_list = await db.event_data.find({}, {"_id": 0, "event_name": 1, "data_rows": 1}).to_list(1000)
    for ed in event_data_list:
        if ed.get("event_name") not in activity_names:
            continue
        for row in ed.get("data_rows", []):
            pd_val = get_field_case_insensitive(row, "postingdate", "")
            if pd_val:
                unique_dates.add(str(pd_val).strip())

    return {"posting_dates": sorted(unique_dates)}


@api_router.get("/event-data/download/{event_name}")
async def download_event_data(event_name: str):
    """Download event data as CSV"""
    event_data = await db.event_data.find_one({"event_name": event_name}, {"_id": 0})
    if not event_data:
        raise HTTPException(status_code=404, detail=f"No data found for event '{event_name}'")

    # Create CSV. Collect ALL unique headers across every row so rows with different
    # field sets (common with JSON-imported data) do not cause DictWriter to crash.
    output = io.StringIO()
    rows = event_data.get('data_rows', [])
    if rows:
        all_keys: list = []
        seen_keys: set = set()
        for row in rows:
            for k in row.keys():
                if k not in seen_keys:
                    all_keys.append(k)
                    seen_keys.add(k)
        writer = csv.DictWriter(
            output, fieldnames=all_keys, extrasaction='ignore', restval=''
        )
        writer.writeheader()
        writer.writerows(rows)

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={event_name}_data.csv"}
    )


@api_router.get("/event-data/{event_name}")
async def get_event_data(event_name: str):
    """Get event data for a specific event"""
    event_data = await db.event_data.find_one({"event_name": event_name}, {"_id": 0})
    if not event_data:
        return {"event_name": event_name, "data_rows": []}
    
    if isinstance(event_data.get('created_at'), str):
        event_data['created_at'] = datetime.fromisoformat(event_data['created_at'])
    
    return event_data


@api_router.patch("/event-data/{event_name}/rows/{row_index}")
async def update_event_data_row(event_name: str, row_index: int, payload: Dict[str, Any] = Body(...)):
    """Update a single cell or multiple cells in a row of event data.

    Body shape: {"updates": {"column": value, ...}} — values are stored as-is
    after best-effort numeric coercion of stringy numbers.
    """
    updates = payload.get("updates") if isinstance(payload, dict) else None
    if not isinstance(updates, dict) or not updates:
        raise HTTPException(status_code=400, detail="Body must include non-empty 'updates' object")

    record = await db.event_data.find_one({"event_name": event_name})
    if not record:
        raise HTTPException(status_code=404, detail=f"Event data not found for '{event_name}'")

    rows = record.get("data_rows") or []
    if not isinstance(rows, list) or row_index < 0 or row_index >= len(rows):
        raise HTTPException(status_code=404, detail=f"Row index {row_index} out of range")

    def _coerce(v):
        if isinstance(v, str):
            s = v.strip()
            if s == "":
                return s
            try:
                if "." in s or "e" in s.lower():
                    return float(s)
                return int(s)
            except (ValueError, TypeError):
                return v
        return v

    target = dict(rows[row_index] or {})
    for k, v in updates.items():
        target[str(k)] = _coerce(v)
    rows[row_index] = target

    await db.event_data.update_one(
        {"event_name": event_name},
        {"$set": {"data_rows": rows}},
    )
    # Refresh in-memory caches used by DSL execution so subsequent runs
    # see the edited values.
    try:
        cursor = db.event_data.find({}, {"_id": 0})
        all_records = await cursor.to_list(length=None)
        raw = {r.get("event_name"): r.get("data_rows", []) for r in all_records if r.get("event_name")}
        set_raw_event_data(raw)
        set_all_event_data(merge_event_data_by_instrument(raw))
    except Exception as _e:
        logger.warning(f"Failed to refresh event-data caches after edit: {_e}")

    return {"event_name": event_name, "row_index": row_index, "row": target}

@api_router.post("/dsl/run")
async def run_dsl_code(request: DSLRunRequest):
    """Run DSL code directly and return results (for console testing)"""
    try:
        dsl_code = request.dsl_code
        
        # Extract all event names referenced in the DSL code
        referenced_events = extract_event_names_from_dsl(dsl_code)
        
        # If no event references, run in standalone mode (for schedule functions, calculations, etc.)
        if not referenced_events:
            # Create standalone execution template
            python_code = dsl_to_python_standalone(dsl_code)
            try:
                try:
                    compile(python_code, '<dsl_standalone>', 'exec')
                except SyntaxError as se:
                    # Log the problematic line for debugging
                    py_lines = python_code.split('\n')
                    err_lineno = se.lineno or 0
                    context_start = max(0, err_lineno - 3)
                    context_end = min(len(py_lines), err_lineno + 2)
                    context = '\n'.join(f"  {'>>>' if i+1 == err_lineno else '   '} {i+1}: {py_lines[i]}" for i in range(context_start, context_end))
                    logger.error(f"Syntax error in generated standalone code at Python line {err_lineno}:\n{context}")
                    # Try to extract DSL line from the error line
                    dsl_line = None
                    if 1 <= err_lineno <= len(py_lines):
                        m = re.search(r'# DSL_LINE:(\d+)', py_lines[err_lineno - 1])
                        if m:
                            dsl_line = int(m.group(1))
                    error_msg = se.msg or "invalid syntax"
                    if dsl_line:
                        error_msg = f"[Line {dsl_line}] SyntaxError: {error_msg}"
                    else:
                        error_msg = f"SyntaxError: {error_msg}"
                    return {
                        "success": False,
                        "error": error_msg,
                        "transactions": []
                    }
                # Provide minimal globals including __file__ so any code
                # referencing __file__ (e.g., os.path.dirname(__file__))
                # does not raise NameError when executed here.
                exec_globals = {
                    '__file__': os.path.abspath(__file__),
                    '__name__': '__dsl_standalone__',
                    '__builtins__': _make_sandbox_builtins(),
                }
                # Defense-in-depth: AST-validate the generated standalone
                # template before exec to block __import__, dunder
                # introspection, and disallowed module imports injected via
                # Custom Code.
                _validate_template_ast(python_code, label='<dsl_standalone>')
                exec(compile(python_code, '<dsl_standalone>', 'exec'), exec_globals)
                
                # Clear any previous print outputs
                clear_prints = exec_globals.get('clear_print_outputs')
                if clear_prints:
                    clear_prints()
                
                # Get the process function
                process_func = exec_globals.get('process_standalone')
                if not process_func:
                    raise ValueError("Generated code does not contain process_standalone function")
                
                # Execute standalone
                results, print_outputs = process_func(request.posting_date, request.effective_date)
                
                # Convert to TransactionOutput models
                transactions = [TransactionOutput(**result) for result in results]
                
                return {
                    "success": True,
                    "transactions": [t.model_dump() for t in transactions],
                    "events_used": [],
                    "row_count": 1,
                    "print_outputs": print_outputs,
                    "mode": "standalone"
                }
            except Exception as e:
                dsl_line = _extract_dsl_line_from_exception(python_code, e)
                error_msg = str(e)
                if dsl_line:
                    error_msg = f"[Line {dsl_line}] {error_msg}"
                logger.error(f"Standalone DSL error: {error_msg}")
                return {
                    "success": False,
                    "error": error_msg,
                    "transactions": []
                }
        
        # Load event definitions and data for all referenced events
        # Build event metadata and raw data maps. Respect eventType: activity vs reference
        all_event_fields = {}
        event_data_dict = {}
        activity_event_data = {}
        activity_events_with_data = []
        events_without_data = []
        reference_events_with_data = []

        for event_name in referenced_events:
            event_def = await db.event_definitions.find_one(
                {"event_name": {"$regex": f"^{event_name}$", "$options": "i"}}, 
                {"_id": 0}
            )
            if not event_def:
                return {
                    "success": False,
                    "error": f"Event definition '{event_name}' not found",
                    "transactions": []
                }

            # Store fields and eventType so generator can behave differently for reference events
            evt_name = event_def['event_name']
            evt_type = event_def.get('eventType', 'activity')
            all_event_fields[evt_name] = {
                'fields': event_def.get('fields', []),
                'eventType': evt_type
            }

            event_data = await db.event_data.find_one(
                {"event_name": {"$regex": f"^{event_name}$", "$options": "i"}}, 
                {"_id": 0}
            )
            rows = event_data['data_rows'] if (event_data and event_data.get('data_rows')) else []
            event_data_dict[evt_name] = rows

            if evt_type == 'activity':
                activity_event_data[evt_name] = rows
                if rows:
                    activity_events_with_data.append(evt_name)
                else:
                    events_without_data.append(evt_name)
            else:
                # reference event
                if rows:
                    reference_events_with_data.append(evt_name)
        
        # Determine merged rows to iterate over:
        # - If we have activity events with data, merge them by instrument
        # - If no activity data but reference events present, create a single dummy row so template runs once
        # - Otherwise, error (no data at all)
        date_fallback_warning = None
        if activity_events_with_data:
            # Filter by requested posting date before merging (Console date-scoped runs)
            scoped_activity = (
                filter_event_data_by_posting_date(activity_event_data, request.posting_date)
                if request.posting_date
                else activity_event_data
            )
            merged_data = merge_event_data_by_instrument(scoped_activity)
            
            # If no data for the requested posting date, fall back to all available data
            if not merged_data and request.posting_date:
                logger.info(f"No data for posting date {request.posting_date}, falling back to all available data")
                merged_data = merge_event_data_by_instrument(activity_event_data)
                if merged_data:
                    # Collect available posting dates for the warning
                    available_dates = set()
                    for rows in activity_event_data.values():
                        for row in rows:
                            pd = str(get_field_case_insensitive(row, "postingdate", "")).strip()
                            if pd:
                                available_dates.add(pd)
                    date_fallback_warning = (
                        f"No data found for posting date {request.posting_date}. "
                        f"Using all available data. Available posting dates: {sorted(available_dates)}"
                    )
        elif reference_events_with_data:
            # No activity rows but we have reference data — run template once with an empty merged row
            merged_data = [{}]
        else:
            return {
                "success": False,
                "error": f"No data found for any referenced events: {referenced_events}",
                "transactions": []
            }
        
        if not merged_data:
            return {
                "success": False,
                "error": "No data found after merging events",
                "transactions": []
            }
        
        # Generate and execute Python code. Pass event metadata (fields + eventType)
        python_code = dsl_to_python_multi_event(dsl_code, all_event_fields)
        # When a posting date is supplied (per-step tests use the earliest posting
        # date), also restrict the raw event data so collect_by_instrument() and
        # collect_all() — which normally span all dates — only see rows for that
        # date. collect() already filters by date, but the broader variants do not.
        # Reference events are passed through unchanged (they have no postingdate).
        raw_for_collect = (
            filter_event_data_by_posting_date(event_data_dict, request.posting_date, all_event_fields)
            if request.posting_date
            else event_data_dict
        )
        execution_result = await execute_python_template(
            python_code, 
            merged_data,
            raw_for_collect,  # Raw event data for collect() functions
            request.posting_date,
            request.effective_date
        )
        
        transactions = execution_result["transactions"]
        print_outputs = execution_result["print_outputs"]
        
        # Build events_used list: activity events with data + reference events with data
        events_used = activity_events_with_data + reference_events_with_data

        result = {
            "success": True,
            "transactions": [t.model_dump() for t in transactions],
            "events_used": events_used,
            "row_count": len(merged_data),
            "print_outputs": print_outputs
        }
        
        # Add warning if some events had no data
        if events_without_data:
            result["warning"] = f"No data for events: {events_without_data}. Their fields defaulted to 0/empty."
            result["events_without_data"] = events_without_data
        
        # Add warning if we fell back to all dates
        if date_fallback_warning:
            existing_warning = result.get("warning", "")
            result["warning"] = (existing_warning + " " + date_fallback_warning).strip()
        
        return result
    except HTTPException as he:
        # Re-extract the detail from execute_python_template which already includes [Line N]
        logger.error(f"DSL run error: {he.detail}")
        return {
            "success": False,
            "error": he.detail,
            "transactions": []
        }
    except Exception as e:
        logger.error(f"DSL run error: {str(e)}")
        return {
            "success": False,
            "error": str(e),
            "transactions": []
        }

# Templates
@api_router.post("/templates")
async def save_template(request: SaveTemplateRequest):
    """Save DSL code as a reusable template"""
    try:
        # Get event definition (try DB, fall back to in-memory)
        event = None
        try:
            event = await db.event_definitions.find_one({"event_name": request.event_name}, {"_id": 0})
        except Exception:
            # DB unavailable - check in-memory definitions
            logger.warning("DB unavailable when fetching event definition, checking in-memory data")
            for e in in_memory_data.get('event_definitions', []):
                if str(e.get('event_name', '')).lower() == request.event_name.lower():
                    event = e
                    break

        if not event:
            raise HTTPException(status_code=404, detail=f"Event '{request.event_name}' not found")

        # Check for existing template with same name
        # Check for existing template with same name (DB first, then in-memory)
        existing = None
        try:
            existing = await db.dsl_templates.find_one({"name": request.name}, {"_id": 0})
        except Exception:
            for t in in_memory_data.get('templates', []):
                if t.get('name', '').lower() == request.name.lower():
                    existing = t
                    break

        if existing:
            if not request.replace:
                raise HTTPException(
                    status_code=409,
                    detail=f"Template with name '{request.name}' already exists. Set replace=true to overwrite."
                )
            # Delete existing template (DB or in-memory)
            try:
                await db.dsl_templates.delete_one({"name": request.name})
            except Exception:
                in_memory_data['templates'] = [t for t in in_memory_data.get('templates', []) if t.get('name', '').lower() != request.name.lower()]

        # Convert DSL to Python using event fields with datatypes (deterministic)
        python_code = dsl_to_python(request.dsl_code, event['fields'])

        # Save template (document stores DSL + latest python_code for convenience)
        template = DSLTemplate(name=request.name, dsl_code=request.dsl_code, python_code=python_code)
        doc = template.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        try:
            await db.dsl_templates.insert_one(doc)
        except Exception:
            # DB unavailable - store in-memory
            global USE_IN_MEMORY
            USE_IN_MEMORY = True
            in_memory_data.setdefault('templates', []).append(doc)

        # Note: dsl_template_artifacts is intentionally NOT written here.
        # The runtime artifact collection is populated only by the explicit
        # Deploy action (POST /user-templates/{id}/deploy).

        return {"message": "Template saved successfully", "template_id": template.id, "replaced": existing is not None}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error saving template: {str(e)}")
        raise HTTPException(status_code=400, detail=str(e))

@api_router.get("/templates/check-name/{name}")
async def check_template_name(name: str):
    """Check if a template name already exists"""
    existing = await db.dsl_templates.find_one({"name": name}, {"_id": 0})
    return {"exists": existing is not None}

@api_router.get("/templates")
async def get_templates():
    """Get all saved templates"""
    try:
        templates = await db.dsl_templates.find({}, {"_id": 0}).to_list(1000)
        if templates:
            # Sanitize for JSON serialization
            return [sanitize_for_json(t) for t in templates]
    except Exception as e:
        logger.warning(f"Could not load templates from database: {str(e)}")
    
    # Return sample data if MongoDB is unavailable
    logger.info("Returning sample templates")
    return [sanitize_for_json(t) for t in SAMPLE_TEMPLATES]

@api_router.delete("/templates/{template_id}")
async def delete_template(template_id: str):
    """Delete a template (robust: tries id, name, ObjectId, in-memory, and sample list)"""
    try:
        # 1) Try deleting by id (DB may be unavailable)
        try:
            result = await db.dsl_templates.delete_one({"id": template_id})
            logger.info(f"delete_one by id result: {getattr(result, 'deleted_count', 'n/a')}")
            if getattr(result, 'deleted_count', 0) == 1:
                logger.info(f"Deleted template by id: {template_id}")
                # Also remove persisted artifacts
                try:
                    await db.dsl_template_artifacts.delete_many({"template_id": template_id})
                except Exception:
                    logger.debug(f"Failed to delete artifacts for template {template_id}")
                return {"message": "Template deleted successfully"}
        except Exception as e:
            logger.debug(f"DB delete by id failed, will try other strategies: {e}")

        # 2) Try deleting by name (in case caller passed a human-readable name)
        try:
            result_by_name = await db.dsl_templates.delete_one({"name": template_id})
            logger.info(f"delete_one by name result: {getattr(result_by_name, 'deleted_count', 'n/a')}")
            if getattr(result_by_name, 'deleted_count', 0) == 1:
                logger.info(f"Deleted template by name: {template_id}")
                # Also remove persisted artifacts by template_name
                try:
                    await db.dsl_template_artifacts.delete_many({"template_name": template_id})
                except Exception:
                    logger.debug(f"Failed to delete artifacts for template name {template_id}")
                return {"message": "Template deleted successfully (by name)"}
        except Exception as e:
            logger.debug(f"DB delete by name failed, will try other strategies: {e}")

        # 3) Try deleting by _id if an ObjectId string was provided
        try:
            from bson import ObjectId
            obj_id = ObjectId(template_id)
            try:
                result_by_obj = await db.dsl_templates.delete_one({"_id": obj_id})
                logger.info(f"delete_one by _id result: {getattr(result_by_obj, 'deleted_count', 'n/a')}")
                if getattr(result_by_obj, 'deleted_count', 0) == 1:
                    logger.info(f"Deleted template by _id: {template_id}")
                    try:
                        await db.dsl_template_artifacts.delete_many({"template_id": template_id})
                    except Exception:
                        logger.debug(f"Failed to delete artifacts for template _id {template_id}")
                    return {"message": "Template deleted successfully (by _id)"}
            except Exception as e:
                logger.debug(f"DB delete by _id failed, will try other strategies: {e}")
        except Exception as e:
            logger.debug(f"Not an ObjectId or failed delete by _id: {str(e)}")

        # 4) If not found in DB, try in-memory storage (useful for local/dev mode)
        if USE_IN_MEMORY:
            before = len(in_memory_data.get("templates", []))
            in_memory_data["templates"] = [t for t in in_memory_data.get("templates", []) if t.get("id") != template_id and t.get("name") != template_id]
            after = len(in_memory_data.get("templates", []))
            if after < before:
                logger.info(f"Deleted template {template_id} from in-memory storage")
                # Also remove in-memory artifacts for this template
                in_memory_data['template_artifacts'] = [a for a in in_memory_data.get('template_artifacts', []) if a.get('template_id') != template_id and a.get('template_name') != template_id]
                return {"message": "Template deleted successfully (in-memory)"}

        # 5) As a last resort, remove from SAMPLE_TEMPLATES
        global SAMPLE_TEMPLATES
        sample_before = len(SAMPLE_TEMPLATES)
        SAMPLE_TEMPLATES = [t for t in SAMPLE_TEMPLATES if t.get("id") != template_id and t.get("name") != template_id]
        if len(SAMPLE_TEMPLATES) < sample_before:
            logger.info(f"Deleted template {template_id} from SAMPLE_TEMPLATES")
            return {"message": "Template deleted successfully (sample data)"}

        # Nothing deleted
        raise HTTPException(status_code=404, detail="Template not found")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting template {template_id}: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@api_router.post("/templates/deploy")
async def deploy_template(request: TemplateDeployRequest):
    """Mark a template as deployed."""
    try:
        result = await db.dsl_templates.update_one(
            {"id": request.template_id},
            {"$set": {"deployed": True, "deployed_at": datetime.now(timezone.utc).isoformat()}}
        )
        if result.matched_count == 0:
            raise HTTPException(status_code=404, detail="Template not found")
        return {"success": True, "message": f"Template {request.template_id} deployed successfully"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deploying template {request.template_id}: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@api_router.post("/templates/execute")
async def execute_template(request: TemplateExecuteRequest):
    """Execute a saved template on event data - supports multiple events"""
    try:
        # Get template (DB first, then in-memory/sample)
        template = None
        try:
            template = await db.dsl_templates.find_one({"id": request.template_id}, {"_id": 0})
        except Exception:
            logger.debug("DB unavailable when fetching template for execution; checking in-memory and sample templates")

        if not template:
            # Check in-memory templates
            for t in in_memory_data.get('templates', []):
                if t.get('id') == request.template_id or t.get('name') == request.template_id:
                    template = t
                    break

        if not template:
            # Check sample templates
            for t in SAMPLE_TEMPLATES:
                if t.get('id') == request.template_id or t.get('name') == request.template_id:
                    template = t
                    break

        if not template:
            raise HTTPException(status_code=404, detail="Template not found")
        
        dsl_code = template['dsl_code']
        
        # Extract all event names referenced in the DSL code
        referenced_events = extract_event_names_from_dsl(dsl_code)
        logger.info(f"Referenced events in DSL: {referenced_events}")
        
        # If no events found in DSL (old format), use the selected event
        if not referenced_events:
            referenced_events = [request.event_name]
        
        # Load event definitions and data for all referenced events
        all_event_fields = {}
        event_data_dict = {}
        
        for event_name in referenced_events:
            # Get event definition
            event_def = await db.event_definitions.find_one(
                {"event_name": {"$regex": f"^{event_name}$", "$options": "i"}}, 
                {"_id": 0}
            )
            if not event_def:
                raise HTTPException(status_code=404, detail=f"Event definition '{event_name}' not found")
            
            all_event_fields[event_def['event_name']] = event_def['fields']
            
            # Get event data
            event_data = await db.event_data.find_one(
                {"event_name": {"$regex": f"^{event_name}$", "$options": "i"}}, 
                {"_id": 0}
            )
            if event_data and event_data.get('data_rows'):
                event_data_dict[event_def['event_name']] = event_data['data_rows']
            else:
                logger.warning(f"No data found for event '{event_name}'")
                event_data_dict[event_def['event_name']] = []
        
        # Merge data from all events by instrumentid
        # Filter by requested posting date before merging (batch date-scoped runs)
        scoped_event_data = (
            filter_event_data_by_posting_date(event_data_dict, request.posting_date)
            if request.posting_date
            else event_data_dict
        )
        merged_data = merge_event_data_by_instrument(scoped_event_data)
        
        if not merged_data:
            raise HTTPException(status_code=404, detail="No data found for the referenced events")
        
        logger.info(f"Merged data: {len(merged_data)} rows from {len(referenced_events)} events")
        
        # Generate Python code for multi-event template
        python_code = dsl_to_python_multi_event(dsl_code, all_event_fields)
        
        # Execute template with optional date overrides
        execution_result = await execute_python_template(
            python_code, 
            merged_data,
            event_data_dict,  # Pass raw event data for collect() functions
            request.posting_date,
            request.effective_date
        )
        
        transactions = execution_result["transactions"]
        print_outputs = execution_result["print_outputs"]
        
        # Save transaction report (DB or in-memory) — append each batch
        transaction_dicts = [t.model_dump() for t in transactions]
        report = TransactionReport(
            template_name=template.get('name', ''),
            event_name=', '.join(referenced_events),
            transactions=transaction_dicts
        )
        doc = report.model_dump()
        doc['executed_at'] = doc['executed_at'].isoformat()
        try:
            # Append new batch — each execution gets its own document
            await db.transaction_reports.insert_one(doc)
        except Exception:
            # Fallback to in-memory storage
            global USE_IN_MEMORY  # noqa: F811
            USE_IN_MEMORY = True
            lst = in_memory_data.setdefault('transaction_reports', [])
            lst.append(doc)
            in_memory_data['transaction_reports'] = lst
        
        return {
            "message": "Template executed successfully",
            "report_id": report.id,
            "transactions": transaction_dicts,
            "events_used": referenced_events,
            "print_outputs": print_outputs
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error executing template: {str(e)}")
        raise HTTPException(status_code=400, detail=str(e))


@api_router.delete("/transaction-reports/all")
async def delete_all_transaction_reports():
    """Wipe all transaction reports"""
    try:
        result = await db.transaction_reports.delete_many({})
        return {"message": f"Deleted {result.deleted_count} report(s)"}
    except Exception:
        in_memory_data['transaction_reports'] = []
        return {"message": "Cleared in-memory reports"}

@api_router.delete("/transaction-reports/{report_id}")
async def delete_transaction_report(report_id: str):
    """Delete all batches for a transaction report"""
    anchor = await db.transaction_reports.find_one({"id": report_id}, {"_id": 0})
    if not anchor:
        raise HTTPException(status_code=404, detail="Report not found")
    # Delete all batches belonging to the same template
    template_name = anchor.get('template_name', '')
    result = await db.transaction_reports.delete_many({"template_name": template_name})
    return {"message": f"Deleted {result.deleted_count} batch(es) for '{template_name}'"}


@api_router.post("/ai/provider/test")
async def test_ai_provider(req: AIProviderTestRequest):
    """Validate an API key and return available models."""
    if req.provider not in PROVIDER_INFO:
        raise HTTPException(status_code=400, detail=f"Unknown provider: {req.provider}")
    try:
        provider = get_provider(req.provider)
        models = await provider.list_models(req.api_key)
        return {
            "valid": True,
            "models": [m.model_dump() for m in models],
        }
    except AIError as e:
        logger.warning(f"AI provider test failed ({req.provider}): {e.error_type} - {e.detail}")
        return {
            "valid": False,
            "error_type": e.error_type,
            "error_message": e.detail,
            "models": [],
        }
    except Exception as e:
        logger.exception(f"Unexpected error testing AI provider ({req.provider})")
        return {
            "valid": False,
            "error_type": "network",
            "error_message": str(e),
            "models": [],
        }

@api_router.post("/ai/provider/save")
async def save_ai_provider(req: AIProviderSaveRequest):
    """Persist the selected provider, encrypted API key, and models."""
    encrypted = encrypt_key(req.api_key)
    doc = {
        "provider": req.provider,
        "encrypted_api_key": encrypted,
        "selected_model": req.selected_model,
        "available_models": req.available_models,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.ai_provider_config.delete_many({})
    await db.ai_provider_config.insert_one(doc)
    return {"success": True}

@api_router.post("/ai/provider/selected-model")
async def set_selected_model(payload: dict = Body(...)):
    """Persist ONLY the selected model (no API key needed). Called when the
    user picks a model in the chat's model dropdown so the choice becomes the
    durable default — surviving hard refresh and localStorage clears.

    We deliberately do NOT validate the id against the cached available_models
    list: that cache can lag behind the live list the dropdown is populated
    from, and rejecting a valid selection would silently drop the choice and
    revert the model on the next refresh. The dropdown already constrains the
    user to real models, so we trust the id and just store it."""
    model_id = (payload or {}).get("selected_model")
    if not model_id or not isinstance(model_id, str):
        raise HTTPException(status_code=400, detail="selected_model (string) is required")
    try:
        result = await db.ai_provider_config.update_one(
            {}, {"$set": {"selected_model": model_id,
                          "updated_at": datetime.now(timezone.utc).isoformat()}},
        )
    except Exception as exc:
        logger.error("Failed to persist selected model: %s", exc)
        raise HTTPException(status_code=500, detail="Could not save the selected model")
    if getattr(result, "matched_count", 0) == 0:
        raise HTTPException(status_code=409, detail="No AI provider configured yet")
    return {"success": True, "selected_model": model_id}


@api_router.get("/ai/provider/status")
async def get_ai_provider_status():
    """Get the current provider config with dynamically refreshed model list."""
    try:
        config = await db.ai_provider_config.find_one({}, {"_id": 0})
    except Exception:
        config = None
    if not config:
        return {"configured": False}

    provider_name = config.get("provider")
    cached_models = config.get("available_models", [])

    # Dynamically refresh models from the provider API
    fresh_models = cached_models
    try:
        api_key = decrypt_key(config["encrypted_api_key"])
        provider = get_provider(provider_name)
        models = await provider.list_models(api_key)
        fresh_models = [m.model_dump() for m in models]
        # Update DB cache in the background
        await db.ai_provider_config.update_one(
            {"provider": provider_name},
            {"$set": {"available_models": fresh_models}},
        )
    except Exception as exc:
        logger.warning(f"Failed to refresh model list for {provider_name}: {exc}")
        # Fall back to cached models

    return {
        "configured": True,
        "provider": provider_name,
        "selected_model": config.get("selected_model"),
        "available_models": fresh_models,
    }

@api_router.delete("/ai/provider")
async def delete_ai_provider():
    """Remove the saved provider configuration."""
    await db.ai_provider_config.delete_many({})
    return {"success": True}

# AI Chat Assistant
@api_router.post("/chat", response_model=ChatResponse)
async def chat_with_assistant(message: ChatMessage):
    """Chat with AI assistant for DSL help - with full context awareness"""
    try:
        session_id = message.session_id or str(uuid.uuid4())

        # --- Load provider config ---
        try:
            provider_config = await db.ai_provider_config.find_one({}, {"_id": 0})
        except Exception:
            provider_config = None

        if not provider_config:
            return ChatResponse(
                response="",
                session_id=session_id,
                error_type="no_provider",
                error_message=ERROR_MESSAGES["no_provider"],
            )

        provider_name = provider_config.get("provider", "")
        selected_model = message.model or provider_config.get("selected_model", "")
        provider_display = PROVIDER_INFO.get(provider_name, {}).get("name", provider_name)

        # Decrypt key
        try:
            api_key = decrypt_key(provider_config["encrypted_api_key"])
        except Exception as e:
            logger.warning(f"Failed to decrypt API key: {e}")
            return ChatResponse(
                response="",
                session_id=session_id,
                error_type="invalid_key",
                error_message=ERROR_MESSAGES["invalid_key"],
            )

        # --- Gather context data ---
        if message.context and message.context.get('events'):
            events = message.context['events']
        else:
            try:
                events = await db.event_definitions.find({}, {"_id": 0}).to_list(1000)
            except Exception:
                events = in_memory_data.get('event_definitions', SAMPLE_EVENTS)

        editor_code = ""
        if message.context and message.context.get('editor_code'):
            editor_code = message.context['editor_code']

        console_output = []
        if message.context and message.context.get('console_output'):
            console_output = message.context['console_output']

        # Rich editor context (cursor, selection, syntax errors)
        editor_cursor = message.context.get('editor_cursor') if message.context else None
        editor_selection = message.context.get('editor_selection') if message.context else None
        editor_syntax_errors = message.context.get('editor_syntax_errors') if message.context else None
        ui_mode = message.context.get('ui_mode') if message.context else None

        # --- Build system prompt via two-tier context engine ---
        system_prompt = build_agent_context(
            dsl_function_metadata=list(DSL_FUNCTION_METADATA),
            events=events,
            editor_code=editor_code,
            editor_cursor=editor_cursor,
            editor_selection=editor_selection,
            editor_syntax_errors=editor_syntax_errors,
            console_output=console_output,
            conversation_history=message.history,
            ui_mode=ui_mode,
        )

        # --- Call the AI provider ---
        try:
            provider = get_provider(provider_name)
            ai_response = await provider.chat(
                api_key=api_key,
                model_id=selected_model,
                system_prompt=system_prompt,
                user_message=message.message,
                history=message.history,
            )
            response_text = ai_response.text
        except AIError as e:
            err_msg = ERROR_MESSAGES.get(e.error_type, e.detail)
            err_msg = err_msg.replace("{provider}", provider_display).replace("{model}", selected_model)
            return ChatResponse(
                response="",
                session_id=session_id,
                error_type=e.error_type,
                error_message=err_msg,
            )

        # --- Try to parse structured JSON response ---
        structured = None
        try:
            import re as _re
            json_match = _re.search(r'```json\s*(\{.*?\})\s*```', response_text, _re.S)
            if json_match:
                parsed = json.loads(json_match.group(1))
                if "dsl_code" in parsed and "explanation" in parsed:
                    structured = {
                        "explanation": parsed.get("explanation", ""),
                        "dsl_code": parsed.get("dsl_code", ""),
                        "insert_mode": parsed.get("insert_mode", "append"),
                        "confidence": parsed.get("confidence", "medium"),
                    }
        except Exception:
            pass

        # --- Post-process response to enforce DSL rules ---
        try:
            import re

            user_msg_lower = (message.message or '').lower()
            user_requested_transactions = any(k in user_msg_lower for k in [
                'createtransaction', 'create transaction', 'createtransactions', 'create transactions', 'include transaction', 'emit transaction', 'include createtransaction'
            ])

            def replace_leading_comments(text: str) -> str:
                return re.sub(r'(^|\n)\s*//', r"\1##", text)

            def process_code_block(code: str) -> str:
                code = replace_leading_comments(code)

                allowed_funcs = set()
                try:
                    allowed_funcs.update(DSL_FUNCTIONS.keys())
                except Exception:
                    pass
                try:
                    for m in DSL_FUNCTION_METADATA:
                        name = m.get('name') if isinstance(m, dict) else None
                        if name:
                            allowed_funcs.add(name)
                except Exception:
                    pass
                extra_allowed = {'print', 'iif', 'collect_by_instrument', 'collect_all', 'collect_by_subinstrument', 'collect_effectivedates_for_subinstrument', 'npv', 'irr', 'sum_field', 'sum', 'len', 'min', 'max', 'abs', 'round', 'lag'}
                allowed_funcs.update(extra_allowed)

                lines = code.splitlines()
                cleaned_lines = []
                for ln in lines:
                    if re.match(r"^\s*(def|class|import|from)\b", ln):
                        cleaned_lines.append('## removed unsupported Python construct')
                        continue
                    if re.match(r"^\s*(for|while)\b.*:\s*$", ln):
                        cleaned_lines.append('## removed unsupported Python loop')
                        continue
                    cleaned_lines.append(ln)
                code = "\n".join(cleaned_lines)

                has_create = re.search(r"\bcreateTransactions?\s*\(", code)
                if has_create and not user_requested_transactions:
                    lines = code.splitlines()
                    lines = [ln for ln in lines if not re.search(r"\bcreateTransactions?\s*\(", ln)]
                    code = "\n".join(lines)
                    assigns = re.findall(r'^\s*([a-z_][a-zA-Z0-9_]*)\s*=.*$', code, flags=re.MULTILINE)
                    if assigns:
                        last_var = assigns[-1]
                        if not re.search(r'print\s*\(\s*' + re.escape(last_var) + r'\s*\)', code):
                            code = code.rstrip() + '\n\nprint(' + last_var + ')'
                    try:
                        code = re.sub(r'(?mi)^\s*##.*create.*transaction.*$', '## Executing the final value', code, flags=re.M)
                    except Exception:
                        pass

                alias_map = {'periods': 'nper', 'num_periods': 'nper', 'number_of_periods': 'nper'}
                for a, b in alias_map.items():
                    code = re.sub(rf"\b{a}\s*\(", f"{b}(", code)

                non_comment_lines = [ln for ln in code.splitlines() if not ln.strip().startswith('##')]
                func_calls = []
                for ln in non_comment_lines:
                    func_calls.extend(re.findall(r"\b([a-zA-Z_][a-zA-Z0-9_]*)\s*\(", ln))
                if func_calls:
                    lines = code.splitlines()
                    new_lines = []
                    for ln in lines:
                        if ln.strip().startswith('##'):
                            new_lines.append(ln)
                            continue
                        code_only = ln.split('##')[0].split('//')[0]
                        called = re.findall(r"\b([a-zA-Z_][a-zA-Z0-9_]*)\s*\(", code_only)
                        if called:
                            illegal = [f for f in called if f not in allowed_funcs]
                            if illegal:
                                new_lines.append('## removed call to unsupported function: ' + ','.join(illegal))
                                continue
                        new_lines.append(ln)
                    code = "\n".join(new_lines)

                return code

            def process_response(text: str) -> str:
                def repl(match):
                    inner = match.group(1)
                    processed = process_code_block(inner)
                    return '```dsl\n' + processed + '\n```'

                text = re.sub(r'```(?:dsl)?\n(.*?)\n```', repl, text, flags=re.S)
                text = replace_leading_comments(text)

                if not user_requested_transactions and re.search(r"\bcreateTransactions?\s*\(", text):
                    lines = text.splitlines()
                    new_lines = []
                    for ln in lines:
                        if re.search(r"\bcreateTransactions?\s*\(", ln):
                            continue
                        new_lines.append(ln)
                    text = "\n".join(new_lines)
                    assigns = re.findall(r'^\s*([a-z_][a-zA-Z0-9_]*)\s*=.*$', text, flags=re.MULTILINE)
                    if assigns:
                        last_var = assigns[-1]
                        if not re.search(r'print\s*\(\s*' + re.escape(last_var) + r'\s*\)', text):
                            text = text.rstrip() + '\n\nprint(' + last_var + ')'

                return text

            response_text = process_response(response_text)

            # Also post-process structured dsl_code if present
            if structured and structured.get("dsl_code"):
                structured["dsl_code"] = process_code_block(structured["dsl_code"])

        except Exception as e:
            logger.warning(f"Post-processing of AI response failed: {e}")

        return ChatResponse(
            response=response_text,
            session_id=session_id,
            structured=structured,
        )
    except Exception as e:
        logger.error(f"Chat error: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Chat error: {str(e)}")


@api_router.post("/chat/stream")
async def chat_stream(message: ChatMessage):
    """SSE streaming chat endpoint — sends tokens as they arrive from the provider."""

    session_id = message.session_id or str(uuid.uuid4())

    async def event_stream():
        try:
            # Load provider config
            try:
                provider_config = await db.ai_provider_config.find_one({}, {"_id": 0})
            except Exception:
                provider_config = None

            if not provider_config:
                yield f"data: {json.dumps({'type': 'error', 'error_type': 'no_provider', 'error_message': ERROR_MESSAGES['no_provider']})}\n\n"
                yield "data: [DONE]\n\n"
                return

            provider_name = provider_config.get("provider", "")
            selected_model = message.model or provider_config.get("selected_model", "")
            provider_display = PROVIDER_INFO.get(provider_name, {}).get("name", provider_name)

            # Decrypt key
            try:
                api_key = decrypt_key(provider_config["encrypted_api_key"])
            except Exception:
                yield f"data: {json.dumps({'type': 'error', 'error_type': 'invalid_key', 'error_message': ERROR_MESSAGES['invalid_key']})}\n\n"
                yield "data: [DONE]\n\n"
                return

            # Emit session info
            yield f"data: {json.dumps({'type': 'session', 'session_id': session_id})}\n\n"

            # Gather context
            if message.context and message.context.get('events'):
                events = message.context['events']
            else:
                try:
                    events = await db.event_definitions.find({}, {"_id": 0}).to_list(1000)
                except Exception:
                    events = in_memory_data.get('event_definitions', SAMPLE_EVENTS)

            editor_code = ""
            if message.context and message.context.get('editor_code'):
                editor_code = message.context['editor_code']

            console_output = []
            if message.context and message.context.get('console_output'):
                console_output = message.context['console_output']

            # Rich editor context
            editor_cursor = message.context.get('editor_cursor') if message.context else None
            editor_selection = message.context.get('editor_selection') if message.context else None
            editor_syntax_errors = message.context.get('editor_syntax_errors') if message.context else None
            ui_mode = message.context.get('ui_mode') if message.context else None

            # Build system prompt via two-tier context engine
            system_prompt = build_agent_context(
                dsl_function_metadata=list(DSL_FUNCTION_METADATA),
                events=events,
                editor_code=editor_code,
                editor_cursor=editor_cursor,
                editor_selection=editor_selection,
                editor_syntax_errors=editor_syntax_errors,
                console_output=console_output,
                conversation_history=message.history,
                ui_mode=ui_mode,
            )

            # Emit context-ready event with summary for the UI
            events_count = len(events) if events else 0
            editor_lines = len(editor_code.split('\n')) if editor_code.strip() else 0
            console_count = len(console_output) if console_output else 0
            yield f"data: {json.dumps({'type': 'context_ready', 'events_count': events_count, 'editor_lines': editor_lines, 'console_count': console_count, 'model': selected_model, 'provider': provider_display})}\n\n"

            # Stream from provider
            provider = get_provider(provider_name)
            full_text = []
            try:
                async for chunk in provider.stream_chat(
                    api_key=api_key,
                    model_id=selected_model,
                    system_prompt=system_prompt,
                    user_message=message.message,
                    history=message.history,
                ):
                    full_text.append(chunk)
                    yield f"data: {json.dumps({'type': 'token', 'token': chunk})}\n\n"
            except AIError as e:
                err_msg = ERROR_MESSAGES.get(e.error_type, e.detail)
                err_msg = err_msg.replace("{provider}", provider_display).replace("{model}", selected_model)
                yield f"data: {json.dumps({'type': 'error', 'error_type': e.error_type, 'error_message': err_msg})}\n\n"
                yield "data: [DONE]\n\n"
                return

            # Post-process the full response (same rules as /chat)
            full_response = ''.join(full_text)
            try:
                import re as _pp_re
                user_msg_lower = (message.message or '').lower()
                user_requested_txn = any(k in user_msg_lower for k in [
                    'createtransaction', 'create transaction', 'createtransactions',
                    'create transactions', 'include transaction', 'emit transaction',
                ])

                def _pp_replace_comments(text):
                    return _pp_re.sub(r'(^|\n)\s*//', r'\1##', text)

                if not user_requested_txn and _pp_re.search(r'\bcreateTransactions?\s*\(', full_response):
                    # AI included transactions the user didn't ask for — flag it
                    yield f"data: {json.dumps({'type': 'post_process', 'warning': 'unrequested_transactions'})}\n\n"
            except Exception:
                pass

            # Send done event
            yield f"data: {json.dumps({'type': 'done'})}\n\n"
            yield "data: [DONE]\n\n"

        except Exception as e:
            logger.error(f"Stream chat error: {str(e)}")
            yield f"data: {json.dumps({'type': 'error', 'error_type': 'network', 'error_message': str(e)})}\n\n"
            yield "data: [DONE]\n\n"

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


REQUIRED_EVENT_FIELDS = {"instrumentId", "eventId", "eventName", "postingDate", "effectiveDate", "status", "eventDetail", "_class"}


# ──────────────────────────────────────────────────────────────────────────
# Autonomous agent endpoints
# ──────────────────────────────────────────────────────────────────────────

# Wire the agent's late-bound helpers exactly once at import time.
try:
    agent_configure_bridge(
        db=db,
        in_memory_data=in_memory_data,
        use_in_memory_getter=lambda: USE_IN_MEMORY,
        helpers={
            "EventDefinition": EventDefinition,
            "EventData": EventData,
            "DSLTemplate": DSLTemplate,
            "DSL_FUNCTION_METADATA": DSL_FUNCTION_METADATA,
            "extract_event_names_from_dsl": extract_event_names_from_dsl,
            "merge_event_data_by_instrument": merge_event_data_by_instrument,
            "filter_event_data_by_posting_date": filter_event_data_by_posting_date,
            "dsl_to_python": dsl_to_python,
            "dsl_to_python_multi_event": dsl_to_python_multi_event,
            "dsl_to_python_standalone": dsl_to_python_standalone,
            "execute_python_template": execute_python_template,
        },
    )
except Exception as _agent_bridge_exc:
    logger.warning("Agent bridge configuration failed: %s", _agent_bridge_exc)


class AgentRunRequest(BaseModel):
    task: str
    model: Optional[str] = None
    max_steps: Optional[int] = 80
    auto_approve_destructive: Optional[bool] = False
    session_id: Optional[str] = None


class AgentApprovalRequest(BaseModel):
    decision: str  # "approve" | "deny"


class AgentRuleApprovalRequest(BaseModel):
    """Maker-checker decision on an agent-authored rule."""
    checker: str                     # identity of the human reviewer
    note: Optional[str] = ""         # rationale / rejection reason


@api_router.post("/agent/run")
async def agent_run_endpoint(req: AgentRunRequest):
    """SSE stream of an autonomous agent run.

    Each SSE event has `data: <json>\\n\\n`. The terminal event has
    `data: [DONE]\\n\\n`.
    """
    async def event_stream():
        # 16 KB pad + immediate "connected" event flushes proxy buffers
        # (including the GitHub Codespaces port-forwarder, which buffers
        # ~8 KB) so the UI shows progress before the first model call
        # returns. Repeat the connected event after the pad to be safe.
        yield ":" + (" " * 16384) + "\n\n"
        yield f"data: {json.dumps({'type': 'connected', 'ts': datetime.now(timezone.utc).isoformat()})}\n\n"
        yield ":" + (" " * 4096) + "\n\n"
        try:
            try:
                provider_config = await db.ai_provider_config.find_one({}, {"_id": 0})
            except Exception:
                provider_config = None
            if not provider_config:
                yield f"data: {json.dumps({'type': 'error', 'error_type': 'no_provider', 'error_message': ERROR_MESSAGES['no_provider']})}\n\n"
                yield "data: [DONE]\n\n"
                return

            provider_name = provider_config.get("provider", "")
            selected_model = req.model or provider_config.get("selected_model", "")
            try:
                api_key = decrypt_key(provider_config["encrypted_api_key"])
            except Exception:
                yield f"data: {json.dumps({'type': 'error', 'error_type': 'invalid_key', 'error_message': ERROR_MESSAGES['invalid_key']})}\n\n"
                yield "data: [DONE]\n\n"
                return

            try:
                provider = get_provider(provider_name)
            except Exception as exc:
                yield f"data: {json.dumps({'type': 'error', 'error_message': f'Unknown provider {provider_name}: {exc}'})}\n\n"
                yield "data: [DONE]\n\n"
                return

            try:
                async for event in agent_run(
                    task=req.task,
                    provider=provider,
                    api_key=api_key,
                    model=selected_model,
                    db=db,
                    in_memory_data=in_memory_data,
                    max_steps=int(req.max_steps or 80),
                    auto_approve_destructive=bool(req.auto_approve_destructive),
                    session_id=req.session_id,
                ):
                    yield f"data: {json.dumps(event, default=str)}\n\n"
            except AIError as ae:
                err_type = getattr(ae, "error_type", "network")
                msg_template = ERROR_MESSAGES.get(err_type, str(ae))
                msg = msg_template.format(provider=provider_name, model=selected_model)
                yield f"data: {json.dumps({'type':'error','error_type':err_type,'error_message':msg})}\n\n"
            except Exception as exc:
                logger.exception("Agent run failed")
                yield f"data: {json.dumps({'type':'error','error_message':str(exc)})}\n\n"
        finally:
            yield "data: [DONE]\n\n"

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@api_router.post("/agent/runs/{run_id}/approve")
async def agent_approve(run_id: str, call_id: str, req: AgentApprovalRequest):
    ok = agent_submit_approval(run_id, call_id, req.decision)
    if not ok:
        raise HTTPException(status_code=404, detail="No pending approval for this call_id")
    return {"ok": True, "run_id": run_id, "call_id": call_id, "decision": req.decision}


@api_router.post("/agent/runs/{run_id}/cancel")
async def agent_cancel(run_id: str):
    ok = agent_cancel_run(run_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Run not active")
    return {"ok": True, "run_id": run_id}


@api_router.get("/agent/runs")
async def agent_list_runs(limit: int = 25):
    try:
        runs = await db.agent_runs.find({}, {"_id": 0, "history": 0}) \
                                  .sort("started_at", -1).to_list(limit)
        if runs:
            return {"runs": runs}
    except Exception:
        pass
    runs = list(in_memory_data.get("agent_runs") or [])
    runs.sort(key=lambda r: r.get("started_at", ""), reverse=True)
    return {"runs": [{k: v for k, v in r.items() if k != "history"} for r in runs[:limit]]}


@api_router.get("/agent/runs/{run_id}")
async def agent_get_run(run_id: str):
    try:
        run = await db.agent_runs.find_one({"run_id": run_id}, {"_id": 0})
        if run:
            return run
    except Exception:
        pass
    for r in (in_memory_data.get("agent_runs") or []):
        if r.get("run_id") == run_id:
            return r
    raise HTTPException(status_code=404, detail="Run not found")


@api_router.get("/agent/destructive-tools")
async def agent_destructive_tools():
    return {"destructive_tools": sorted(AGENT_DESTRUCTIVE_TOOLS)}


@api_router.post("/agent/sessions/{session_id}/reset")
async def agent_reset_session(session_id: str):
    """Drop the agent's persisted conversation memory for a chat session.
    The next agent run for this session_id starts from a clean slate."""
    try:
        from backend.agent import reset_session_history
    except Exception:
        try:
            from .agent import reset_session_history  # type: ignore
        except Exception:
            from agent import reset_session_history  # type: ignore
    cleared = await reset_session_history(session_id, db=db)
    return {"ok": True, "session_id": session_id, "cleared": cleared}


# ── Excel workbook upload for the agent's model-import workflow ─────────────

def _agent_workbook_module():
    try:
        from backend.agent import workbook as _wb
    except Exception:
        try:
            from .agent import workbook as _wb  # type: ignore
        except Exception:
            from agent import workbook as _wb  # type: ignore
    return _wb


@api_router.post("/agent/workbooks/upload")
async def agent_upload_workbook(file: UploadFile = File(...)):
    """Upload an .xlsx model workbook for the agent to analyse and translate
    into DSL rules. The file is stored on disk; the agent inspects it via the
    list_workbooks / get_workbook_overview / get_sheet_formulas tools."""
    wb = _agent_workbook_module()
    content = await file.read()
    try:
        meta = wb.save_workbook_bytes(file.filename or "workbook.xlsx", content)
    except wb.WorkbookError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    if meta.get("duplicate_of_existing"):
        message = (
            f"This file is already uploaded as '{meta['filename']}' "
            f"(workbook_id {meta['workbook_id']}) — reusing the existing "
            f"copy instead of duplicating it."
        )
    else:
        message = (
            f"Workbook '{meta['filename']}' uploaded "
            f"({len(meta['sheets'])} sheets). Ask the agent to analyse it — "
            f"it will confirm which sheets are inputs / calculations / "
            f"outputs, translate the formulas into rules, and reconcile the "
            f"results against the workbook's numbers."
        )
    return {**meta, "message": message}


@api_router.get("/agent/workbooks")
async def agent_list_workbooks():
    wb = _agent_workbook_module()
    items = wb.list_workbooks()
    return {"workbooks": items, "count": len(items)}


@api_router.delete("/agent/workbooks/{workbook_id}")
async def agent_delete_workbook(workbook_id: str):
    wb = _agent_workbook_module()
    try:
        return wb.delete_workbook(workbook_id)
    except wb.WorkbookError as exc:
        raise HTTPException(status_code=404, detail=str(exc))


# ── Requirement documents (PDF / Word) the agent reads & builds from ────────

def _agent_document_module():
    try:
        from backend.agent import requirements_doc as _rd
    except Exception:
        try:
            from .agent import requirements_doc as _rd  # type: ignore
        except Exception:
            from agent import requirements_doc as _rd  # type: ignore
    return _rd


@api_router.post("/agent/documents/upload")
async def agent_upload_document(file: UploadFile = File(...)):
    """Upload a business-requirements document (PDF or Word .docx). The server
    extracts its text; the agent reads it via the list_requirement_documents /
    read_requirement_document tools, analyses it, asks clarifying questions,
    then authors the rules."""
    rd = _agent_document_module()
    content = await file.read()
    try:
        meta = rd.save_document_bytes(file.filename or "document", content)
    except rd.DocumentError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    if meta.get("duplicate_of_existing"):
        message = (
            f"This document is already uploaded as '{meta['filename']}' "
            f"(document_id {meta['document_id']}) — reusing the existing copy."
        )
    else:
        unit = (f"{meta.get('pages')} pages" if meta.get("pages")
                else f"{meta.get('paragraphs', 0)} paragraphs")
        message = (
            f"Requirements document '{meta['filename']}' uploaded ({unit}). "
            f"Ask the agent to read it — it will summarise what it understands, "
            f"confirm the details with you, then build the rules."
        )
    return {**meta, "message": message}


@api_router.get("/agent/documents")
async def agent_list_documents():
    rd = _agent_document_module()
    items = rd.list_documents()
    return {"documents": items, "count": len(items)}


@api_router.delete("/agent/documents/{document_id}")
async def agent_delete_document(document_id: str):
    rd = _agent_document_module()
    try:
        return rd.delete_document(document_id)
    except rd.DocumentError as exc:
        raise HTTPException(status_code=404, detail=str(exc))


# ── Maker-checker: human review queue for agent-authored rules ──────────────

@api_router.get("/agent/approvals")
async def agent_list_approvals():
    """List rules awaiting human approval (maker-checker). Each entry carries the
    agent's change_summary and submission time so a reviewer can triage."""
    try:
        rules = await db.saved_rules.find(
            {"approval_status": "pending"},
            {"_id": 0, "id": 1, "name": 1, "priority": 1, "approval": 1,
             "updated_at": 1},
        ).sort("updated_at", -1).to_list(200)
    except Exception:
        rules = []
    return {"pending": rules, "count": len(rules),
            "enforced": settings.require_agent_approval}


async def _decide_rule_approval(rule_id: str, req: AgentRuleApprovalRequest,
                                decision: str) -> dict:
    checker = (req.checker or "").strip()
    if not checker:
        raise HTTPException(status_code=400,
                            detail="`checker` (reviewer identity) is required.")
    rule = await db.saved_rules.find_one(
        {"id": rule_id}, {"_id": 0, "approval": 1, "approval_status": 1, "name": 1}
    )
    if not rule:
        raise HTTPException(status_code=404, detail="Rule not found.")
    maker = ((rule.get("approval") or {}).get("maker") or "").strip()
    if checker.lower() == maker.lower() and maker:
        raise HTTPException(
            status_code=400,
            detail="Checker must differ from maker (segregation of duties).",
        )
    now = datetime.now(timezone.utc).isoformat()
    await db.saved_rules.update_one({"id": rule_id}, {"$set": {
        "approval_status": decision,
        "approval.checker": checker,
        "approval.decided_at": now,
        "approval.note": (req.note or "").strip(),
    }})
    return {"ok": True, "rule_id": rule_id, "name": rule.get("name"),
            "approval_status": decision, "checker": checker, "decided_at": now}


@api_router.post("/agent/approvals/{rule_id}/approve")
async def agent_approve_rule(rule_id: str, req: AgentRuleApprovalRequest):
    """Approve an agent-authored rule so its template can be deployed."""
    return await _decide_rule_approval(rule_id, req, "approved")


@api_router.post("/agent/approvals/{rule_id}/reject")
async def agent_reject_rule(rule_id: str, req: AgentRuleApprovalRequest):
    """Reject an agent-authored rule (records the reviewer + reason)."""
    return await _decide_rule_approval(rule_id, req, "rejected")


@api_router.post("/import/transactions")
async def import_transactions(file: UploadFile = File(...)):
    """Load transaction definitions from a JSON array.

    Each entry is shaped like:
        { "name": "PAYMENT_UPB", "exclusive": 1, "isGL": 1, "isReplayable": 0 }

    Replaces the entire `transaction_definitions` collection.
    """
    if not file.filename or not file.filename.lower().endswith(".json"):
        raise HTTPException(status_code=400, detail="Only .json files are accepted.")
    try:
        content = await file.read()
        records = json.loads(content.decode("utf-8"))
    except Exception:
        raise HTTPException(status_code=400, detail="The file is not valid JSON.")
    if not isinstance(records, list):
        raise HTTPException(status_code=422, detail="File must contain a JSON array.")

    docs = []
    for i, item in enumerate(records):
        if not isinstance(item, dict):
            raise HTTPException(status_code=422, detail=f"Item at index {i} is not an object.")
        name = (item.get("name") or item.get("transactiontype") or "").strip()
        if not name:
            raise HTTPException(status_code=422, detail=f"Item at index {i} is missing 'name'.")
        docs.append({
            "transactiontype": name,
            "exclusive": item.get("exclusive", 1),
            "isGL": item.get("isGL", 1),
            "isReplayable": item.get("isReplayable", 0),
        })

    try:
        await db.transaction_definitions.delete_many({})
        if docs:
            await db.transaction_definitions.insert_many([dict(d) for d in docs])
    except Exception as e:
        logger.error(f"Failed to persist transaction definitions: {e}")
        raise HTTPException(status_code=500, detail=f"Could not save transaction definitions: {e}")

    return {
        "success": True,
        "count": len(docs),
        "transaction_types": [d["transactiontype"] for d in docs],
    }


# ---------------------------------------------------------------------------
# Event Configuration → Event Definition transformer
# ---------------------------------------------------------------------------

# Keywords that mark a field name as numeric (decimal). Lookup is case-
# insensitive and substring-based on the assembled UPPERCASE field name.
_DECIMAL_KEYWORDS = (
    "AMOUNT", "BALANCE", "PRINCIPAL", "RATE", "INTEREST", "PRICE",
    "QUANTITY", "TERM", "COUPON", "YIELD", "SPREAD", "FEE", "PAYMENT",
    "CASH", "ACCRUAL", "RECEIVABLE", "CF", "LOAN",
)
_DECIMAL_NAME_SUFFIXES = ("_ID", "ID")  # ProductId / customer_id → decimal


def _infer_field_dt(name: str) -> str:
    """Infer datatype from the assembled (already prefixed) field name."""
    n = name.upper()
    if "DATE" in n:
        return "date"
    if any(k in n for k in _DECIMAL_KEYWORDS):
        return "decimal"
    # Treat *Id / *_ID / *_ID_* as decimal (matches sample reference data:
    # ProductId, ATTRIBUTE_PRODUCT_ID_CURRENT, etc.)
    if n.endswith("ID") or n.endswith("_ID") or "_ID_" in n:
        return "decimal"
    return "string"


# `MeasurementType` is a Fyntrac convention: when an operational-trigger event
# exposes a column literally named `MeasurementType`, the column is also
# lifted into its own reference event (a measurement-type lookup table).
_LIFTED_REFERENCE_FIELD_NAMES = {"measurementtype"}


def _transform_event_configurations(records: list) -> list:
    """Convert an EventConfiguration JSON array into EventDefinition docs.

    Naming rules (validated against the sample reference data in
    `Importflow/`):

    * Reference event (triggerSource includes `reference_table`):
      eventType=reference, eventTable=custom; field name = bare
      `sourceColumns[].value`.
    * Operational-trigger event (triggerSource includes `operational_table`):
      eventType=activity, eventTable=standard; field name = bare
      `sourceColumns[].value`. If a column is `MeasurementType`, also emit
      a separate reference event for it.
    * Otherwise (model-execution / replay / transaction-post triggers):
      - If `versionType` is non-empty: name =
        `{TABLE}_{COLUMN}_{VERSIONTYPE}` for each (column, version) pair.
      - Else if table is `Balances`: emit three phases per dataMapping —
        `BALANCES_BEGINNINGBALANCE_{DM}`,
        `BALANCES_ENDINGBALANCE_{DM}`,
        `BALANCES_ACTIVITY_{DM}` (decimal).
      - Else if `dataMapping` is non-empty: name =
        `{TABLE}_{COLUMN}_{DM}` per dataMapping value. When
        `fieldType=AGGREGATED`, only the first dataMapping is emitted.
      - Else: name = `{TABLE}_{COLUMN}`.
    """
    out: list = []
    ts = datetime.now(timezone.utc).isoformat()
    lifted_refs: dict = {}  # event_name -> list[(field_name, datatype)]

    for cfg in records:
        if not isinstance(cfg, dict):
            continue
        eid = (cfg.get("eventId") or cfg.get("eventName") or "UNKNOWN").strip() or "UNKNOWN"

        trig = cfg.get("triggerSetup") or {}
        trig_sources = [
            ((s.get("value") or "").strip().lower())
            for s in (trig.get("triggerSource") or []) if isinstance(s, dict)
        ]
        is_reference = "reference_table" in trig_sources
        is_operational = "operational_table" in trig_sources

        evt_type = "reference" if is_reference else "activity"
        evt_table = "custom" if is_reference else "standard"

        ordered_fields: list = []
        seen_names: set = set()

        def _add(name: str, dt: str) -> None:
            if not name or name in seen_names:
                return
            seen_names.add(name)
            ordered_fields.append({"name": name, "datatype": dt})

        for sm in (cfg.get("sourceMappings") or []):
            if not isinstance(sm, dict):
                continue
            table_up = ((sm.get("sourceTable") or "").strip()).upper()
            cols = sm.get("sourceColumns") or []
            ver_types = sm.get("versionType") or []
            data_map = sm.get("dataMapping") or []
            field_type = (sm.get("fieldType") or "NONE").upper()

            for col in cols:
                if not isinstance(col, dict):
                    continue
                col_val = (col.get("value") or "").strip()
                if not col_val:
                    continue
                col_up = col_val.upper()

                if is_reference or is_operational:
                    _add(col_val, _infer_field_dt(col_val))
                    if is_operational and col_val.lower() in _LIFTED_REFERENCE_FIELD_NAMES:
                        lifted_refs.setdefault(col_val, []).append((col_val, _infer_field_dt(col_val)))
                    continue

                if ver_types:
                    for vt in ver_types:
                        suf = ((vt.get("value") or "").strip()).upper()
                        name = f"{table_up}_{col_up}_{suf}" if suf else f"{table_up}_{col_up}"
                        _add(name, _infer_field_dt(name))
                    continue

                if data_map:
                    targets = data_map[:1] if field_type == "AGGREGATED" else data_map
                    if table_up == "BALANCES":
                        phases = ("BEGINNINGBALANCE", "ENDINGBALANCE", "ACTIVITY")
                        for dm in targets:
                            dm_v = ((dm.get("value") or "").strip()).upper()
                            for ph in phases:
                                _add(f"{table_up}_{ph}_{dm_v}", "decimal")
                    else:
                        for dm in targets:
                            dm_v = ((dm.get("value") or "").strip()).upper()
                            _add(f"{table_up}_{col_up}_{dm_v}", _infer_field_dt(f"{table_up}_{col_up}_{dm_v}"))
                    continue

                _add(f"{table_up}_{col_up}", _infer_field_dt(f"{table_up}_{col_up}"))

        out.append({
            "id": str(uuid.uuid4()),
            "event_name": eid,
            "fields": ordered_fields,
            "eventType": evt_type,
            "eventTable": evt_table,
            "created_at": ts,
        })

    for name, items in lifted_refs.items():
        seen2: set = set()
        unique: list = []
        for fn, dt in items:
            if fn in seen2:
                continue
            seen2.add(fn)
            unique.append({"name": fn, "datatype": dt})
        out.append({
            "id": str(uuid.uuid4()),
            "event_name": name,
            "fields": unique,
            "eventType": "reference",
            "eventTable": "custom",
            "created_at": ts,
        })

    return out


@api_router.post("/import/event-configurations")
async def import_event_configurations(file: UploadFile = File(...)):
    """Load event definitions from an EventConfiguration JSON array.

    See `_transform_event_configurations` for naming rules. Replaces the
    entire `event_definitions` collection.
    """
    if not file.filename or not file.filename.lower().endswith(".json"):
        raise HTTPException(status_code=400, detail="Only .json files are accepted.")
    try:
        content = await file.read()
        records = json.loads(content.decode("utf-8"))
    except Exception:
        raise HTTPException(status_code=400, detail="The file is not valid JSON.")
    if not isinstance(records, list) or not records:
        raise HTTPException(status_code=422, detail="File must contain a non-empty JSON array.")

    try:
        defs = _transform_event_configurations(records)
        if not defs:
            raise HTTPException(status_code=422, detail="No event definitions could be derived from the file.")
        await db.event_definitions.delete_many({})
        await db.event_definitions.insert_many([dict(d) for d in defs])
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Event configuration import failed: {e}")
        raise HTTPException(status_code=500, detail=f"Could not import event configurations: {e}")

    return {
        "success": True,
        "count": len(defs),
        "names": [d["event_name"] for d in defs],
        "types": {d["event_name"]: d["eventTable"] for d in defs},
    }


# ── Saved Rules CRUD ────────────────────────────────────────────────────

@api_router.get("/saved-rules")
async def list_saved_rules(summary: int = 0):
    """List all saved rule builder configurations.
    Pass ?summary=1 to exclude generatedCode (fast list for UI display).
    """
    try:
        projection = {"_id": 0}
        if summary:
            projection["generatedCode"] = 0
        rules = await db.saved_rules.find({}, projection).sort("updated_at", -1).to_list(500)
        return rules
    except Exception as e:
        logger.error(f"Error listing saved rules: {e}")
        return []

@api_router.post("/saved-rules")
async def save_rule(request: dict):
    """Save or update a rule builder configuration. Rule name must be unique."""
    name = (request.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Rule name is required.")

    rule_id = request.get("id")
    now = datetime.now(timezone.utc).isoformat()
    existing_doc = None
    if rule_id:
        existing_doc = await db.saved_rules.find_one({"id": rule_id}, {"_id": 0})

    # Check uniqueness: no other rule with same name (case-insensitive)
    existing = await db.saved_rules.find_one(
        {"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}},
        {"_id": 0, "id": 1},
    )
    if existing and (not rule_id or existing["id"] != rule_id):
        raise HTTPException(
            status_code=409,
            detail=f"A rule named \"{name}\" already exists. Please choose a different name.",
        )

    # Priority uniqueness across rules AND schedules
    priority = request.get("priority")
    if priority is not None:
        priority = int(priority)
        # Check other rules
        rule_with_priority = await db.saved_rules.find_one(
            {"priority": priority, **({"id": {"$ne": rule_id}} if rule_id else {})},
            {"_id": 0, "id": 1, "name": 1},
        )
        if rule_with_priority:
            raise HTTPException(
                status_code=409,
                detail=f"Priority {priority} is already used by rule \"{rule_with_priority['name']}\". Please choose a different priority.",
            )
        # Check schedules collection
        sched_with_priority = await db.saved_schedules.find_one(
            {"priority": priority},
            {"_id": 0, "id": 1, "name": 1},
        )
        if sched_with_priority:
            raise HTTPException(
                status_code=409,
                detail=f"Priority {priority} is already used by schedule \"{sched_with_priority['name']}\". Please choose a different priority.",
            )

    doc = {
        "name": name,
        "priority": priority,
        "disabled": bool(request.get("disabled", (existing_doc or {}).get("disabled", False))),
        "ruleType": request.get("ruleType", "simple_calc"),
        "variables": request.get("variables", []),
        "conditions": request.get("conditions", []),
        "elseFormula": request.get("elseFormula", ""),
        "conditionResultVar": request.get("conditionResultVar", "result"),
        "iterations": request.get("iterations", []),
        "iterConfig": request.get("iterConfig", {}),
        "outputs": request.get("outputs", {}),
        "inlineComment": request.get("inlineComment", False),
        "commentText": request.get("commentText", ""),
        "customCode": request.get("customCode", ""),
        "generatedCode": request.get("generatedCode", ""),
        "steps": request.get("steps", []),
        "updated_at": now,
    }

    # Backfill missing postingDate/effectiveDate/subInstrumentId on every
    # transaction entry, mirroring the agent-side normaliser, so that UI
    # saves don't bypass the date-defaulting logic.
    try:
        from backend.agent.tools import _normalise_transaction_outputs
        _normalise_transaction_outputs(doc.get("steps") or [], doc.get("outputs") or {})
    except Exception as _norm_err:
        logger.warning(f"transaction normalisation skipped: {_norm_err}")

    # Always regenerate generatedCode from the current persisted shape so UI
    # preview/runtime combined code reflect disabled step/transaction comments
    # even if the client sends stale generatedCode.
    try:
        from backend.agent.tools import _generate_rule_code
        doc["generatedCode"] = _generate_rule_code(doc)
    except Exception as _gen_err:
        logger.warning(f"generatedCode regeneration skipped: {_gen_err}")

    if rule_id:
        doc["id"] = rule_id
        await db.saved_rules.replace_one({"id": rule_id}, doc, upsert=True)
    else:
        doc["id"] = str(uuid.uuid4())
        doc["created_at"] = now
        await db.saved_rules.insert_one(doc)

    return {"success": True, "id": doc["id"], "message": f"Rule \"{name}\" saved."}

# NOTE: Static routes (/saved-rules/reorder) MUST be declared BEFORE the
# parameterized routes (/saved-rules/{rule_id}) — otherwise FastAPI matches
# "reorder" as a {rule_id} path parameter.
@api_router.put("/saved-rules/reorder")
async def reorder_saved_rules(request: dict):
    """Batch-update priorities for saved rules based on drag-and-drop ordering.
    Expects: { "order": [ { "id": "...", "priority": 1 }, ... ] }
    """
    order = request.get("order", [])
    if not order:
        raise HTTPException(status_code=400, detail="No ordering provided.")
    try:
        for item in order:
            rule_id = item.get("id")
            priority = item.get("priority")
            if rule_id is not None and priority is not None:
                await db.saved_rules.update_one(
                    {"id": rule_id},
                    {"$set": {"priority": int(priority)}},
                )
        return {"success": True, "message": f"Updated priorities for {len(order)} rules."}
    except Exception as e:
        logger.error(f"Error reordering rules: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@api_router.delete("/saved-rules/{rule_id}")
async def delete_saved_rule(rule_id: str):
    """Delete a saved rule by its id."""
    result = await db.saved_rules.delete_one({"id": rule_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Rule not found.")
    return {"success": True, "message": "Rule deleted."}

@api_router.post("/saved-rules/{rule_id}/revert")
async def revert_saved_rule(rule_id: str):
    """Restore the most recent pre-save snapshot for a rule.

    Every time the agent saves a rule (update_step, patch_step, update_saved_rule,
    etc.) the previous version is saved to the rule_history collection.
    This endpoint pops the newest snapshot and restores it so the user can
    undo the last agent change.
    """
    # Find the most recent snapshot for this rule.
    snap = await db.rule_history.find_one(
        {"rule_id": rule_id},
        sort=[("snapshot_at", -1)],
    )
    if not snap:
        raise HTTPException(status_code=404, detail="No history found for this rule.")
    rule_doc = snap.get("rule_doc")
    if not rule_doc:
        raise HTTPException(status_code=500, detail="Snapshot has no rule document.")
    rule_doc.pop("_id", None)
    # Before restoring, snapshot the CURRENT state so the revert itself is
    # also reversible (undo-undo).
    current = await db.saved_rules.find_one({"id": rule_id}, {"_id": 0})
    if current:
        current.pop("_id", None)
        await db.rule_history.insert_one({
            "rule_id": rule_id,
            "snapshot_at": datetime.now(timezone.utc).isoformat(),
            "rule_doc": current,
        })
    # Restore the snapshot.
    await db.saved_rules.replace_one({"id": rule_id}, rule_doc, upsert=True)
    # Remove the snapshot we just restored (it's now the live doc).
    await db.rule_history.delete_one({"_id": snap["_id"]})
    # Prune: keep only the 20 most recent snapshots for this rule.
    try:
        all_snaps = await db.rule_history.find(
            {"rule_id": rule_id}, {"_id": 1}
        ).sort("snapshot_at", -1).to_list(None)
        if len(all_snaps) > 20:
            ids_to_delete = [s["_id"] for s in all_snaps[20:]]
            await db.rule_history.delete_many({"_id": {"$in": ids_to_delete}})
    except Exception:
        pass
    rule_doc.pop("_id", None)
    return {
        "success": True,
        "message": f"Rule '{rule_doc.get('name', rule_id)}' reverted to snapshot from {snap.get('snapshot_at', '?')}",
        "rule": rule_doc,
    }

@api_router.get("/saved-rules/{rule_id}/history")
async def get_rule_history(rule_id: str):
    """Return the list of available revert snapshots for a rule (newest first)."""
    snaps = await db.rule_history.find(
        {"rule_id": rule_id},
        {"_id": 0, "rule_doc": 0},
        sort=[("snapshot_at", -1)],
    ).to_list(20)
    return {"rule_id": rule_id, "snapshots": snaps}

@api_router.get("/saved-rules/{rule_id}")
async def get_saved_rule(rule_id: str):
    """Fetch a single saved rule by id (used by the Rule Builder to refresh
    after agent-driven mutations)."""
    doc = await db.saved_rules.find_one({"id": rule_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Rule not found.")
    return doc

@api_router.put("/saved-rules/{rule_id}")
async def update_saved_rule(rule_id: str, request: dict):
    """Patch specific fields of a saved rule (generatedCode, outputs, steps, etc.)."""
    allowed = {"generatedCode", "outputs", "steps", "name", "priority", "variables",
               "conditions", "elseFormula", "conditionResultVar", "iterations",
               "iterConfig", "inlineComment", "commentText", "ruleType", "disabled"}
    update_fields = {k: v for k, v in request.items() if k in allowed}
    if not update_fields:
        raise HTTPException(status_code=400, detail="No valid fields to update.")
    existing = await db.saved_rules.find_one({"id": rule_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Rule not found.")

    merged = {**existing, **update_fields}

    # If caller patches steps/outputs/name/etc, regenerate generatedCode so
    # disabled step/transaction comments stay in sync.
    if any(k in update_fields for k in {"steps", "outputs", "name"}):
        try:
            from backend.agent.tools import _normalise_transaction_outputs, _generate_rule_code
            _normalise_transaction_outputs(merged.get("steps") or [], merged.get("outputs") or {})
            merged["generatedCode"] = _generate_rule_code(merged)
        except Exception as _gen_err:
            logger.warning(f"update_saved_rule regeneration skipped: {_gen_err}")

    result = await db.saved_rules.update_one({"id": rule_id}, {"$set": merged})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Rule not found.")
    return {"success": True, "message": "Rule updated."}

@api_router.delete("/saved-rules")
async def delete_all_saved_rules():
    """Delete ALL saved rules."""
    result = await db.saved_rules.delete_many({})
    return {"success": True, "deleted": result.deleted_count, "message": f"Deleted {result.deleted_count} rule(s)."}


@api_router.put("/saved-schedules/reorder")
async def reorder_saved_schedules(request: dict):
    """Batch-update priorities for saved schedules based on drag-and-drop ordering.
    Expects: { "order": [ { "id": "...", "priority": 1 }, ... ] }
    """
    order = request.get("order", [])
    if not order:
        raise HTTPException(status_code=400, detail="No ordering provided.")
    try:
        for item in order:
            sched_id = item.get("id")
            priority = item.get("priority")
            if sched_id is not None and priority is not None:
                await db.saved_schedules.update_one(
                    {"id": sched_id},
                    {"$set": {"priority": int(priority)}},
                )
        return {"success": True, "message": f"Updated priorities for {len(order)} schedules."}
    except Exception as e:
        logger.error(f"Error reordering schedules: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ── User Templates CRUD ─────────────────────────────────────────────────

async def _mirror_user_template_to_dsl(name: str, combined_code: str, rules: list) -> None:
    """Mirror a user_template into dsl_templates (upsert by name) and append a
    versioned artifact in dsl_template_artifacts.

    `combined_code` from the rule builder is *DSL* source (not Python). To make
    the artifact directly executable by FyntracPythonModel.ModelRunner, we
    compile it through the same path the playground uses
    (`dsl_to_python_multi_event` / `dsl_to_python_standalone`). The compiled
    Python defines `process_event_data(...)` (or `process_standalone(...)`),
    which is the entry point ModelRunner expects.
    """
    if not name:
        return
    try:
        now_iso = datetime.now(timezone.utc).isoformat()

        # Compile DSL → Python so the artifact is runnable by ModelRunner.
        python_code = ""
        compile_error = None
        try:
            referenced_events = extract_event_names_from_dsl(combined_code or "")
            all_event_fields: Dict[str, Dict[str, Any]] = {}
            for evt_name in referenced_events:
                evt = await db.event_definitions.find_one(
                    {"event_name": evt_name}, {"_id": 0}
                )
                if evt:
                    all_event_fields[evt_name] = {
                        "fields": evt.get("fields", []),
                        "eventType": evt.get("eventType", "activity"),
                    }
            if all_event_fields:
                python_code = dsl_to_python_multi_event(
                    combined_code or "", all_event_fields
                )
            else:
                python_code = dsl_to_python_standalone(combined_code or "")
        except Exception as e:
            compile_error = str(e)
            logger.warning(
                f"DSL→Python compile failed for user template '{name}': {e}"
            )

        # Upsert dsl_templates by name. Preserve existing id when updating so
        # the artifact's template_id remains stable across edits.
        existing = await db.dsl_templates.find_one({"name": name}, {"_id": 0, "id": 1})
        template_id = (existing or {}).get("id") or str(uuid.uuid4())
        dsl_doc = {
            "id": template_id,
            "name": name,
            "dsl_code": combined_code or "",
            "python_code": python_code,
            "updated_at": now_iso,
        }
        if not existing:
            dsl_doc["created_at"] = now_iso
        await db.dsl_templates.update_one(
            {"name": name},
            {"$set": dsl_doc, "$setOnInsert": {"source": "user_template"}},
            upsert=True,
        )

        # Determine next artifact version for this template.
        latest = await db.dsl_template_artifacts.find_one(
            {"template_id": template_id},
            {"_id": 0, "version": 1},
            sort=[("version", -1)],
        )
        next_version = 1
        if latest and isinstance(latest.get("version"), int):
            next_version = latest["version"] + 1

        artifact_doc = {
            "id": str(uuid.uuid4()),
            "template_id": template_id,
            "template_name": name,
            "version": next_version,
            "python_code": python_code,
            "rules_count": len(rules or []),
            "created_at": now_iso,
            "read_only": True,
        }
        if compile_error:
            artifact_doc["compile_error"] = compile_error
        await db.dsl_template_artifacts.insert_one(artifact_doc)

        keep_versions = os.environ.get(
            "KEEP_TEMPLATE_ARTIFACT_VERSIONS", "false"
        ).lower() in ("1", "true", "yes")
        if not keep_versions:
            await db.dsl_template_artifacts.delete_many(
                {"template_id": template_id, "version": {"$lt": next_version}}
            )
    except Exception as e:
        logger.warning(f"Failed to mirror user template '{name}' to dsl_templates: {e}")


@api_router.get("/user-templates")
async def list_user_templates():
    """List all user-created templates."""
    try:
        templates = await db.user_templates.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)
        return templates
    except Exception as e:
        logger.error(f"Error listing user templates: {e}")
        return []

@api_router.post("/user-templates")
async def save_user_template(request: dict):
    """Create a user template from saved rules."""
    name = (request.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Template name is required.")
    description = (request.get("description") or "").strip()
    category = (request.get("category") or "User Created").strip()
    rules = request.get("rules", [])
    schedules = request.get("schedules", [])
    combined_code = request.get("combinedCode", "")

    # Check name uniqueness
    existing = await db.user_templates.find_one(
        {"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}},
        {"_id": 0, "id": 1},
    )
    if existing:
        raise HTTPException(status_code=409, detail=f"A template named \"{name}\" already exists.")

    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()),
        "name": name,
        "description": description,
        "category": category,
        "rules": rules,
        "schedules": schedules,
        "combinedCode": combined_code,
        "created_at": now,
        "updated_at": now,
    }
    await db.user_templates.insert_one(doc)
    # Note: dsl_templates / dsl_template_artifacts are populated only when the
    # user explicitly clicks Deploy (POST /user-templates/{id}/deploy).
    return {"success": True, "id": doc["id"], "message": f"Template \"{name}\" created."}

@api_router.delete("/user-templates/{template_id}")
async def delete_user_template(template_id: str):
    """Delete a user template by id."""
    existing = await db.user_templates.find_one({"id": template_id}, {"_id": 0, "name": 1})
    result = await db.user_templates.delete_one({"id": template_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Template not found.")
    # Cascade delete the mirrored dsl_template + its artifacts.
    if existing and existing.get("name"):
        try:
            mirrored = await db.dsl_templates.find_one(
                {"name": existing["name"]}, {"_id": 0, "id": 1}
            )
            await db.dsl_templates.delete_one({"name": existing["name"]})
            if mirrored and mirrored.get("id"):
                await db.dsl_template_artifacts.delete_many(
                    {"template_id": mirrored["id"]}
                )
        except Exception as e:
            logger.warning(
                f"Failed to cascade-delete dsl_template mirror for '{existing['name']}': {e}"
            )
    return {"success": True, "message": "Template deleted."}

@api_router.put("/user-templates/{template_id}")
async def update_user_template(template_id: str, request: dict):
    """Overwrite an existing user template's rules and code (keeps name/description/category)."""
    existing = await db.user_templates.find_one({"id": template_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Template not found.")
    now = datetime.now(timezone.utc).isoformat()
    update_fields = {"updated_at": now}
    if "rules" in request:
        update_fields["rules"] = request["rules"]
    if "schedules" in request:
        update_fields["schedules"] = request["schedules"]
    if "combinedCode" in request:
        update_fields["combinedCode"] = request["combinedCode"]
    # Allow optional metadata updates
    if "description" in request:
        update_fields["description"] = request["description"]
    if "category" in request:
        update_fields["category"] = request["category"]
    await db.user_templates.update_one({"id": template_id}, {"$set": update_fields})
    # Note: dsl_templates / dsl_template_artifacts are NOT auto-updated here.
    # The user must click Deploy (POST /user-templates/{id}/deploy) to push
    # changes to the runtime.
    return {"success": True, "id": template_id, "message": f"Template \"{existing['name']}\" updated."}


@api_router.post("/user-templates/{template_id}/deploy")
async def deploy_user_template(template_id: str):
    """
    Deploy a single user template to the runtime.

    Compiles the user_template's DSL using the *current* event_definitions
    schema and writes:
      - dsl_templates: one document keyed by template name (DSL + ready-to-run
        Python in `python_code`).
      - dsl_template_artifacts: a new versioned snapshot scoped to this
        template's id (older versions of the same template are pruned unless
        KEEP_TEMPLATE_ARTIFACT_VERSIONS=true).

    Only the named template's documents are touched; all other templates and
    their artifacts are left untouched.
    """
    existing = await db.user_templates.find_one({"id": template_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Template not found.")

    name = existing.get("name") or ""
    if not name:
        raise HTTPException(status_code=400, detail="Template has no name; cannot deploy.")

    # Maker-checker gate: refuse to deploy if any rule in the template is still
    # awaiting (or was denied) human approval. Only enforced when enabled.
    if settings.require_agent_approval:
        try:
            from backend.agent.tools import unapproved_rules_in_template
        except Exception:
            try:
                from agent.tools import unapproved_rules_in_template  # type: ignore
            except Exception:
                from .agent.tools import unapproved_rules_in_template  # type: ignore
        blocked = await unapproved_rules_in_template(db, existing)
        if blocked:
            listing = ", ".join(
                f"{b.get('name')} ({b.get('approval_status')})" for b in blocked
            )
            raise HTTPException(
                status_code=409,
                detail=(
                    f"Cannot deploy \"{name}\": {len(blocked)} rule(s) await "
                    f"approval: {listing}. A reviewer must approve each via "
                    f"POST /api/agent/approvals/{{rule_id}}/approve before deploy."
                ),
            )

    combined_code = existing.get("combinedCode") or ""
    rules = existing.get("rules") or []

    await _mirror_user_template_to_dsl(name, combined_code, rules)

    # Read back the freshly-written rows so the caller gets confirmation of
    # what the runtime will see.
    dsl_doc = await db.dsl_templates.find_one(
        {"name": name},
        {"_id": 0, "id": 1, "updated_at": 1},
    ) or {}
    artifact = await db.dsl_template_artifacts.find_one(
        {"template_id": dsl_doc.get("id")},
        {"_id": 0, "id": 1, "version": 1, "created_at": 1, "compile_error": 1},
        sort=[("version", -1)],
    ) or {}

    return {
        "success": True,
        "message": f"Template \"{name}\" deployed.",
        "template": {"id": dsl_doc.get("id"), "name": name, "updated_at": dsl_doc.get("updated_at")},
        "artifact": artifact,
    }


# ── Template Sample Data ────────────────────────────────────────────────

@api_router.post("/template-sample-data/{template_id}")
async def load_template_sample_data(template_id: str):
    """Load pre-defined sample event definitions and event data for a specific template."""
    import importlib, sys, os
    backend_dir = os.path.dirname(os.path.abspath(__file__))
    if backend_dir not in sys.path:
        sys.path.insert(0, backend_dir)
    from template_sample_data import TEMPLATE_SAMPLE_DATA

    if template_id not in TEMPLATE_SAMPLE_DATA:
        raise HTTPException(status_code=404, detail=f"No sample data available for template '{template_id}'")

    sample = TEMPLATE_SAMPLE_DATA[template_id]

    for evt in sample["events"]:
        existing = await db.event_definitions.find_one({"event_name": evt["event_name"]})
        if not existing:
            doc = {
                "id": str(uuid.uuid4()),
                "event_name": evt["event_name"],
                "fields": evt["fields"],
                "eventType": evt.get("eventType", "activity"),
                "eventTable": evt.get("eventTable", "standard"),
                "created_at": datetime.utcnow().isoformat(),
            }
            await db.event_definitions.insert_one(doc)

    for ed in sample["event_data"]:
        await db.event_data.delete_many({"event_name": ed["event_name"]})
        doc = {
            "event_name": ed["event_name"],
            "data_rows": ed["data_rows"],
            "created_at": datetime.utcnow().isoformat(),
        }
        await db.event_data.insert_one(doc)

    # Populate transaction definitions if provided in sample data
    if sample.get("transaction_types"):
        await db.transaction_definitions.delete_many({})
        for txn_type in sample["transaction_types"]:
            await db.transaction_definitions.insert_one({"transactiontype": txn_type})

    events = await db.event_definitions.find({}, {"_id": 0}).to_list(1000)
    return {"success": True, "events": events}

# ── Saved Schedules CRUD ────────────────────────────────────────────────

@api_router.get("/saved-schedules")
async def list_saved_schedules(summary: int = 0):
    """List all saved schedule builder configurations.
    Pass ?summary=1 to exclude generatedCode (fast list for UI display).
    """
    try:
        projection = {"_id": 0}
        if summary:
            projection["generatedCode"] = 0
        schedules = await db.saved_schedules.find({}, projection).sort("updated_at", -1).to_list(500)
        return schedules
    except Exception as e:
        logger.error(f"Error listing saved schedules: {e}")
        return []

@api_router.post("/saved-schedules")
async def save_schedule(request: dict):
    """Save or update a schedule builder configuration."""
    name = (request.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Schedule name is required.")

    schedule_id = request.get("id")
    now = datetime.now(timezone.utc).isoformat()

    # Check uniqueness: no other schedule with same name (case-insensitive)
    existing = await db.saved_schedules.find_one(
        {"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}},
        {"_id": 0, "id": 1},
    )
    if existing and (not schedule_id or existing["id"] != schedule_id):
        raise HTTPException(
            status_code=409,
            detail=f"A schedule named \"{name}\" already exists. Please choose a different name.",
        )

    # Priority uniqueness across rules AND schedules
    priority = request.get("priority")
    if priority is not None:
        priority = int(priority)
        # Check rules collection
        rule_with_priority = await db.saved_rules.find_one(
            {"priority": priority, **({"id": {"$ne": schedule_id}} if schedule_id else {})},
            {"_id": 0, "id": 1, "name": 1},
        )
        if rule_with_priority:
            raise HTTPException(
                status_code=409,
                detail=f"Priority {priority} is already used by rule \"{rule_with_priority['name']}\". Please choose a different priority.",
            )
        # Check schedules collection
        sched_with_priority = await db.saved_schedules.find_one(
            {"priority": priority, **({"id": {"$ne": schedule_id}} if schedule_id else {})},
            {"_id": 0, "id": 1, "name": 1},
        )
        if sched_with_priority:
            raise HTTPException(
                status_code=409,
                detail=f"Priority {priority} is already used by schedule \"{sched_with_priority['name']}\". Please choose a different priority.",
            )

    doc = {
        "name": name,
        "priority": priority,
        "generatedCode": request.get("generatedCode", ""),
        "config": request.get("config", {}),
        "updated_at": now,
    }

    if schedule_id:
        doc["id"] = schedule_id
        await db.saved_schedules.replace_one({"id": schedule_id}, doc, upsert=True)
    else:
        doc["id"] = str(uuid.uuid4())
        doc["created_at"] = now
        await db.saved_schedules.insert_one(doc)

    return {"success": True, "id": doc["id"], "message": f"Schedule \"{name}\" saved."}

@api_router.delete("/saved-schedules/{schedule_id}")
async def delete_saved_schedule(schedule_id: str):
    """Delete a saved schedule by its id."""
    result = await db.saved_schedules.delete_one({"id": schedule_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Schedule not found.")
    return {"success": True, "message": "Schedule deleted."}

@api_router.delete("/saved-schedules")
async def delete_all_saved_schedules():
    """Delete ALL saved schedules."""
    result = await db.saved_schedules.delete_many({})
    return {"success": True, "deleted": result.deleted_count, "message": f"Deleted {result.deleted_count} schedule(s)."}


# ── Combined code endpoint (rules + schedules ordered by priority) ──────

@api_router.get("/combined-code")
async def get_combined_code():
    """Return generated code from all saved rules and schedules, ordered by priority (ascending).

    The 'Dependencies from saved rules' section inside each rule's generatedCode
    re-emits variables that were already defined (and correctly ordered) by earlier
    rules.  When the combined code is executed, those re-emissions overwrite the
    correct values with potentially wrong-ordered ones (e.g. totalssp used before
    it is computed).  To prevent this we track every variable name that has already
    been assigned and strip any re-assignment from later rules' dependency sections.
    """
    import re as _re
    try:
        rules = await db.saved_rules.find({}, {"_id": 0}).to_list(500)
        schedules = await db.saved_schedules.find({}, {"_id": 0}).to_list(500)

        from backend.agent.tools import _generate_rule_code as _gen_code

        items = []
        for r in rules:
            p = r.get("priority")
            try:
                code = _gen_code(r)
            except Exception:
                code = r.get("generatedCode", "")
            items.append({
                "priority": p if p is not None else float('inf'),
                "code": code,
                "name": r.get("name", ""),
                "disabled": bool(r.get("disabled", False)),
            })
        for s in schedules:
            p = s.get("priority")
            try:
                code = _gen_code(s)
            except Exception:
                code = s.get("generatedCode", "")
            items.append({"priority": p if p is not None else float('inf'), "code": code, "name": s.get("name", ""), "disabled": False})

        items.sort(key=lambda x: (x["priority"], x["name"]))

        _assign_re = _re.compile(r'^([A-Za-z_][A-Za-z0-9_]*)\s*=(?!=)')

        def strip_dependencies_section(code: str) -> str:
            """Remove lines between '## Dependencies from saved rules' and the next '##' heading."""
            out = []
            in_deps = False
            for line in code.split('\n'):
                stripped = line.strip()
                if stripped == '## Dependencies from saved rules':
                    in_deps = True
                    out.append(line)
                    continue
                if in_deps:
                    if stripped.startswith('## ') and not stripped.startswith('## ═'):
                        in_deps = False
                        out.append(line)
                    continue
                out.append(line)
            return '\n'.join(out)

        # Strip deps sections from EVERY rule first. The deps section is a
        # snapshot meant for standalone execution; in combined execution every
        # rule runs anyway, and keeping deps creates phantom cross-rule
        # dependencies (e.g. Stage 1's deps referring to a Schedule defined by
        # Stage 2) that lead to NameError at runtime.
        for it in items:
            it['code'] = strip_dependencies_section(it.get('code', ''))

        # If a rule is disabled, keep it visible in combined/runtime code but
        # comment out every non-empty line so it never executes.
        for it in items:
            if not it.get('disabled'):
                continue
            code = it.get('code', '') or ''
            it['code'] = '\n'.join(
                (f"# [DISABLED RULE] {line}" if line.strip() else line)
                for line in code.split('\n')
            )

        # ── Topological reorder: if rule A's body references a symbol that
        # rule B defines, B must come before A — even if A has a lower priority
        # number. This prevents "cannot access local variable 'X'" errors when
        # a high-priority rule references something defined by a lower-priority
        # rule (e.g., a Schedule defined later).
        _ident_re = _re.compile(r'\b([A-Za-z_][A-Za-z0-9_]*)\b')

        def _parse_defines(code: str) -> set:
            defines = set()
            for line in code.split('\n'):
                m = _assign_re.match(line.lstrip())
                if m:
                    defines.add(m.group(1))
            return defines

        def _parse_uses(code: str) -> set:
            names = set()
            for line in code.split('\n'):
                s = line.strip()
                if not s or s.startswith('#'):
                    continue
                if '=' in line and not line.lstrip().startswith('=='):
                    rhs = line.split('=', 1)[1]
                else:
                    rhs = line
                names.update(_ident_re.findall(rhs))
            return names

        for it in items:
            it['_defs'] = _parse_defines(it['code'])
            it['_uses'] = _parse_uses(it['code'])

        n = len(items)
        indeg = [0] * n
        edges = [[] for _ in range(n)]
        for j in range(n):
            needed = items[j]['_uses'] - items[j]['_defs']
            for i in range(n):
                if i == j:
                    continue
                if items[i]['_defs'] & needed:
                    edges[i].append(j)
                    indeg[j] += 1

        import heapq as _heapq
        heap = [(items[i]['priority'], items[i]['name'], i) for i in range(n) if indeg[i] == 0]
        _heapq.heapify(heap)
        ordered_idx = []
        local_indeg = list(indeg)
        while heap:
            _, _, i = _heapq.heappop(heap)
            ordered_idx.append(i)
            for j in edges[i]:
                local_indeg[j] -= 1
                if local_indeg[j] == 0:
                    _heapq.heappush(heap, (items[j]['priority'], items[j]['name'], j))

        if len(ordered_idx) == n:
            items = [items[i] for i in ordered_idx]
        # else: cycle — fall back to the priority order already in place

        for it in items:
            it.pop('_defs', None)
            it.pop('_uses', None)

        code_blocks = [it['code'] for it in items if it.get('code')]
        combined = "\n\n".join(code_blocks)
        return {"success": True, "code": combined, "count": len(code_blocks)}
    except Exception as e:
        logger.error(f"Error generating combined code: {e}")
        raise HTTPException(status_code=500, detail=str(e))


from contextlib import asynccontextmanager

@asynccontextmanager
async def lifespan(app):
    """Application lifespan: startup and shutdown hooks."""
    logger.info("Application startup")
    yield
    client.close()
    logger.info("Application shutdown — MongoDB client closed")

# Include router under /api so frontend proxying to /api/* resolves correctly
app.include_router(api_router, prefix="/api")

# Also include the same routes at root (no prefix) for dev environments where
# the frontend proxy or external clients may strip the `/api` prefix. This
# makes the backend tolerant to both `/api/...` and `/<route>` requests and
# prevents 404s when the proxy rewrites paths unexpectedly.
app.include_router(api_router)

# Set lifespan on app
app.router.lifespan_context = lifespan

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=settings.cors_origins,
    allow_methods=["*"],
    allow_headers=["*"],
)

# WebSocket endpoint for development (supports hot reload, live updates)
@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    """WebSocket endpoint for dev client connections and hot reload"""
    await websocket.accept()
    try:
        while True:
            # Receive and echo messages to keep connection alive
            data = await websocket.receive_text()
            await websocket.send_text(data)
    except WebSocketDisconnect:
        logger.info("WebSocket client disconnected")
    except Exception as e:
        logger.error(f"WebSocket error: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host=settings.host, port=settings.port)
