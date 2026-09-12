"""Business-requirements document ingestion for the agent.

Companion to workbook.py. Where workbook.py handles Excel *models* (formulas to
translate), this module handles requirement *documents* — a PDF or Word file
describing, in prose, what the user wants built. The agent reads the extracted
text, analyses it, asks the user clarifying questions, then authors the rules.

Storage mirrors workbook.py exactly: uploaded files live on disk under
`uploads/documents/` ({document_id}.{ext}), with a {document_id}.meta.json
sidecar (original filename, sha256, extracted-text length, page/paragraph
counts) and a {document_id}.txt holding the extracted plain text. The directory
scan IS the registry — no DB dependency, so it works identically in Mongo and
in-memory mode.

Everything here is synchronous and side-effect-free apart from the explicit
save/delete entry points; the async agent tools in tools.py are thin wrappers.
No LLM access, no bridge access.
"""

from __future__ import annotations

import hashlib
import json
import logging
import uuid
from datetime import datetime, timezone
from pathlib import Path

logger = logging.getLogger(__name__)

UPLOAD_DIR = Path(__file__).parent / "uploads" / "documents"

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024      # refuse uploads beyond 20 MB
MAX_TEXT_CHARS = 400_000                    # cap stored extracted text

# extension -> canonical kind
_PDF_EXTS = {".pdf"}
_WORD_EXTS = {".docx"}
# .doc (legacy binary Word) needs a different parser we don't ship — reject it
# with a clear message rather than producing garbage text.
_LEGACY_WORD_EXTS = {".doc"}


class DocumentError(Exception):
    """User-facing failure (bad id, unsupported type, corrupt file, ...)."""


# ──────────────────────────────────────────────────────────────────────────
# Storage / registry
# ──────────────────────────────────────────────────────────────────────────

def _ext_of(filename: str) -> str:
    return Path((filename or "").strip().lower()).suffix


def _doc_path(document_id: str, ext: str) -> Path:
    return UPLOAD_DIR / f"{document_id}{ext}"


def _meta_path(document_id: str) -> Path:
    return UPLOAD_DIR / f"{document_id}.meta.json"


def _text_path(document_id: str) -> Path:
    return UPLOAD_DIR / f"{document_id}.txt"


def _read_meta(document_id: str) -> dict:
    try:
        return json.loads(_meta_path(document_id).read_text(encoding="utf-8"))
    except Exception:
        return {}


def _write_meta(document_id: str, meta: dict) -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    _meta_path(document_id).write_text(
        json.dumps(meta, indent=2, default=str), encoding="utf-8"
    )


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def list_documents() -> list[dict]:
    """Return every stored requirement document's registry entry (no text)."""
    if not UPLOAD_DIR.is_dir():
        return []
    out: list[dict] = []
    for meta_file in sorted(UPLOAD_DIR.glob("*.meta.json")):
        try:
            meta = json.loads(meta_file.read_text(encoding="utf-8"))
        except Exception:
            continue
        if meta.get("document_id"):
            # never leak the full text into a listing
            meta.pop("text", None)
            out.append(meta)
    out.sort(key=lambda m: m.get("uploaded_at") or "", reverse=True)
    return out


# ──────────────────────────────────────────────────────────────────────────
# Text extraction
# ──────────────────────────────────────────────────────────────────────────

def _extract_pdf(path: Path) -> tuple[str, int]:
    """Return (text, page_count). Raises DocumentError on unreadable PDFs."""
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - dependency guard
        raise DocumentError(
            "PDF support is not installed on the server (pypdf missing). "
            "Please upload a Word (.docx) file instead, or ask an administrator "
            "to install the PDF reader."
        ) from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentError(
            "This PDF could not be opened — it may be corrupt or password-"
            "protected. Please re-save it and try again."
        ) from exc
    if getattr(reader, "is_encrypted", False):
        # Try an empty-password decrypt (common for "protected" but not
        # truly locked PDFs); fail clearly if it doesn't work.
        try:
            reader.decrypt("")
        except Exception:
            raise DocumentError(
                "This PDF is password-protected. Please remove the password "
                "and upload it again."
            )
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n\n".join(parts).strip(), len(reader.pages)


def _extract_docx(path: Path) -> tuple[str, int]:
    """Return (text, paragraph_count). Includes table cell text."""
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - dependency guard
        raise DocumentError(
            "Word support is not installed on the server (python-docx missing). "
            "Please upload a PDF instead, or ask an administrator to install "
            "the Word reader."
        ) from exc
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise DocumentError(
            "This Word file could not be opened — it may be corrupt or an "
            "older .doc format. Save it as .docx and try again."
        ) from exc
    lines: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Flatten tables into pipe-separated rows so requirement tables survive.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines).strip(), len(doc.paragraphs)


def _extract_text(path: Path, ext: str) -> tuple[str, dict]:
    """Dispatch extraction by extension. Returns (text, unit_stats)."""
    if ext in _PDF_EXTS:
        text, pages = _extract_pdf(path)
        return text, {"pages": pages}
    if ext in _WORD_EXTS:
        text, paras = _extract_docx(path)
        return text, {"paragraphs": paras}
    raise DocumentError(f"Unsupported document type '{ext}'")


# ──────────────────────────────────────────────────────────────────────────
# Save / read / delete
# ──────────────────────────────────────────────────────────────────────────

def save_document_bytes(filename: str, content: bytes) -> dict:
    """Persist an uploaded requirement document + its extracted text.

    Returns the registry entry (without the full text). Raises DocumentError
    on any validation / extraction problem so the API can surface a clear 400.
    """
    name = (filename or "document").strip()
    ext = _ext_of(name)

    if ext in _LEGACY_WORD_EXTS:
        raise DocumentError(
            "Legacy .doc files aren't supported. Open the file in Word and "
            "use 'Save As' → Word Document (.docx), then upload the .docx."
        )
    if ext not in _PDF_EXTS and ext not in _WORD_EXTS:
        raise DocumentError(
            "Only PDF (.pdf) and Word (.docx) documents are supported."
        )
    if not content:
        raise DocumentError("Uploaded file is empty")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentError(
            f"File is {len(content) // (1024 * 1024)} MB — the limit is "
            f"{MAX_DOCUMENT_BYTES // (1024 * 1024)} MB"
        )
    # Magic-byte sniff so a mislabelled/renamed file is caught up front.
    if ext in _PDF_EXTS and not content.startswith(b"%PDF"):
        raise DocumentError(
            "This file isn't a valid PDF (missing the PDF header). It may be "
            "renamed or corrupt — re-save it as a PDF and try again."
        )
    if ext in _WORD_EXTS and content[:4] not in (
        b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"
    ):
        raise DocumentError(
            "This file isn't a valid Word .docx (it's not a Word ZIP "
            "container). Re-save it as .docx in Word and try again."
        )

    # Deduplicate identical bytes — re-uploading the same doc returns the
    # existing entry rather than minting a second id.
    sha256 = hashlib.sha256(content).hexdigest()
    for existing in list_documents():
        if existing.get("sha256") == sha256 and \
                _doc_path(existing.get("document_id", ""),
                          existing.get("ext", "")).is_file():
            existing["duplicate_of_existing"] = True
            return existing

    document_id = uuid.uuid4().hex[:12]
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    path = _doc_path(document_id, ext)
    path.write_bytes(content)

    # Extract text now (fail → clean up the orphaned file).
    try:
        text, unit_stats = _extract_text(path, ext)
    except DocumentError:
        try:
            path.unlink()
        except Exception:
            pass
        raise
    except Exception as exc:
        try:
            path.unlink()
        except Exception:
            pass
        raise DocumentError(f"Could not read the document: {exc}") from exc

    truncated = False
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        truncated = True

    if not text.strip():
        try:
            path.unlink()
        except Exception:
            pass
        raise DocumentError(
            "No readable text was found in this document. If it's a scanned "
            "PDF (images of pages), it needs to be run through OCR first, or "
            "re-typed as a text document."
        )

    _text_path(document_id).write_text(text, encoding="utf-8")

    meta = {
        "document_id": document_id,
        "filename": name,
        "ext": ext,
        "kind": "pdf" if ext in _PDF_EXTS else "word",
        "sha256": sha256,
        "size_bytes": len(content),
        "text_chars": len(text),
        "text_truncated": truncated,
        "uploaded_at": _now_iso(),
        **unit_stats,
    }
    _write_meta(document_id, meta)
    return meta


def get_meta(document_id: str) -> dict:
    meta = _read_meta(document_id)
    if not meta.get("document_id"):
        raise DocumentError(f"No document found with id '{document_id}'")
    return meta


def get_document_text(document_id: str, *, offset: int = 0,
                      limit: int = 40_000) -> dict:
    """Return a slice of the extracted text plus paging info.

    The agent reads long documents in chunks so a huge requirements PDF can't
    blow the context window in one shot.
    """
    meta = get_meta(document_id)
    try:
        full = _text_path(document_id).read_text(encoding="utf-8")
    except Exception as exc:
        raise DocumentError(
            f"The extracted text for '{meta.get('filename')}' is missing — "
            f"please re-upload the document."
        ) from exc
    offset = max(0, int(offset or 0))
    limit = max(1, min(int(limit or 40_000), 120_000))
    chunk = full[offset:offset + limit]
    end = offset + len(chunk)
    return {
        "document_id": document_id,
        "filename": meta.get("filename"),
        "kind": meta.get("kind"),
        "total_chars": len(full),
        "offset": offset,
        "returned_chars": len(chunk),
        "has_more": end < len(full),
        "next_offset": end if end < len(full) else None,
        "text": chunk,
    }


def delete_document(document_id: str) -> dict:
    meta = _read_meta(document_id)
    if not meta.get("document_id"):
        raise DocumentError(f"No document found with id '{document_id}'")
    for p in (_doc_path(document_id, meta.get("ext", "")),
              _meta_path(document_id), _text_path(document_id)):
        try:
            p.unlink()
        except Exception:
            pass
    return {"deleted": True, "document_id": document_id,
            "filename": meta.get("filename")}
